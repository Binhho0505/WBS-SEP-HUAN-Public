"""Native app tests; synthetic fixtures explicitly cannot be CC1 prices or approval."""
import base64,copy,io,json,tempfile,threading,unittest,urllib.request,urllib.error
from decimal import Decimal
from pathlib import Path
from openpyxl import load_workbook
from domain import new_project,questions,status,calculate,digest,catalog,upgrade,needed_price_keys
from server import App,imported
from store import Store,Conflict
import ai

def fixture():
    s=new_project('TEST ONLY - NOT REAL PROJECT');s['context']={k:'TEST ONLY' for k in ['project_code','legal_entity','contract_ref','baseline','location','owner','reviewer','legal_ref','rounding_ref']};s['context'].update(period='2026-10',currency='VND',cost_basis='ESTIMATE',money_places=2,estimate_purpose='INTERNAL_COST',relation_profile='ESTIMATE_MIN',rounding_level='TOTAL')
    s['boq']=[dict(id='test-source',source_code='1.1',description='TEST ONLY concrete',quantity='200',uom='m3',unit_price='100',kind='MEASURED',source_ref='TEST_ONLY:1',source_hash='a'*64)]
    wr=next(r['business_code'] for r in catalog()['WORK_RESULTS'] if r['business_code'].startswith('H') and r['readiness']!='HOLD' and r['node_kind']=='LEAF_CANDIDATE')
    s['wbs']=[dict(id='TEST-WP',work_result_code=wr,lbs='TEST-LBS',scope='TEST scope in/out',acceptance='TEST acceptance',ca='TEST-CA',cam='TEST-CAM')]
    s['mapping']=[dict(id='test-item',source_line_id='test-source',quantity='200',uom='m3',weight='1',wbs_id='TEST-WP',mapping_ref='TEST mapping evidence',technical_spec='TEST spec',method='TEST method',norm_match_ref='TEST condition check')]
    s['applications']=[dict(id='TEST-APP',item_id='test-item',norm_id='TEST-RECIPE',qty_factor='1')]
    s['norms']=[dict(id='TEST-RECIPE',code='TEST-CUSTOM-NOT-TT38',kind='CC1_APPROVED_CUSTOM',edition='TEST ONLY',uom='m3',denominator='100',conditions='TEST ONLY conditions',source_ref='TEST norm source',selection_ref='TEST local check',tt37_ref='TEST-only method reference')]
    s['resources']=[dict(id='TEST-RESOURCE-LINE',norm_id='TEST-RECIPE',line_type='QTY',resource_code='TEST-MAT',name='TEST ONLY material',spec='TEST ONLY SPEC',uom='t',coefficient='1.5',cbs='C10',source_ref='TEST source coefficient')];s['resource_completeness_ref']='TEST all resource rows checked'
    s['prices']=[dict(id='TEST-PRICE',resource_code='TEST-MAT',spec='TEST ONLY SPEC',uom='t',unit_price='10.01',location='TEST ONLY',period='2026-10',currency='VND',tax_basis='EX_VAT',delivery='AT_SITE',included='TEST transport included; not a real price',source_ref='TEST PRICE',price_ref='TEST CHECK')]
    s['adjustments']={k:'TEST ONLY' for k in ['cbs_scheme_ref','cost_boundary_ref','tax_ref','completeness_ref']};s['adjustments']['allocation_driver']='DIRECT_PRORATA';s['adjustments']['rules']=[dict(id='TEST_OH',name='TEST overhead',method='PERCENT',basis='DIRECT',value='10',source_ref='TEST',nature='COST',cbs='C50'),dict(id='TEST_VAT',name='TEST VAT',method='PERCENT',basis='DIRECT+TEST_OH',value='8',source_ref='TEST',nature='TAX')]
    s['export']=dict(prepared_by='TEST ONLY',checked_by='TEST ONLY',review_ref='TEST ONLY',template_id='TPL03_COMPOSITION',local_draft_ack=True)
    return s
def composite():
    """TEST ONLY: one item composed from two recipes (concrete + rebar with 0.12 t/m3), PERCENT line and a CBS-based rule."""
    s=fixture()
    s['norms'].append(dict(s['norms'][0],id='TEST-REBAR',code='TEST-REBAR-CUSTOM',uom='t',denominator='1'))
    s['applications'].append(dict(id='TEST-APP-2',item_id='test-item',norm_id='TEST-REBAR',qty_factor='0.12',factor_ref='TEST drawing schedule',conversion_ref='TEST t per m3'))
    s['resources']+= [dict(id='TEST-STEEL',norm_id='TEST-REBAR',line_type='QTY',resource_code='TEST-STEEL',name='TEST steel',spec='TEST D16',uom='kg',coefficient='1020',cbs='C10',source_ref='TEST'),dict(id='TEST-LAB',norm_id='TEST-REBAR',line_type='QTY',resource_code='TEST-LAB',name='TEST labour',spec='TEST 3.5/7',uom='công',coefficient='10',cbs='C20',source_ref='TEST'),dict(id='TEST-OTHER',norm_id='TEST-REBAR',line_type='PERCENT',resource_code='',name='TEST vật liệu khác',percent='2',percent_of='C10',cbs='C10',source_ref='TEST')]
    s['prices']+= [dict(s['prices'][0],id='TEST-P2',resource_code='TEST-STEEL',spec='TEST D16',uom='kg',unit_price='0.02'),dict(s['prices'][0],id='TEST-P3',resource_code='TEST-LAB',spec='TEST 3.5/7',uom='công',unit_price='0.5')]
    return s
class DomainTests(unittest.TestCase):
    def test_all_steps_and_arithmetic(self):
        s=fixture()
        for i in range(8):self.assertEqual(questions(s,i),[],str(i))
        r=calculate(s);self.assertEqual(r['direct_cost'],'30.030');self.assertEqual(r['total_unrounded'],'35.67564');self.assertEqual(r['total'],'35.68')
    def test_no_price_no_total_but_known_amount(self):
        s=composite();s['prices'][2]['unit_price']=None;self.assertTrue(questions(s,5));r=calculate(s)
        self.assertIsNone(r['total']);self.assertFalse(r['complete']);self.assertEqual(r['coverage'],{'needed':4,'priced':3});self.assertEqual(Decimal(r['known_amount']),Decimal('30.03')+Decimal('489.6')*Decimal('1.02'))
        s['prices'][1]['unit_price']=None;self.assertEqual(calculate(s)['coverage']['priced'],1)
    def test_wrong_price_context_and_spec(self):
        s=fixture();s['prices'][0].update(period='2025-01',spec='WRONG',tax_basis='VAT_INCLUDED');self.assertTrue(questions(s,5))
        s=fixture();s['prices'][0].update(period='2025-01',tax_basis='VAT_INCLUDED');self.assertGreaterEqual(len(questions(s,5)),2)
    def test_composition_percent_line_and_cbs_basis(self):
        s=composite();s['adjustments']['rules']=[dict(id='LAB_SITE',name='TEST % on labour',method='PERCENT',basis='C20',value='5',source_ref='TEST',nature='COST',cbs='C50')]
        for i in range(8):self.assertEqual(questions(s,i),[],str(i))
        r=calculate(s);steel=Decimal(200)*Decimal('0.12')*1020*Decimal('0.02');labour=Decimal(200)*Decimal('0.12')*10*Decimal('0.5')
        self.assertEqual(Decimal(r['cbs']['C10']),Decimal('30.03')+steel+steel*Decimal('0.02'));self.assertEqual(Decimal(r['rules'][0]['amount']),labour*Decimal('0.05'))
        self.assertEqual(sum(1 for l in r['lines'] if l['line_type']=='PERCENT'),1)
    def test_composition_requires_factor_and_conversion_refs(self):
        s=composite();s['applications'][1].pop('factor_ref');s['applications'][1].pop('conversion_ref');qs=[x['field'] for x in questions(s,3)]
        self.assertIn('applications[1].factor_ref',qs);self.assertIn('applications[1].conversion_ref',qs)
        s=composite();s['applications'].append(dict(s['applications'][1],id='dup'));self.assertTrue(questions(s,3))
    def test_uom_typography_is_not_a_conversion(self):
        from domain import same_uom
        self.assertTrue(same_uom('m³','m3') and same_uom('M2','m²') and same_uom(' tấn.','tấn'))
        self.assertFalse(same_uom('tấn','kg') or same_uom('m3','m2') or same_uom('100m3','m3'))
        s=fixture();s['mapping'][0]['uom']='m³';s['boq'][0]['uom']='m³';self.assertFalse(any(q['field'].endswith('conversion_ref') for q in questions(s,3)))
        s['mapping'][0]['uom']='m2';self.assertTrue(any(q['field'].endswith('conversion_ref') for q in questions(s,3)))
    def test_percent_line_needs_group(self):
        s=composite();s['resources'][-1]['percent_of']='C30';self.assertTrue(any('C30' in q['question'] for q in questions(s,4)))
    def test_price_zone_from_wbs(self):
        s=fixture();s['wbs'][0]['price_zone']='TEST ZONE B';self.assertIn(('TEST-MAT','TEST ONLY SPEC','t','TEST ZONE B'),needed_price_keys(s));self.assertTrue(questions(s,5))
        s['prices'][0]['location']='TEST ZONE B';self.assertEqual(questions(s,5),[])
    def test_earlier_period_needs_reference(self):
        s=fixture();s['prices'][0]['period']='2026-07';self.assertTrue(any(q['field'].endswith('period_ref') for q in questions(s,5)));self.assertFalse(calculate(s)['complete'])
        s['prices'][0]['period_ref']='TEST annual labour list carried forward';self.assertEqual(questions(s,5),[]);self.assertEqual(calculate(s)['direct_cost'],'30.030')
        s['prices'][0]['period']='2026-12';self.assertTrue(questions(s,5))  # a later period can never be carried back
    def test_library_search_accent_insensitive(self):
        with tempfile.TemporaryDirectory() as d:
            store=Store(Path(d)/'db.sqlite');store.library_add('PRICES','TEST ONLY','a'*64,'t.csv','',[{'id':'1','name':'Thép thanh vằn Ø14 CB300','unit_price':'1'},{'id':'2','name':'Xi măng PCB40','unit_price':'2'}])
            store.library_add('REFERENCE','TEST ONLY ref','b'*64,'r.csv','',[{'id':'r1','series':'TT36_RATE','value':'6.0'}])
            self.assertEqual([r['id'] for r in store.library_search('PRICES','thep van cb300')['rows']],['1'])
            self.assertEqual([r['id'] for r in store.library_search('PRICES','phi14')['rows']],['1'])
            store.library_add('PRICES','TEST ONLY NC','c'*64,'n.csv','',[{'id':'n1','name':'Nhân công nhóm 1','location':'Vùng III','unit_price':'248737'},{'id':'n2','name':'Nhân công nhóm 2','location':'Vùng II','unit_price':'298012'}])
            self.assertEqual([r['id'] for r in store.library_search('PRICES','nhan cong nhom 2 vung ii')['rows']],['n2'])
            self.assertEqual(store.library_search('PRICES','xi mang')['total'],1);self.assertEqual(store.library_search('REFERENCE','tt36_rate')['total'],1)
    def test_duplicate_price_key_ambiguous(self):
        s=fixture();s['prices'].append(dict(s['prices'][0],id='TEST-PRICE-2',unit_price='11'));self.assertTrue(questions(s,5));self.assertFalse(calculate(s)['complete'])
    def test_line_rounding_and_allocation_comparison(self):
        s=fixture();s['context'].update(rounding_level='LINE',money_places=0);r=calculate(s)
        self.assertEqual(r['lines'][0]['amount'],'30');self.assertEqual(r['total'],'36')
        item=r['items'][0];self.assertEqual(item['contract_amount'],'20000');self.assertEqual(Decimal(item['allocated_indirect']),Decimal('3'));self.assertEqual(Decimal(item['margin']),Decimal(20000)-Decimal(33))
    def test_purpose_mix_requires_reference(self):
        s=fixture();s['context']['estimate_purpose']='STATE_NORM_CHECK';self.assertTrue(any(q['field']=='context.norm_mix_ref' for q in questions(s,3)))
        s['context']['norm_mix_ref']='TEST reason';self.assertEqual(questions(s,3),[])
    def test_profile_full_requires_relations(self):
        s=fixture();self.assertEqual(questions(s,2),[]);s['context']['relation_profile']='FULL';self.assertTrue(any(q['field']=='wbs[0].bim' for q in questions(s,2)))
    def test_unassigned_item_and_split(self):
        s=fixture();s['mapping'][0]['wbs_id']='';self.assertTrue(any(q['field']=='mapping[0].wbs_id' for q in questions(s,2)))
        s=fixture();s['mapping']*=2;s['mapping'][1]=dict(s['mapping'][0],id='different',weight='.5',quantity='100');self.assertTrue(questions(s,2))
    def test_source_coverage_cannot_drop(self):
        s=fixture();s['mapping']=[];self.assertTrue(questions(s,2))
    def test_master_conflict(self):
        s=fixture();s['wbs'][0]['work_result_code']='I7010.50';self.assertTrue(any('HOLD' in r['question'] for r in questions(s,2)))
    def test_atomic_group_rejected(self):
        s=fixture();s['norms'][0].update(kind='TT38_ATOMIC',code='AF.65100',table_verified=True);self.assertTrue(questions(s,3))
    def test_conditions_required(self):
        s=fixture();s['mapping'][0]['technical_spec']='';self.assertTrue(questions(s,3))
    def test_cycle_and_reserved_rule(self):
        s=fixture();s['adjustments']['rules'][0].update(id='DIRECT',basis='DIRECT+LATER');self.assertTrue(questions(s,6))
        s=fixture();s['adjustments']['rules'][0]['id']='C20';self.assertTrue(questions(s,6))
    def test_dirty_predecessor_blocks_all(self):
        s=fixture();s['confirmed']=list(range(8));s['context']['legal_ref']='';self.assertFalse(any(r['confirmed'] for r in status(s)))
    def test_upgrade_from_010(self):
        s=fixture();old=copy.deepcopy(s);old['schema_version']='0.1.0';old.pop('wbs');old.pop('applications');old['confirmed']=list(range(8))
        old['mapping'][0].update(norm_id='TEST-RECIPE',**{k:v for k,v in s['wbs'][0].items() if k!='id'});old['prices'][0]=dict(old['prices'][0],resource_id='TEST-RESOURCE-LINE');old['prices'][0].pop('resource_code');[r.pop('line_type') for r in old['resources']];[r.pop('resource_code') for r in old['resources']]
        u=upgrade(old);self.assertEqual(u['confirmed'],[]);self.assertEqual(u['wbs'][0]['id'],'TEST-WP');self.assertEqual(u['applications'][0]['norm_id'],'TEST-RECIPE');self.assertEqual(u['prices'][0]['resource_code'],'TEST-RESOURCE-LINE')
        self.assertEqual(calculate(u)['direct_cost'],'30.030')
    def test_import_csv_preserves_raw_hash(self):
        b='Code,Description,Qty,UOM\nA,=UNTRUSTED(),3,m\n'.encode();r,_=imported(dict(filename='test.csv',data=base64.b64encode(b).decode(),columns={'source_code':'Code','description':'Description','quantity':'Qty','uom':'UOM'}));self.assertEqual(r['rows'][0]['raw']['Description'],'=UNTRUSTED()');self.assertEqual(len(r['sha256']),64)
    def test_import_xlsx_keeps_formula_and_warns(self):  # D79 FR02: formula kept beside the cached value; subtotal / duplicate code warned, nothing dropped
        from openpyxl import Workbook
        wb=Workbook();ws=wb.active;ws.title='BOQ'
        for r in (['Code','Description','Qty','UOM'],['A','TEST ONLY a',3,'m'],['B','TEST ONLY b','=1+1','m'],['A','TEST ONLY tổng','=SUM(C2:C3)','m']):ws.append(r)
        bio=io.BytesIO();wb.save(bio)
        r,_=imported(dict(filename='test.xlsx',data=base64.b64encode(bio.getvalue()).decode(),columns={'source_code':'Code','description':'Description','quantity':'Qty','uom':'UOM'}))
        rows=r['rows'];self.assertEqual(len(rows),3);self.assertEqual(rows[1]['raw_formulas'],{'Qty':'=1+1'});self.assertNotIn('raw_formulas',rows[0])
        self.assertIsNone(rows[1]['quantity'])  # no cached value in the file → stays empty, never computed by the app
        self.assertTrue(any('SUM(' in w and 'Dòng 4' in w for w in r['warnings']));self.assertTrue(any('Mã A xuất hiện 2 dòng' in w for w in r['warnings']))
    def test_store_conflict_and_restart(self):
        with tempfile.TemporaryDirectory() as d:
            store=Store(Path(d)/'db.sqlite');s=store.create('TEST');store.put(s,0,'TEST')
            with self.assertRaises(Conflict):store.put(s,0,'STALE')
            self.assertEqual(Store(Path(d)/'db.sqlite').get(s['id'])['revision'],1)
    def test_anthropic_sdk_request_shape_and_refusal(self):
        """Real Anthropic SDK against a loopback mock: fallbacks default + beta header + effort; refusal raises."""
        from http.server import BaseHTTPRequestHandler,ThreadingHTTPServer
        seen=[];reply={'stop_reason':'end_turn'}
        class Mock(BaseHTTPRequestHandler):
            def log_message(self,*a):pass
            def do_POST(self):
                body=json.loads(self.rfile.read(int(self.headers['Content-Length'])));seen.append((self.path,dict(self.headers),body))
                data=json.dumps({'id':'msg_test','type':'message','role':'assistant','model':body['model'],'content':[{'type':'text','text':'TEST ONLY proposal'}],'stop_reason':reply['stop_reason'],'stop_sequence':None,'usage':{'input_tokens':1,'output_tokens':1}}).encode()
                self.send_response(200);self.send_header('Content-Type','application/json');self.send_header('Content-Length',str(len(data)));self.end_headers();self.wfile.write(data)
        server=ThreadingHTTPServer(('127.0.0.1',0),Mock);threading.Thread(target=server.serve_forever,daemon=True).start()
        try:
            import anthropic,unittest.mock as um
            real=anthropic.Anthropic
            with um.patch.object(anthropic,'Anthropic',lambda **k:real(base_url=f'http://127.0.0.1:{server.server_port}',max_retries=0,**k)):
                text,model=ai.call_anthropic({'api_key':'sk-ant-TESTONLY'},'claude-opus-5-5','SYS','USER')
                self.assertEqual(text,'TEST ONLY proposal');path,headers,body=seen[-1]
                self.assertIn('server-side-fallback-2026-07-01',headers.get('anthropic-beta',''));self.assertEqual(body['fallbacks'],'default');self.assertEqual(body['output_config'],{'effort':'medium'});self.assertNotIn('thinking',body)
                ai.call_anthropic({'api_key':'sk-ant-TESTONLY'},'claude-haiku-5-5','SYS','USER');self.assertNotIn('fallbacks',seen[-1][2])
                reply['stop_reason']='refusal'
                with self.assertRaises(ai.AIError):ai.call_anthropic({'api_key':'sk-ant-TESTONLY'},'claude-opus-5-5','SYS','USER')
        finally:server.shutdown();server.server_close()
    def test_ai_prompt_minimised(self):
        s=fixture();system,user=ai.build_prompt(s,0,'TEST question',True);self.assertNotIn('"owner"',user);self.assertIn('TEST question',user);self.assertIn('Không bịa',system)

class HttpTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.app=App(('127.0.0.1',0),Path(self.tmp.name)/'db.sqlite');self.thread=threading.Thread(target=self.app.serve_forever,daemon=True);self.thread.start();self.base=f'http://127.0.0.1:{self.app.server_port}';self.token=self.call('/api/bootstrap')['token']
    def tearDown(self):self.app.shutdown();self.app.server_close();self.tmp.cleanup()
    def call(self,path,body=None,token=True):
        req=urllib.request.Request(self.base+path,data=None if body is None else json.dumps(body).encode(),headers={} if body is None else {'Content-Type':'application/json','X-CC1-Session':self.app.token if token else 'WRONG'})
        with urllib.request.urlopen(req) as r:return json.load(r)
    def test_reject_skip_export_and_csrf(self):
        s=self.call('/api/projects',{'name':'TEST ONLY'})['project']
        for path,body in [(f'/api/projects/{s["id"]}/stage/7',{'revision':0,'data':{},'confirm':True}),(f'/api/projects/{s["id"]}/export',{'revision':0})]:
            with self.assertRaises(urllib.error.HTTPError) as e:self.call(path,body)
            self.assertEqual(e.exception.code,409)
        with self.assertRaises(urllib.error.HTTPError) as e:self.call('/api/projects',{'name':'TEST'},False)
        self.assertEqual(e.exception.code,403)
    def test_boq_manual_edit_and_line_keep_mapping(self):
        """D73: BOQ code/description/price/quantity editable by hand with a reason (raw/hash/source kept, original_* server-side);
        a typed line gets a server hash; adding/editing lines keeps WBS mapping and norm applications of the other lines."""
        f=fixture();s=self.save_all(f);pid=s['id'];stage=lambda i,data,**kw:self.call(f'/api/projects/{pid}/stage/{i}',dict({'revision':s['revision'],'data':data,'confirm':True},**kw))
        boq=copy.deepcopy(s['boq']);boq[0].update(source_code='1.1A',description='TEST ONLY concrete B30 (sửa tay)',unit_price='120',quantity='210')
        d=stage(1,boq);self.assertFalse(d['confirmed']);r=d['project']['boq'][0]
        self.assertEqual((r['original_source_code'],r['original_description'],r['original_unit_price'],r['original_quantity']),('1.1','TEST ONLY concrete','100','200'))
        self.assertEqual((r['raw'] if 'raw' in r else None,r['source_hash'],r['source_ref']),(None,'a'*64,'TEST_ONLY:1'))
        self.assertTrue(any(q['field']=='boq[0].disposition_ref' and 'sửa tay' in q['question'] for q in d['questions']))
        m=d['project']['mapping'];self.assertEqual((len(m),m[0]['wbs_id'],m[0]['quantity']),(1,'TEST-WP','210'))  # mapping kept, unsplit quantity follows
        self.assertEqual(len(d['project']['applications']),1);s=d['project']
        boq=copy.deepcopy(s['boq']);boq[0]['disposition_ref']='TEST phụ lục 01'
        boq.append({'id':'manual-1','source_code':'1.2','description':'TEST ONLY dòng thêm tay','kind':'MEASURED','quantity':'5','uom':'m3','unit_price':'90','source_ref':'MANUAL:TEST','source_hash':'b'*64,'raw':{'x':1}})
        d=stage(1,boq);r2=d['project']['boq'][1];self.assertEqual(r2['raw']['MANUAL'],True);self.assertNotEqual(r2['source_hash'],'b'*64);self.assertEqual(len(r2['source_hash']),64)
        self.assertTrue(any(q['field']=='boq[1].disposition_ref' and 'nhập tay' in q['question'] for q in d['questions']))
        m={x['source_line_id']:x for x in d['project']['mapping']};self.assertEqual((m['test-source']['wbs_id'],m['manual-1']['wbs_id'],len(d['project']['applications'])),('TEST-WP','',1))
        s=d['project'];boq=copy.deepcopy(s['boq']);boq[1]['disposition_ref']='TEST biên bản bổ sung';d=stage(1,boq);self.assertTrue(d['confirmed'],d['questions']);s=d['project']
        boq=copy.deepcopy(s['boq']);boq[0]['source_hash']='c'*64
        with self.assertRaises(urllib.error.HTTPError) as e:stage(1,boq)
        self.assertEqual(e.exception.code,400)
        boq=copy.deepcopy(s['boq']);boq.append({'id':'forged','description':'x','kind':'MEASURED','quantity':'1','uom':'m3','source_ref':'X'})
        with self.assertRaises(urllib.error.HTTPError) as e:stage(1,boq)
        self.assertEqual(e.exception.code,400)
        d=stage(1,[x for x in s['boq'] if x['id']!='manual-1']);self.assertEqual([x['source_line_id'] for x in d['project']['mapping']],['test-source'])  # removed line → its item goes
    def test_restore_resets_confirmations(self):
        f=fixture();f['confirmed']=list(range(8));s=self.call('/api/restore',{'schema_version':'0.2.0','project':f,'audit':[]})['project']
        self.assertEqual(s['confirmed'],[]);self.assertNotEqual(s['id'],f['id']);self.assertEqual(s['mapping'],f['mapping']);self.assertEqual(s['wbs'],f['wbs'])
    def save_all(self,f):
        s=self.call('/api/projects',{'name':f['name']})['project']
        for i,key in enumerate(['context','boq','mapping','norms','resources','prices','adjustments','export']):
            data={'revision':s['revision'],'data':f[key],'confirm':True}
            if i==2:data['wbs']=f['wbs']
            if i==3:data['applications']=f['applications'];data['assignments']=[{k:m.get(k) for k in ['id','technical_spec','method','norm_match_ref']} for m in f['mapping']]
            if i==4:data['resource_completeness_ref']=f['resource_completeness_ref']
            d=self.call(f'/api/projects/{s["id"]}/stage/{i}',data);self.assertTrue(d['confirmed'],d.get('questions'));s=d['project']
        return s
    def test_full_persisted_flow_and_export(self):
        s=self.save_all(fixture())
        export=self.call(f'/api/projects/{s["id"]}/export',{'revision':s['revision']});same=self.call(f'/api/projects/{s["id"]}/export',{'revision':s['revision']});self.assertEqual(export,same)
        with urllib.request.urlopen(self.base+export['url']) as r:binary=r.read()
        wb=load_workbook(io.BytesIO(binary),data_only=False);self.assertIn('CBS_CC1',wb.sheetnames);self.assertIn('SO_SANH_HOP_DONG',wb.sheetnames);self.assertIn('DINH_MUC_AP_DUNG',wb.sheetnames)
        tong={r[0].value:r[1].value for r in wb['TONG_HOP'].iter_rows(min_row=2)};self.assertEqual(tong['Tổng sau làm tròn'],'35.68');self.assertEqual(wb['10_WP_BOQ_COMPOSITION']['H5'].value,'20000')
        self.assertFalse(any(c.data_type=='f' for ws in wb for row in ws for c in row));wb.close()
        d=self.call(f'/api/projects/{s["id"]}/stage/0',{'revision':s['revision'],'data':s['context'],'confirm':False});self.assertEqual(d['project']['confirmed'],[])
        with self.assertRaises(urllib.error.HTTPError):self.call(f'/api/projects/{s["id"]}/export',{'revision':d['project']['revision']})
    def test_library_dataset_and_price_pull(self):
        f=fixture();csv='id,resource_code,spec,uom,unit_price,location,period,currency,tax_basis,delivery,included,source_ref,price_ref\nL1,TEST-MAT,TEST ONLY SPEC,t,10.01,TEST ONLY,2026-10,VND,EX_VAT,AT_SITE,TEST,TEST LIB,TEST\nL2,OTHER,X,t,5,TEST ONLY,2026-10,VND,EX_VAT,AT_SITE,TEST,TEST,TEST\n'
        d=self.call('/api/library',{'kind':'PRICES','name':'TEST ONLY price list','filename':'p.csv','data':base64.b64encode(csv.encode()).decode()})
        self.assertEqual((d['version'],d['row_count']),(1,2));again=self.call('/api/library',{'kind':'PRICES','name':'TEST ONLY price list','filename':'p.csv','data':base64.b64encode(csv.encode()).decode()});self.assertEqual(again['version'],2)
        f['prices']=[];s=self.call('/api/projects',{'name':'TEST'})['project']
        for i,key in enumerate(['context','boq','mapping','norms','resources']):
            data={'revision':s['revision'],'data':f[key],'confirm':True,'wbs':f['wbs'] if i==2 else None,'applications':f['applications'] if i==3 else None,'resource_completeness_ref':f['resource_completeness_ref']}
            if i==3:data['assignments']=[{k:m.get(k) for k in ['id','technical_spec','method','norm_match_ref']} for m in f['mapping']]
            s=self.call(f'/api/projects/{s["id"]}/stage/{i}',data)['project']
        pulled=self.call(f'/api/projects/{s["id"]}/library-prices',{'dataset':d['id']});self.assertEqual([r['id'] for r in pulled['rows']],['L1']);self.assertEqual(pulled['rows'][0]['library_sha256'],d['sha256'])
        stubs=self.call(f'/api/projects/{s["id"]}/price-stubs',{});self.assertEqual(len(stubs['rows']),1);self.assertIsNone(stubs['rows'][0]['unit_price'])
        self.assertEqual(len(self.call('/api/library')['datasets']),2)
    def code_library_fixture(self):
        """TEST ONLY library: one TT38-like norm whose lines carry CC1 master codes, a resource master and prices."""
        st=self.app.store
        st.library_add('NORMS','TEST ONLY TT38','1'*64,'n.json','',[{'id':'TT38-2026:TS.00011','code':'TS.00011','kind':'TT38_ATOMIC','edition':'TEST','uom':'m3','denominator':'1','conditions':'TEST ONLY bê tông thử','source_ref':'TEST','selection_ref':'','table_verified':False,'parse_status':'PARSED','parse_flags':''}])
        st.library_add('RESOURCES','TEST ONLY TT38 hao phí','2'*64,'r.json','',[
            {'id':'TT38-2026:TS.00011:1','norm_id':'TT38-2026:TS.00011','line_type':'QTY','resource_code':'TT38.C10.XI','master_code':'VL.XIMANG.TEST01','name':'Xi măng','spec':'TEST PCB40','uom':'kg','coefficient':'300','cbs':'C10','source_ref':'TEST'},
            {'id':'TT38-2026:TS.00011:2','norm_id':'TT38-2026:TS.00011','line_type':'QTY','resource_code':'TT38.C20.NC2','master_code':'NC.N2','name':'Nhân công nhóm 2','spec':'Nhóm 2','uom':'công','coefficient':'1','cbs':'C20','source_ref':'TEST'},
            {'id':'TT38-2026:TS.00011:3','norm_id':'TT38-2026:TS.00011','line_type':'QTY','resource_code':'TT38.C30.MAY','master_code':'M999.0001','name':'Máy thử','spec':'TEST máy','uom':'ca','coefficient':'0.1','cbs':'C30','source_ref':'TEST'}])
        prices=st.library_add('PRICES','TEST ONLY giá','3'*64,'p.json','',[
            {'id':'P-XI','resource_code':'X.TEST','name':'Xi măng PCB40 TEST','uom':'kg','unit_price':'2','location':'TEST TỈNH','period':'2026-10','currency':'VND','tax_basis':'EX_VAT','delivery':'AT_SITE'},
            {'id':'P-NC','resource_code':'NC.N2','name':'Nhân công nhóm 2','uom':'công','unit_price':'300','location':'TEST TỈNH - Vùng II','period':'2026-07','currency':'VND','tax_basis':'EX_VAT','delivery':'AT_SITE'},
            {'id':'P-NC-LATER','resource_code':'NC.N2','name':'Nhân công nhóm 2','uom':'công','unit_price':'999','location':'TEST TỈNH - Vùng II','period':'2026-12','currency':'VND','tax_basis':'EX_VAT','delivery':'AT_SITE'}])
        st.library_add('REFERENCE','Danh mục nguồn lực CC1 TEST ONLY','4'*64,'m.json','',[{'id':'m1','series':'RESOURCE_MASTER','code':'VL.XIMANG.TEST01','name':'Xi măng','status':'STRONG','value':'Xi măng PCB40 TEST','candidates':[{'price_id':'P-XI'}]}])
    def test_code_library_pull_persists_resources_and_auto_prices(self):
        self.code_library_fixture();f=fixture();f['context']['location']='TEST TỈNH - Vùng II';s=self.call('/api/projects',{'name':'TEST ONLY codelib'})['project']
        for i,key in enumerate(['context','boq','mapping']):
            s=self.call(f'/api/projects/{s["id"]}/stage/{i}',{'revision':s['revision'],'data':f[key],'confirm':True,'wbs':f['wbs'] if i==2 else None})['project']
        hits=self.call('/api/code-library/search?q=ts.00011');self.assertEqual(hits['rows'][0]['code'],'TS.00011')
        detail=self.call('/api/code-library/norm?id=TT38-2026:TS.00011');self.assertEqual(detail['lines'][0]['map']['status'],'STRONG')
        x=self.call(f'/api/projects/{s["id"]}/code-norm',{'id':'TT38-2026:TS.00011'});self.assertEqual({r['resource_code'] for r in x['resources']},{'VL.XIMANG.TEST01','NC.N2','M999.0001'})
        norm=dict(x['norm'],selection_ref='TEST check',table_verified=True,code='TS.00011',kind='CC1_APPROVED_CUSTOM',tt37_ref='TEST')
        s=self.call(f'/api/projects/{s["id"]}/stage/3',{'revision':s['revision'],'data':[norm],'confirm':True,'applications':[{'id':'a1','item_id':f['mapping'][0]['id'],'norm_id':norm['id'],'qty_factor':'1'}],'assignments':[{k:m.get(k) for k in ['id','technical_spec','method','norm_match_ref']} for m in f['mapping']],'resources':x['resources']})['project']
        self.assertEqual(len(s['resources']),3,'resources sent with step 3 must persist (library pull regression)')
        s=self.call(f'/api/projects/{s["id"]}/stage/4',{'revision':s['revision'],'data':s['resources'],'confirm':True,'resource_completeness_ref':'TEST'})['project']
        auto=self.call(f'/api/projects/{s["id"]}/auto-prices',{})
        got={r['resource_code']:r for r in auto['rows']};self.assertEqual(set(got),{'VL.XIMANG.TEST01','NC.N2'})
        self.assertEqual(got['NC.N2']['unit_price'],'300');self.assertTrue(got['NC.N2']['mapped_from']['earlier_period']);self.assertEqual(got['VL.XIMANG.TEST01']['location'],'TEST TỈNH - Vùng II')
        self.assertEqual([u['resource_code'] for u in auto['unresolved']],['M999.0001'])
        d=self.call(f'/api/projects/{s["id"]}/stage/5',{'revision':s['revision'],'data':auto['rows'],'confirm':True});self.assertFalse(d['confirmed'])
        self.assertTrue(any(q['field'].endswith('price_ref') for q in d['questions']) and any(q['field'].endswith('period_ref') for q in d['questions']))
    def suggest_fixture(self):
        """TEST ONLY TT38-like table TS.651 (two columns) plus decoys: under-water, demolition, other unit, survey appendix."""
        n=lambda code,cond,uom='t',ed='TEST PL3':{'id':'TT38-2026:'+code,'code':code,'kind':'TT38_ATOMIC','edition':ed,'uom':uom,'denominator':'1','conditions':cond,'source_ref':'TEST','parse_status':'PARSED','parse_flags':''}
        self.app.store.library_add('NORMS','TEST ONLY TT38 gợi ý','5'*64,'n.json','',[
            n('TS.65110','Công tác sản xuất lắp dựng cốt thép bệ trụ cầu trên cạn; Đường kính ≤18mm'),n('TS.65120','Công tác sản xuất lắp dựng cốt thép bệ trụ cầu trên cạn; Đường kính >18mm'),
            n('TS.65210','Công tác sản xuất lắp dựng cốt thép bệ trụ cầu dưới nước; Đường kính ≤18mm'),n('TS.21110','Phá dỡ kết cấu cốt thép bệ trụ cầu'),
            n('TS.11110','Bê tông bệ trụ cầu trên cạn',uom='m3'),n('TS.00110','Khảo sát cốt thép bệ trụ cầu trên cạn',ed='TEST PL1')])
    def test_suggest_norms_ranks_candidates_without_changing_project(self):
        import codelib
        self.suggest_fixture();st=self.app.store
        c=codelib.suggest(st,'Cốt thép bệ trụ','Kết cấu phần dưới > Trụ cầu trên cạn','tấn',5)
        self.assertEqual(c[0]['code'][:6],'TS.651');self.assertEqual([x['code'] for x in c[0]['columns']],['TS.65110','TS.65120']);self.assertTrue(c[0]['unit_match'])
        rank={x['code'][:6]:i for i,x in enumerate(c)}
        for decoy in ('TS.652','TS.211','TS.111','TS.001'):self.assertGreater(rank.get(decoy,99),0,decoy)
        self.assertFalse(next(x for x in c if x['code'].startswith('TS.111'))['unit_match'])
        self.assertEqual(codelib.suggest(st,'','',''),[]);self.assertLessEqual(len(codelib.suggest(st,'cốt thép bệ trụ cầu','','t',2)),2)
        boq=[{'source_code':'1','description':'Cầu'},{'source_code':'1.2','description':'Trụ cầu trên cạn'},{'source_code':'1.2.1','description':'Cốt thép bệ trụ'}]
        self.assertEqual(codelib.boq_path(boq,boq[2]),'Cầu > Trụ cầu trên cạn')
        f=fixture();f['boq']=[dict(f['boq'][0],id='h',source_code='1',kind='HEADING',description='Trụ cầu trên cạn',quantity=None,uom=''),dict(f['boq'][0],source_code='1.1',description='Cốt thép bệ trụ',uom='tấn')]
        f['mapping'][0]['uom']='tấn';s=self.call('/api/projects',{'name':'TEST ONLY gợi ý'})['project']
        for i,key in enumerate(['context','boq']):s=self.call(f'/api/projects/{s["id"]}/stage/{i}',{'revision':s['revision'],'data':f[key],'confirm':True})['project']
        self.assertTrue(s['mapping']);d=self.call(f'/api/projects/{s["id"]}/suggest-norms',{'limit':3})
        items=d['items'];self.assertTrue(items);self.assertEqual(items[0]['path'],'Trụ cầu trên cạn');self.assertEqual(items[0]['candidates'][0]['code'][:6],'TS.651')
        self.assertEqual(self.call(f'/api/projects/{s["id"]}/suggest-norms',{'item_ids':['nope']})['items'],[])
        self.assertEqual(self.call(f'/api/projects/{s["id"]}')['project']['revision'],s['revision'],'suggestion must not write the project')
    def test_suggest_tests_lump_sum_and_synonyms(self):
        """BL55 regressions (TEST ONLY rows modelled on the TT38 noise seen in the Mễ Sở AB batch)."""
        import codelib
        n=lambda code,cond,uom,ed='TEST PL2':{'id':'TT38-2026:'+code,'code':code,'kind':'TT38_ATOMIC','edition':ed,'uom':uom,'denominator':'1','conditions':cond,'source_ref':'TEST','parse_status':'PARSED','parse_flags':''}
        self.app.store.library_add('NORMS','TEST ONLY TT38 BL55','6'*64,'n.json','',[
            n('TC.12910','Thí nghiệm ép cọc biến dạng lớn PDA; Thí nghiệm siêu âm chất lượng cọc','mặt cắt',ed='TEST PL1'),n('TS.37510','Bơm vữa xi măng lấp đầy ống siêu âm cọc khoan nhồi; Trên cạn','m3'),
            n('TS.25111','Lắp đặt gối cầu; Loại gối cầu / Gối thép','bộ'),n('TB.91103','Lắp đặt chậu rửa - lắp đặt thuyền tắm; Có hương / sen','bộ',ed='TEST PL3'),n('TS.99110','Công tác khảo sát địa hình','100m')])
        st=self.app.store
        self.assertEqual(codelib.suggest(st,'Siêu âm cọc D2.0m','Cọc khoan nhồi','cọc')[0]['code'],'TC.12910')
        self.assertEqual(codelib.suggest(st,'Gối chậu song hướng 31000kN','Kết cấu phần trên','bộ')[0]['code'],'TS.25111')
        for u in ('Trọn gói','TG','khoản'):self.assertEqual(codelib.suggest(st,'Chi phí khảo sát','',u),[],u)
        self.assertTrue(codelib.lump_sum('trọn gói') and not codelib.lump_sum('m3'))
    def test_suggest_bl58_method_groups_components(self):
        """BL58 regressions on TEST ONLY rows: concrete method preference, grouping by source group heading,
        component titles, bracketed specs, drilling synonym and 'Móng' headings."""
        import codelib
        n=lambda code,cond,uom,grp:{'id':'TT38-2026:'+code,'code':code,'kind':'TT38_ATOMIC','edition':'TEST PL2','uom':uom,'denominator':'1','conditions':cond,'source_ref':'TEST trang 1; nhóm '+grp+'; sha256 x','parse_status':'PARSED','parse_flags':''}
        self.app.store.library_add('NORMS','TEST ONLY TT38 BL58','7'*64,'n.json','',[
            n('AF.14210','BÊ TÔNG MŨ MỐ, MŨ TRỤ CẦU; Bê tông mũ trụ cầu; Trên cạn','m3','AF.14200'),n('AF.33210','BÊ TÔNG MŨ MỐ, MŨ TRỤ CẦU; Bê tông mũ trụ cầu; Trên cạn','m3','AF.33000'),
            n('AF.31110','BÊ TÔNG MÓNG; Bê tông móng; ≤250','m3','AF.31100'),n('AF.32410','BÊ TÔNG TƯỜNG CHẮN; Bê tông tường chắn','m3','AF.32400'),
            n('TD.23221','RẢI THẢM BÊ TÔNG NHỰA CHẶT (LOẠI BTNC19); 3','m2','TD.23220'),n('TD.23271','RẢI THẢM BÊ TÔNG NHỰA CHẶT (LOẠI BTNC16); 3','m2','TD.23270'),n('TD.23272','RẢI THẢM BÊ TÔNG NHỰA CHẶT (LOẠI BTNC16); 4','m2','TD.23270'),
            n('TC.32210','KHOAN VÀO ĐẤT DƯỚI NƯỚC; Khoan vào đất dưới nước; 2000','m','TC.32200'),n('TC.34511','DƯỚI NƯỚC; Lắp đặt ống vách cọc khoan nhồi dưới nước; ≤800','m','TC.34510')])
        st=self.app.store;s=lambda *a,**k:[x['code'] for x in codelib.suggest(st,*a,**k)]
        self.assertEqual(s('Bê tông C40','Kết cấu > Xà mũ','m3',method='BOM')[0],'AF.33210')
        self.assertEqual(s('Bê tông C40','Kết cấu > Xà mũ','m3',method='THU_CONG')[0],'AF.14210')
        b=codelib.suggest(st,'Bê tông nhựa chặt C16 dày 4cm','Mặt đường','m2')
        self.assertTrue(b[0]['code'].startswith('TD.2327'));self.assertEqual({c['code'] for c in b[0]['columns']},{'TD.23271','TD.23272'},'grouped by source group heading, not code[:6]')
        self.assertEqual(s('Thi công cọc khoan nhồi D2.0m (trong phạm vi ống vách để lại)','Thi công dưới nước > Cọc khoan nhồi','m')[0],'TC.32210')
        self.assertEqual(s('Bê tông C30','Tường chắn BTCT > Móng tường','m3',method='BOM')[0],'AF.31110')
    def test_suggest_bl75_units_method_components_survey(self):
        """BL75 regressions on TEST ONLY rows: count unit vs volume table, method written in the description,
        named component missing from the title, survey heading → PL1, fly-ash chapter AN, guardrail synonym."""
        import codelib
        n=lambda code,cond,uom,ed='TEST PL2':{'id':'TT38-2026:'+code,'code':code,'kind':'TT38_ATOMIC','edition':ed,'uom':uom,'denominator':'1','conditions':cond,'source_ref':'TEST; nhóm '+code[:5]+'00; x','parse_status':'PARSED','parse_flags':''}
        self.app.store.library_add('NORMS','TEST ONLY TT38 BL75','a'*64,'n.json','',[
            n('TX.71120','NẠO VÉT BẰNG TÀU HÚT; Nạo vét, chiều cao ống xả ≤6m','m3'),n('TX.24201','LẮP ĐẶT XÀ; Lắp đặt xà; Lắp thủ công','bộ'),
            n('AF.11210','BÊ TÔNG MÓNG; Bê tông móng; ≤250','m3'),n('AF.31110','BÊ TÔNG MÓNG; Bê tông móng; ≤250','m3'),n('TX.25143','CHIỀU RỘNG MÓNG ≤ 6M; Đào móng bằng máy; chiều rộng móng 250','m3'),
            n('TX.61110','CỐT THÉP MÓNG; Cốt thép móng; ≤10','tấn'),n('TX.01101','LẮP ĐẶT HỆ THỐNG ĐƯỜNG ỐNG THÉP BẰNG PHƯƠNG PHÁP HÀN; Lắp đặt hệ thống đường ống thép ITS','tấn',ed='TEST PL4'),
            n('TY.21130','ĐỘ SÂU HỐ KHOAN TỪ 0M ĐẾN 30M; Độ sâu hố khoan; cấp IV','m khoan',ed='TEST PL1'),n('TX.31341','KHOAN ĐÁ TRÊN CẠN; Khoan vào đá cấp IV, trên cạn','m'),
            n('AN.11213','ĐẮP NỀN ĐƯỜNG BẰNG HỖN HỢP TRO XỈ BẰNG MÁY LU; Đắp nền đường K95','m3'),n('TX.64113','ĐẮP ĐẤT NỀN ĐƯỜNG BẰNG MÁY LU; Đắp đất nền đường; K=0,95','m3'),
            n('TX.27110','LẮP ĐẶT HỆ THỐNG AN TOÀN HỘ LAN BÁNH XOAY; Hộ lan bánh xoay','m'),n('TX.34110','LẮP ĐẶT DẢI PHÂN CÁCH; Lắp đặt dải phân cách bằng tôn lượn sóng','m')])
        st=self.app.store;s=lambda *a,**k:[x['code'] for x in codelib.suggest(st,*a,**k)]
        self.assertEqual(s('Xà bắt đèn pha','Nút giao','bộ')[0],'TX.24201','count unit must not match a m3 dredging table')
        self.assertEqual(s('Đổ bê tông thủ công bằng máy trộn, bê tông móng, chiều rộng > 250 cm','Nút giao','m3')[0],'AF.11210')
        self.assertEqual(s('Cốt thép','Hệ thống ITS','tấn')[0],'TX.61110')
        self.assertEqual(s('Khoan trên cạn độ sâu 0-30m, cấp IV','Chi phí tư vấn > Khảo sát địa chất cống','m khoan')[0],'TY.21130')
        self.assertEqual(s('Đắp nền bằng đất K95','Đường','m3')[0],'TX.64113')
        self.assertEqual(s('Hộ lan tôn sóng mạ kẽm 3 sóng','An toàn giao thông','m')[0],'TX.34110')
    def test_suggest_bl60_dictionary_rules(self):
        """BL60: versioned BOQ→TT38 dictionary (REFERENCE) steers candidates and flags lines; bad patterns are skipped."""
        import codelib
        n=lambda code,cond,uom:{'id':'TT38-2026:'+code,'code':code,'kind':'TT38_ATOMIC','edition':'TEST PL3','uom':uom,'denominator':'1','conditions':cond,'source_ref':'TEST; nhóm '+code[:5]+'00; x','parse_status':'PARSED','parse_flags':''}
        st=self.app.store
        st.library_add('NORMS','TEST ONLY TT38 BL60','8'*64,'n.json','',[n('TA.14401','LẮP ĐẶT ỐNG NHỰA ĐẶT CHÌM BẢO HỘ DÂY DẪN; Lắp đặt ống nhựa đặt chìm bảo hộ dây dẫn','m'),n('TB.45111','LẮP ĐẶT ỐNG NHỰA HDPE NỐI BẰNG MĂNG SÔNG; Lắp đặt ống nhựa HDPE','m')])
        before=[c['code'] for c in codelib.suggest(st,'Ống nhựa xoắn HDPE D85','Ụ chân cột điện','m')]
        self.assertEqual(before[0],'TB.45111')
        rule=lambda i,m,**k:dict({'id':i,'series':'BOQ_TT38_RULE','match':m,'ctx':'','prefer':[],'rewrite':'','flag':'','note':'TEST','status':'TEST ONLY'},**k)
        st.library_add('REFERENCE','Từ điển BOQ→TT38 TEST ONLY','9'*64,'d.json','',[rule('R1','ong (nhua )?(xoan|ruot ga)|ong (nhua )?hdpe xoan',prefer=['TA.144'],rewrite='ong nhua dat chim bao ho day dan'),rule('R2','phong nuoc mat cau',flag='NO_TT38'),rule('BAD','(unclosed')])
        self.assertEqual(codelib.suggest(st,'Ống nhựa xoắn HDPE D85','Ụ chân cột điện','m')[0]['code'],'TA.14401')
        self.assertEqual([r['id'] for r in codelib.matched_rules(st,'Lớp phòng nước mặt cầu','')],['R2'])
        self.assertEqual(codelib.matched_rules(st,'Khoan cọc (trong phạm vi ống nhựa xoắn)',''),[],'bracketed spec must not trigger a rule')
        f=fixture();f['boq'][0].update(description='Lớp phòng nước mặt cầu',uom='m2');f['mapping'][0]['uom']='m2';s=self.call('/api/projects',{'name':'TEST ONLY BL60'})['project']
        for i,key in enumerate(['context','boq']):s=self.call(f'/api/projects/{s["id"]}/stage/{i}',{'revision':s['revision'],'data':f[key],'confirm':True})['project']
        it=self.call(f'/api/projects/{s["id"]}/suggest-norms',{})['items'][0];self.assertEqual([(r['id'],r['flag']) for r in it['rules']],[('R2','NO_TT38')])
    def test_qs_memory_remembers_gate_passed_choices_and_applies(self):
        """BL63 (D37 'tự áp dụng'): gate-confirmed choices are remembered append-only and applied to identical lines;
        drafts are not remembered; disagreeing decisions are CONFLICT; item facts are never copied."""
        import codelib
        self.code_library_fixture();st=self.app.store
        st.library_add('NORMS','TEST ONLY TT38 thứ hai','a'*64,'n2.json','',[{'id':'TT38-2026:TS.00022','code':'TS.00022','kind':'TT38_ATOMIC','edition':'TEST','uom':'m3','denominator':'1','conditions':'TEST ONLY bê tông khác','source_ref':'TEST','selection_ref':'','table_verified':False,'parse_status':'PARSED','parse_flags':''}])
        def project(name,nid,confirm=True):
            f=fixture();s=self.call('/api/projects',{'name':name})['project']
            for i,key in enumerate(['context','boq','mapping']):s=self.call(f'/api/projects/{s["id"]}/stage/{i}',{'revision':s['revision'],'data':f[key],'confirm':True,'wbs':f['wbs'] if i==2 else None})['project']
            x=self.call(f'/api/projects/{s["id"]}/code-norm',{'id':nid});norm=dict(x['norm'],selection_ref='TEST QS chọn '+name,table_verified=True,kind='CC1_APPROVED_CUSTOM',tt37_ref='TEST')
            d=self.call(f'/api/projects/{s["id"]}/stage/3',{'revision':s['revision'],'data':[norm],'confirm':confirm,'applications':[{'id':'a-'+name,'item_id':f['mapping'][0]['id'],'norm_id':nid,'qty_factor':'1'}],'assignments':[{k:m.get(k) for k in ['id','technical_spec','method','norm_match_ref']} for m in f['mapping']],'resources':x['resources']})
            self.assertEqual(d['confirmed'],confirm,d.get('questions'));return d['project'],f
        key=codelib.memory_key('TEST ONLY concrete','m3')
        project('TEST ONLY nháp','TT38-2026:TS.00011',confirm=False);self.assertEqual(st.choices_for([key]),[],'drafts are not remembered')
        a,_=project('TEST ONLY A','TT38-2026:TS.00011');self.assertEqual([h['norm_code'] for h in st.choices_for([key])],['TS.00011'])
        f=fixture();b=self.call('/api/projects',{'name':'TEST ONLY B'})['project']
        for i,k in enumerate(['context','boq','mapping']):
            data=f[k]
            if k=='mapping':data=[{x:y for x,y in data[0].items() if x not in ('technical_spec','method','norm_match_ref')}]
            b=self.call(f'/api/projects/{b["id"]}/stage/{i}',{'revision':b['revision'],'data':data,'confirm':True,'wbs':f['wbs'] if i==2 else None})['project']
        sg=self.call(f'/api/projects/{b["id"]}/suggest-norms',{})['items'][0];self.assertEqual(sg['memory']['status'],'UNIQUE')
        m=self.call(f'/api/projects/{b["id"]}/apply-memory',{})
        self.assertEqual([x['norm_id'] for x in m['applications']],['TT38-2026:TS.00011']);self.assertEqual(len(m['resources']),3)
        n=m['norms'][0];self.assertTrue(n['selection_ref'].startswith('Theo lựa chọn QS đã nhớ: hồ sơ «TEST ONLY A»'));self.assertIs(n['table_verified'],True)
        d=self.call(f'/api/projects/{b["id"]}/stage/3',{'revision':b['revision'],'data':[dict(n,kind='CC1_APPROVED_CUSTOM',tt37_ref='TEST')],'confirm':True,'applications':m['applications'],'resources':m['resources'],'assignments':[{'id':b['mapping'][0]['id']}]})
        self.assertFalse(d['confirmed']);self.assertTrue(any(q['field'].endswith('technical_spec') for q in d['questions']),'item facts are not copied from memory')
        project('TEST ONLY C','TT38-2026:TS.00022')
        mem=self.call(f'/api/projects/{b["id"]}/suggest-norms',{})['items'][0]['memory'];self.assertEqual((mem['status'],mem['decisions']),('CONFLICT',2),'two different decisions -> conflict')
        again=self.call(f'/api/projects/{b["id"]}/apply-memory',{});self.assertEqual((again['applications'],len(again['skipped_has_recipe'])),([],1),'an item that already has a recipe is never overwritten')
        self.assertEqual(len(st.choices_for([key])),2,'memory is append-only')
        dd=self.call('/api/projects',{'name':'TEST ONLY D'})['project']
        for i,k in enumerate(['context','boq','mapping']):dd=self.call(f'/api/projects/{dd["id"]}/stage/{i}',{'revision':dd['revision'],'data':f[k],'confirm':True,'wbs':f['wbs'] if i==2 else None})['project']
        dm=self.call(f'/api/projects/{dd["id"]}/apply-memory',{});self.assertEqual((dm['applications'],len(dm['conflicts'])),([],1),'conflict -> nothing applied')
    def test_qs_memory_ignores_unconfirmed_and_revoked_choices(self):
        """BL70 (D45): a basis that says 'Cần QS xác nhận' is not remembered; revoked choices stay as evidence but are
        never offered or applied again."""
        import codelib
        self.code_library_fixture();st=self.app.store;key=codelib.memory_key('TEST ONLY concrete','m3')
        def confirm(name,sel):
            f=fixture();s=self.call('/api/projects',{'name':name})['project']
            for i,k in enumerate(['context','boq','mapping']):s=self.call(f'/api/projects/{s["id"]}/stage/{i}',{'revision':s['revision'],'data':f[k],'confirm':True,'wbs':f['wbs'] if i==2 else None})['project']
            x=self.call(f'/api/projects/{s["id"]}/code-norm',{'id':'TT38-2026:TS.00011'});norm=dict(x['norm'],selection_ref=sel,table_verified=True,kind='CC1_APPROVED_CUSTOM',tt37_ref='TEST')
            d=self.call(f'/api/projects/{s["id"]}/stage/3',{'revision':s['revision'],'data':[norm],'confirm':True,'applications':[{'id':'a-'+name,'item_id':f['mapping'][0]['id'],'norm_id':norm['id'],'qty_factor':'1'}],'assignments':[{k:m.get(k) for k in ['id','technical_spec','method','norm_match_ref']} for m in f['mapping']],'resources':x['resources']})
            self.assertTrue(d['confirmed']);return d['project']
        confirm('TEST ONLY auto','PILOT auto: ứng viên #1 từ Gợi ý mã TT38. Cần QS xác nhận.')
        self.assertEqual(st.choices_for([key]),[],'self-declared unconfirmed basis is not a QS choice')
        confirm('TEST ONLY QS thật','TEST QS đã kiểm bảng gốc trang 1')
        hits=st.choices_for([key]);self.assertEqual(len(hits),1)
        self.assertEqual(st.choices_revoke([hits[0]['id']],'TEST thu hồi','TEST','TEST D45'),1);self.assertEqual(st.choices_revoke([hits[0]['id']],'again','TEST','TEST'),0,'revocation is idempotent')
        self.assertEqual(st.choices_for([key]),[],'revoked choice is never offered again')
        with st.connect() as db:self.assertEqual(db.execute('select count(*) from qs_choices').fetchone()[0],1,'the revoked row itself is kept as evidence')
        # D47: a QS basis on the application wins over an unconfirmed norm-level basis shared with PILOT items
        p3=confirm('TEST ONLY app-level','PILOT auto: ứng viên #1. Cần QS xác nhận.')
        self.assertEqual(st.choices_for([key]),[])
        s3=self.call(f'/api/projects/{p3["id"]}')['project'];s3['applications'][0].update(selection_ref='TEST QS chọn theo gói pilot',table_verified=True)
        d=self.call(f'/api/projects/{p3["id"]}/stage/3',{'revision':s3['revision'],'data':s3['norms'],'confirm':True,'applications':s3['applications'],'resources':s3['resources'],'assignments':[{k:m.get(k) for k in ['id','technical_spec','method','norm_match_ref']} for m in s3['mapping']]})
        self.assertTrue(d['confirmed']);self.assertEqual([h['selection_ref'] for h in st.choices_for([key])],['TEST QS chọn theo gói pilot'])
        f=fixture();n=self.call('/api/projects',{'name':'TEST ONLY mới'})['project']
        for i,k in enumerate(['context','boq','mapping']):n=self.call(f'/api/projects/{n["id"]}/stage/{i}',{'revision':n['revision'],'data':f[k],'confirm':True,'wbs':f['wbs'] if i==2 else None})['project']
        self.assertEqual(len(self.call(f'/api/projects/{n["id"]}/apply-memory',{})['applications']),1,'only the confirmed QS choice is applied')
    def test_ai_settings_never_expose_key_and_consent(self):
        p=self.call('/api/ai/settings',{'provider':'openai','api_key':'sk-TESTONLY-1234567890','model':'test-model'})
        self.assertNotIn('sk-TESTONLY',json.dumps(p));self.assertTrue(next(x for x in p['providers'] if x['id']=='openai')['key_set'])
        s=self.save_all(fixture());calls=[]
        original=ai.CALLERS['openai'];ai.CALLERS['openai']=lambda cfg,model,system,user:(calls.append((cfg['api_key'],model,user)) or ('TEST ONLY proposal',model))
        try:
            with self.assertRaises(urllib.error.HTTPError) as e:self.call('/api/ai/ask',{'project_id':s['id'],'step':5,'provider':'openai','model':'test-model','question':'Kiểm giá?','consent':False})
            self.assertEqual(e.exception.code,502)
            d=self.call('/api/ai/ask',{'project_id':s['id'],'step':5,'provider':'openai','model':'test-model','question':'Kiểm giá?','consent':True,'include_rows':True})
        finally:ai.CALLERS['openai']=original
        self.assertEqual(d['proposal']['status'],'AI_PROPOSAL_NOT_APPLIED');self.assertEqual(calls[0][0],'sk-TESTONLY-1234567890');self.assertIn('TEST-MAT',calls[0][2])
        after=self.call(f'/api/projects/{s["id"]}');self.assertEqual(after['project']['revision'],s['revision']);self.assertTrue(all(x['confirmed'] for x in after['stages']))
        self.assertEqual(len(self.call(f'/api/projects/{s["id"]}/ai')['proposals']),1)
        with urllib.request.urlopen(self.base+f'/api/projects/{s["id"]}/backup') as r:self.assertNotIn(b'sk-TESTONLY',r.read())
if __name__=='__main__':unittest.main(verbosity=2)
