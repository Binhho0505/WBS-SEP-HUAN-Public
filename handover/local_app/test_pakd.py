"""E17 (D76) native tests for PAKD / BM01-TH. TEST ONLY fixtures; every expected value is computed by hand in the comments.

Fixture (test_subcontract.multi + prices): AB 1.1 = 200 m3 × 100; 1.2 = 10 t × 50; 2.1 = 100 m3 × 20 → B = 22,500.
Estimate VL: 1.1 = 200 × 1.5/100 × 10.01 = 30.03; 2.1 = 100 × 1.5/100 × 10.01 = 15.015; 1.2 has no recipe (incomplete).
"""
import base64,copy,io,json,tempfile,threading,unittest,urllib.request,urllib.error
from decimal import Decimal
from pathlib import Path
from openpyxl import load_workbook
import subcontract as sc,pakd
from server import App
from test_subcontract import multi,pk,GOP
def base():
    s=multi();ab3=next(r for r in s['boq'] if r['id']=='ab3');ab3['unit_price']='20'
    s['applications'].append(dict(id='TEST-APP-3',item_id='item3',norm_id='TEST-RECIPE',qty_factor='1'))
    gt1=pk('GT-01',procurement_method='CHI_DINH_THAU',delivery=dict(GOP),pricing={'k_all':'0.9'},
           lines=[{'id':'l1','bb_code':'GT-01.1','ab_line_id':'test-source','alloc_pct':'100'},{'id':'l2','bb_code':'GT-01.2','ab_line_id':'ab2','alloc_pct':'100'}])
    cc1=pk('CC1','CC1 tự làm',lines=[{'id':'l3','bb_code':'CC1.1','ab_line_id':'ab3','alloc_pct':'40'}])
    s=sc.save(s,{'packages':[gt1,cc1],'pm_model':'PMBOK6'})
    s=sc.save_award(s,{'package_id':'id-GT-01','vendor':'TEST NTP','reason':'TEST','proposed_by':'TEST','prices':{'GT-01.1':'95','GT-01.2':'40'}})
    s['revision']=1;return s
LINES=[{'code':'C.1.1','method':'AMOUNT','value':'100','source_ref':'TEST báo giá thiết kế'},
       {'code':'C.2','method':'PERCENT','base':'DT','value':'2','source_ref':'TEST quy chế gián tiếp'},
       {'code':'C.4','method':'PERCENT','base':'C1','value':'1','source_ref':'TEST bảo hành'}]
def row(res,code):return next(r for r in res['rows'] if r['code']==code)
class PakdDomain(unittest.TestCase):
    def setUp(self):
        self.s=base();sid=pakd.state(self.s)['current']
        self.s=pakd.act(self.s,{'action':'save','scenario_id':sid,'basis':{'id-GT-01':'AWARD'},'lines':LINES});self.sid=sid
    def res(self,s=None):st=pakd.state(s or self.s);return pakd.compute(s or self.s,pakd._scen(st,self.sid))
    def test_bm01_hand_checked(self):
        r=self.res();self.assertTrue(r['complete'],r['missing'])
        # C.1.5 = 95×200 + 40×10 = 19,400; CPXD = 6.006 + 19,400 + 9.009 = 19,415.015; C.1 = 100 + CPXD = 19,515.015
        # C.2 = 2% × 22,500 = 450; C.4 = 1% × 19,515.015 = 195.15015; C = 20,160.16515; D = 2,339.83485 (10.399…%)
        want={'B':'22500','C.1.1':'100','C.1.4':'6.01','C.1.5':'19400','C.1.6':'0','C.1.6b':'0','C.1.7':'9.01','C.2':'450','C.4':'195.15','D':'2339.83'}
        self.assertEqual({k:row(r,k)['value'] for k in want},want)
        self.assertEqual((row(r,'D')['margin_pct'],r['totals']['CPXD'],r['totals']['C1']),('10.40','19415.02','19515.02'))
        g={p['code']:p for p in r['packages']}
        self.assertEqual((g['GT-01']['ab_value'],g['GT-01']['cost'],g['GT-01']['direct_margin'],g['GT-01']['direct_margin_pct']),('20500','19400','1100','5.4'))
        self.assertEqual((g['CC1']['ab_value'],g['CC1']['cost']),('800','6.01'))
    def test_missing_is_never_zero(self):
        s=copy.deepcopy(self.s);next(r for r in s['boq'] if r['id']=='ab3')['unit_price']=None
        r=self.res(s);self.assertIsNone(row(r,'B')['value']);self.assertIsNone(row(r,'D')['value']);self.assertEqual(row(r,'B')['known'],'20500')
        self.assertIsNone(row(r,'C.2')['value'])  # 2% of an unknown revenue stays unknown
        self.assertTrue(any('đơn giá hợp đồng' in m['text'] for m in r['missing']))
        s=copy.deepcopy(self.s);s['packages'][0]['award']['prices'].pop('GT-01.2')
        r=self.res(s);self.assertIsNone(row(r,'C.1.5')['value']);self.assertEqual(row(r,'C.1.5')['known'],'19000');self.assertIsNone(row(r,'D')['value'])
    def test_basis_choice_and_no_ab_ratio_split(self):
        s=pakd.act(copy.deepcopy(self.s),{'action':'save','scenario_id':self.sid,'basis':{'id-GT-01':'A'}})
        r=self.res(s);self.assertEqual(row(r,'C.1.5')['value'],'18450')  # A = AB price × K 0.9: 90×200 + 45×10
        s=pakd.act(s,{'action':'save','scenario_id':self.sid,'basis':{'id-GT-01':'B'}})
        r=self.res(s);self.assertIsNone(row(r,'C.1.5')['value']);self.assertTrue(any('căn cứ B' in m['text'] for m in r['missing']))
        s=pakd.act(s,{'action':'save','scenario_id':self.sid,'basis':{}});r=self.res(s)
        self.assertTrue(any('chưa chọn căn cứ giá giao' in m['text'] for m in r['missing']));self.assertTrue(any('GT-01' in q['question'] for q in pakd.questions(s,pakd._scen(pakd.state(s),self.sid))))
    def test_tach_cc1_supplies_material(self):
        s=copy.deepcopy(self.s);s['packages'][0]['delivery']=dict(GOP,mode='TACH',VL='CC1',NC='NTP',M='NTP')
        r=self.res(s);self.assertEqual((row(r,'C.1.6')['known'],row(r,'C.1.5')['value']),('19400','0'))
        self.assertTrue(any('GT-01: dòng GT-01.2: dự toán CC1' in m['text'] for m in r['missing']))  # 1.2 has no recipe: CC1 material cost unknown
        self.assertEqual(row(r,'C.1.4')['known'],'36.04')  # 30.03 (1.1 VL by CC1) + 6.006 (CC1 self)
    def test_user_lines_rules(self):
        st=pakd.state(self.s)
        with self.assertRaises(ValueError):pakd.check_lines([{'code':'C.1.1','method':'PERCENT','base':'DT','value':'1'}])  # C.1.1 amount only
        with self.assertRaises(ValueError):pakd.check_lines([{'code':'C.1.8','method':'PERCENT','base':'C1','value':'1'}])  # cycle base refused
        with self.assertRaises(ValueError):pakd.check_lines([{'code':'C.9','method':'AMOUNT','value':'1'}])
        s=pakd.act(copy.deepcopy(self.s),{'action':'save','scenario_id':self.sid,'lines':LINES+[{'code':'C.3','method':'PERCENT','base':'DT','value':'','source_ref':''}]})
        q=pakd.questions(s,pakd._scen(pakd.state(s),self.sid));self.assertEqual(len([x for x in q if x['field'].startswith('pakd.lines[3]')]),2)
        self.assertIsNone(row(self.res(s),'C.3')['value'])  # an empty % is not 0
        s=pakd.act(s,{'action':'save','scenario_id':self.sid,'lines':LINES+[{'code':'C.3','method':'PERCENT','base':'DT','value':'5','source_ref':'x','on':False}]})
        self.assertEqual(row(self.res(s),'C.3')['value'],'0')  # switched off → not counted, not missing
    def test_scenarios_and_frozen_versions(self):
        s=pakd.act(copy.deepcopy(self.s),{'action':'copy','scenario_id':self.sid,'name':'TEST căn cứ A'});st=pakd.state(s);self.assertEqual(len(st['scenarios']),2)
        new=st['current'];s=pakd.act(s,{'action':'save','scenario_id':new,'basis':{'id-GT-01':'A'}})
        v=pakd.view(s);d={c['name']:c['result']['totals']['D'] for c in v['scenarios']};self.assertEqual(d,{'Kịch bản 1':'2339.83','TEST căn cứ A':'3299.33'})  # A: C.1.5 18,450 → C 19,200.66515 → D 3,299.33485
        with self.assertRaises(ValueError) as e:pakd.act(s,{'action':'freeze','scenario_id':self.sid,'name':'TEST lần 1','by':'TEST QS'})  # D86: baseline needs the cash flow
        self.assertIn('dòng tiền',str(e.exception));s=pakd.act(s,{'action':'save','scenario_id':self.sid,'cashflow':self.CF})
        s=pakd.act(s,{'action':'freeze','scenario_id':self.sid,'name':'TEST lần 1','by':'TEST QS'});v0=pakd.state(s)['versions'][0]
        self.assertEqual((v0['status'],len(v0['inputs_sha256']),row(v0,'D')['value']),('FROZEN_LOCAL',64,'2339.83'))
        with self.assertRaises(ValueError):pakd.act(s,{'action':'freeze','scenario_id':self.sid,'name':'test LẦN 1','by':'x'})  # never overwritten
        s2=copy.deepcopy(s);s2['packages'][0]['award']['prices']['GT-01.1']='99'
        vv=pakd.view(s2)['versions'][0];self.assertTrue(vv['stale']);self.assertEqual(row(vv,'D')['value'],'2339.83')  # frozen rows unchanged
        self.assertFalse(pakd.view(s)['versions'][0]['stale'])
        bad=pakd.act(copy.deepcopy(s),{'action':'save','scenario_id':self.sid,'basis':{}})
        with self.assertRaises(ValueError):pakd.act(bad,{'action':'freeze','scenario_id':self.sid,'name':'TEST thiếu','by':'x'})  # incomplete cannot be frozen
        one=pakd.act(copy.deepcopy(s),{'action':'delete','scenario_id':new});self.assertEqual(len(pakd.state(one)['scenarios']),1)
        with self.assertRaises(ValueError):pakd.act(one,{'action':'delete','scenario_id':self.sid})  # the last scenario stays
    # ---------- BL100 (D77): parts by work group (prompt 6.11) and cash flow (PRD FR22) ----------
    def test_work_groups_from_boq_headings(self):
        g=pakd.work_groups(self.s);self.assertEqual(g['test-source'],('1','TEST ONLY heading'));self.assertEqual(g['ab3'][0],pakd.NOGROUP)
        pg=pakd.package_groups(self.s,self.s['packages'][0]);self.assertEqual([(x['code'],x['bb_codes'],x['value']) for x in pg],[('1',['GT-01.1','GT-01.2'],'20500')])
    def test_parts_by_group_and_line(self):
        part={'id':'p-thep','name':'TEST thép','mode':'GOP','basis':'A'}
        # whole group "1" → part on basis A: 90×200 + 45×10 = 18,450 (same as scenario A) → D 3,299.33; the package basis is no longer needed
        s=pakd.act(copy.deepcopy(self.s),{'action':'save','scenario_id':self.sid,'basis':{},'parts':{'id-GT-01':{'parts':[part],'by_group':{'1':'p-thep'}}}})
        sc_=pakd._scen(pakd.state(s),self.sid);r=self.res(s)
        self.assertTrue(r['complete'],r['missing']);self.assertEqual(row(r,'C.1.5')['value'],'18450');self.assertEqual(row(r,'D')['value'],'3299.33')
        self.assertFalse([q for q in pakd.questions(s,sc_) if q['field']=='pakd.basis.GT-01'])
        # a line choice beats the group: GT-01.1 back to the common part (AWARD 95×200 = 19,000) + part 45×10 = 450 → C.1.5 19,450
        # C.4 = 1% × (19,515.015 + 50) → D = 2,339.83485 − 50 − 0.5 = 2,289.33485
        s=pakd.act(s,{'action':'save','scenario_id':self.sid,'basis':{'id-GT-01':'AWARD'},'parts':{'id-GT-01':{'parts':[part],'by_group':{'1':'p-thep'},'by_line':{'GT-01.1':'main'}}}})
        r=self.res(s);self.assertEqual((row(r,'C.1.5')['value'],row(r,'D')['value']),('19450','2289.33'))
        g=next(p for p in r['packages'] if p['code']=='GT-01');self.assertEqual([(x['name'],x['lines'],x['cost']) for x in g['parts']],[('Phần chung',1,'19000'),('TEST thép',1,'450')])
        # a part without arrangement is missing (never 0) and asked for by name
        s2=pakd.act(copy.deepcopy(s),{'action':'save','scenario_id':self.sid,'parts':{'id-GT-01':{'parts':[dict(part,mode=None,basis=None)],'by_line':{'GT-01.2':'p-thep'}}}})
        r=self.res(s2);self.assertIsNone(row(r,'C.1.5')['value']);self.assertEqual(row(r,'C.1.5')['known'],'19000')
        self.assertTrue(any('phần "TEST thép"' in m['text'] for m in r['missing']))
        self.assertEqual({q['field'] for q in pakd.questions(s2,pakd._scen(pakd.state(s2),self.sid)) if 'parts' in q['field']},{'pakd.parts.GT-01[0].mode','pakd.parts.GT-01[0].basis'})
        # an empty part deletes itself; unknown targets are refused; the CC1 package cannot be split
        s3=pakd.act(copy.deepcopy(s),{'action':'save','scenario_id':self.sid,'parts':{'id-GT-01':{'parts':[part]}}})
        self.assertEqual(pakd._scen(pakd.state(s3),self.sid)['parts'],{})
        for bad in ({'id-GT-01':{'parts':[part],'by_line':{'GT-01.1':'nope'}}},{'id-GT-01':{'parts':[part],'by_group':{'9':'p-thep'}}},{'id-CC1':{'parts':[part]}},
                    {'id-GT-01':{'parts':[dict(part,name='')]}},{'id-GT-01':{'parts':[dict(part,mode='X')]}}):
            with self.assertRaises(ValueError):pakd.act(copy.deepcopy(s),{'action':'save','scenario_id':self.sid,'parts':bad})
        # parts are inputs: changing them makes a frozen version stale; copying a scenario copies them
        s=pakd.act(s,{'action':'save','scenario_id':self.sid,'cashflow':self.CF});s=pakd.act(s,{'action':'freeze','scenario_id':self.sid,'name':'TEST phần','by':'TEST'})
        s4=pakd.act(copy.deepcopy(s),{'action':'save','scenario_id':self.sid,'parts':{}});self.assertTrue(pakd.view(s4)['versions'][0]['stale'])
        s5=pakd.act(copy.deepcopy(s),{'action':'copy','scenario_id':self.sid,'name':'TEST bản sao'});self.assertEqual(pakd.state(s5)['scenarios'][1]['parts'],pakd._scen(pakd.state(s),self.sid)['parts'])
    CF={'periods':3,'start':'2027-01','schedule_ref':'TEST tiến độ','revenue_pct':['20','50','30'],'ntp_pct':['30','50','20'],'own_pct':['100','0','0'],
        'ab':{'advance_pct':'10','advance_period':'1','recovery_pct':'50','retention_pct':'5','retention_release':'4','delay':'1','ref':'TEST HĐ AB'},
        'ntp':{'advance_pct':'0','retention_pct':'0','delay':'0','ref':'TEST HĐ NTP'},'own':{'delay':'0','ref':'TEST'},'rate_pct':'12','rate_ref':'TEST lãi vay'}
    def test_cashflow_hand_checked(self):
        s=pakd.act(copy.deepcopy(self.s),{'action':'save','scenario_id':self.sid,'cashflow':self.CF})
        v=pakd.view(s)['scenarios'][0];cf=v['cashflow'];self.assertTrue(cf['complete'],cf['missing']);self.assertEqual(v['questions'],[])
        # money_places 2: B 22,500, C 20,160.17, NTP = C.1.5 19,400, CC1 own = 760.17.
        # AB: advance 2,250 in T1 (not revenue); earned 4,500 / 11,250 / 6,750; T1 recovers 2,250 (50% capped at the advance), retention 5%;
        #     paid one period later: T2 2,025, T3 10,687.5, T4 6,412.5 + retention 1,125 = 7,537.5 → Σ = 22,500 = B
        # NTP: 5,820 / 9,700 / 3,880; CC1: 760.17 in T1. Cumulative −4,330.17 / −12,005.17 / −5,197.67 / 2,339.83 = B − C
        p=cf['periods'];self.assertEqual([x['label'] for x in p],['2027-01','2027-02','2027-03','2027-04'])
        self.assertEqual([x['ab_in'] for x in p],['2250','2025','10687.50','7537.50']);self.assertEqual([x['earned'] for x in p],['4500','11250','6750','0'])
        self.assertEqual([x['ntp_out'] for x in p],['5820','9700','3880','0']);self.assertEqual([x['cum'] for x in p],['-4330.17','-12005.17','-5197.67','2339.83'])
        t=cf['totals'];self.assertEqual((t['in'],t['out'],t['net'],t['peak_funding'],t['peak_period']),('22500','20160.17','2339.83','12005.17','2027-02'))
        self.assertEqual(t['financing'],'215.33')  # (4,330.17 + 12,005.17 + 5,197.67) × 12%/12 = 215.3301
        self.assertEqual(row(v['result'],'D')['value'],'2339.83')  # cash flow never changes the profit table
    def test_cashflow_missing_terms_stay_open(self):
        cf=dict(self.CF,ab={'advance_pct':'10','retention_pct':'5','ref':''},revenue_pct=['20','50',''],rate_pct='',rate_ref='')
        s=pakd.act(copy.deepcopy(self.s),{'action':'save','scenario_id':self.sid,'cashflow':cf})
        v=pakd.view(s)['scenarios'][0];f={q['field'] for q in v['questions']}
        self.assertTrue({'pakd.cashflow.revenue_pct','pakd.cashflow.ab.delay','pakd.cashflow.ab.advance_period','pakd.cashflow.ab.recovery_pct','pakd.cashflow.ab.retention_release','pakd.cashflow.ab.ref'}<=f,f)
        self.assertFalse(v['cashflow']['complete']);self.assertEqual(v['cashflow']['periods'],[])
        s=pakd.act(s,{'action':'save','scenario_id':self.sid,'cashflow':dict(self.CF,revenue_pct=['20','50','20'])})
        self.assertTrue(any('tổng % phải bằng 100' in q['question'] for q in pakd.questions(s,pakd._scen(pakd.state(s),self.sid))))
        s=pakd.act(s,{'action':'save','scenario_id':self.sid,'cashflow':dict(self.CF,rate_pct='')})
        c=pakd.view(s)['scenarios'][0]['cashflow'];self.assertTrue(c['complete']);self.assertIsNone(c['totals']['financing'])  # no rate → no financing figure, not 0
        for bad in (dict(self.CF,periods=0),dict(self.CF,start='2027/01'),dict(self.CF,ab=dict(self.CF['ab'],delay='-1'))):
            with self.assertRaises(ValueError):pakd.act(copy.deepcopy(s),{'action':'save','scenario_id':self.sid,'cashflow':bad})
    def test_process_a_i_mapping(self):  # D86: BM01 → A–I of the PAKD process v1.0; H only when every line is grouped
        r=self.res();pv=r['process'];g={x['group']:x for x in pv['rows']}
        self.assertEqual((g['A']['value'],g['B']['value'],g['G']['value'],g['C']['value'],g['E']['value'],g['F']['value']),('22500','19415.02','19400','450','0','0'))
        self.assertEqual([u['code'] for u in pv['untagged']],['C.1.1','C.4']);self.assertIsNone(g['H']['value']);self.assertEqual(pv['untagged_amount'],'295.15')
        tagged=[dict(LINES[0],process_group='D'),LINES[1],dict(LINES[2],process_group='E')]
        s=pakd.act(copy.deepcopy(self.s),{'action':'save','scenario_id':self.sid,'lines':tagged});r=self.res(s);g={x['group']:x for x in r['process']['rows']}
        self.assertEqual((g['D']['value'],g['E']['value'],g['H']['value'],g['H']['margin_pct'],r['process']['untagged']),('100','195.15','2339.83','10.40',[]))
        self.assertEqual(g['H']['value'],row(r,'D')['value']);self.assertIsNone(g['I']['value'])
        with self.assertRaises(ValueError):pakd.check_lines([dict(LINES[0],process_group='A')])   # A/G/H/I are not cost groups
class PakdHttp(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.app=App(('127.0.0.1',0),Path(self.tmp.name)/'db.sqlite');threading.Thread(target=self.app.serve_forever,daemon=True).start()
        self.base=f'http://127.0.0.1:{self.app.server_port}';s=base();p=self.app.store.create('TEST PAKD');s.update(id=p['id'],revision=p['revision'],name=p['name'])
        self.s=self.app.store.put(s,p['revision'],'TEST_SEED')
    def tearDown(self):self.app.shutdown();self.app.server_close();self.tmp.cleanup()
    def call(self,path,body=None,raw=False):
        req=urllib.request.Request(self.base+path,data=None if body is None else json.dumps(body).encode(),headers={} if body is None else {'Content-Type':'application/json','X-CC1-Session':self.app.token})
        with urllib.request.urlopen(req) as r:return r.read() if raw else json.load(r)
    def test_api_save_freeze_export(self):
        pid=self.s['id'];v=self.call(f'/api/projects/{pid}/subcontract');sid=v['pakd']['current']
        self.assertTrue(any('GT-01' in q['question'] for q in v['pakd']['scenarios'][0]['questions']))
        v=self.call(f'/api/projects/{pid}/pakd',{'revision':v['revision'],'action':'save','scenario_id':sid,'basis':{'id-GT-01':'AWARD'},'lines':LINES,'cashflow':PakdDomain.CF})
        c=v['pakd']['scenarios'][0];self.assertEqual((c['result']['totals']['D'],c['questions']),('2339.83',[]))
        with self.assertRaises(urllib.error.HTTPError) as e:self.call(f'/api/projects/{pid}/pakd',{'revision':0,'action':'select','scenario_id':sid})
        self.assertEqual(e.exception.code,409)
        v=self.call(f'/api/projects/{pid}/pakd',{'revision':v['revision'],'action':'freeze','scenario_id':sid,'name':'TEST chốt','by':'TEST'})
        ver=v['pakd']['versions'][0]['id'];self.assertFalse(v['pakd']['versions'][0]['stale'])
        for q in (f'?scenario={sid}',f'?version={ver}'):
            wb=load_workbook(io.BytesIO(self.call(f'/api/projects/{pid}/pakd.xlsx'+q,raw=True)),read_only=True)
            self.assertEqual(wb.sheetnames,['BM01_TH','GOI_THAU','DONG_NHAP','DONG_TIEN','QUY_TRINH_A_I','CON_THIEU'])
            rows=list(wb['BM01_TH'].iter_rows(values_only=True));self.assertIn('DỰ THẢO TRÊN MÁY',rows[0][0])
            d=next(r for r in rows if r and r[0]=='D');self.assertAlmostEqual(d[3],2339.83,2)
        ev=[e['action'] for e in self.app.store.events(pid)];self.assertIn('PAKD_FREEZE',ev)
        import zipfile  # BL100: the procurement ZIP carries the live scenario and every frozen version
        z=zipfile.ZipFile(io.BytesIO(self.call(f'/api/projects/{pid}/procurement-pack.zip',raw=True)))
        names=[n for n in z.namelist() if n.startswith('PAKD/')];self.assertEqual(len(names),2,names)
        self.assertTrue(any('DANG_LAP' in n for n in names) and any('CHOT_TEST_chốt' in n for n in names),names)
        man=json.loads(z.read('MANIFEST.json'));self.assertTrue(all(any(m['file']==n for m in man['files']) for n in names))
    def test_restore_keeps_pakd_and_excel_never_runs_formulas(self):
        pid=self.s['id'];v=self.call(f'/api/projects/{pid}/subcontract');sid=v['pakd']['current']
        lines=LINES[:1]+[dict(LINES[1],source_ref='=HYPERLINK("http://x","TEST")')]+LINES[2:]
        v=self.call(f'/api/projects/{pid}/pakd',{'revision':v['revision'],'action':'save','scenario_id':sid,'basis':{'id-GT-01':'AWARD'},'lines':lines,'cashflow':PakdDomain.CF})
        v=self.call(f'/api/projects/{pid}/pakd',{'revision':v['revision'],'action':'freeze','scenario_id':sid,'name':'TEST chốt','by':'TEST'})
        wb=load_workbook(io.BytesIO(self.call(f'/api/projects/{pid}/pakd.xlsx?scenario={sid}',raw=True)))
        c=next(c for r in wb['DONG_NHAP'].iter_rows() for c in r if isinstance(c.value,str) and 'HYPERLINK' in c.value);self.assertEqual(c.data_type,'s')
        p=self.app.store.get(pid);r=self.call('/api/restore',{'schema_version':p.get('schema_version') or '0.2.0','project':p,'audit':[]})['project']
        self.assertEqual(len(r['pakd']['versions']),1);self.assertTrue(r['pakd']['versions'][0]['restored'])
if __name__=='__main__':unittest.main(verbosity=1)
