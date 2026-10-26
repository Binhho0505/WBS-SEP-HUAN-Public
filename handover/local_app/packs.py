"""E16 (D59): export packs.

- bb_workbook(): BOQ BB of one package for subcontractors (BB code, WBS keys, quantities, empty price column) or the
  internal version with the three suggested-price bases.
- tender_pack(): ZIP for the tender department (estimate file): status of the 8 gates, AB BOQ, WBS links, CBS, open items,
  PM-model checklist, the full estimate workbook when all gates passed, project backup, manifest.
- procurement_pack(): ZIP for the subcontract procurement department: package plan (CC1 template columns), BB per package
  (send + internal), quote comparison and review findings, DRAFT award proposals, coverage and WBS reconciliation
  AB = Σ BB + CC1, checklist, manifest.
Everything is labelled DỰ THẢO LOCAL — not a corporate release, not an approval.
"""
import hashlib,io,json,re,zipfile
from datetime import datetime
from decimal import Decimal
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from domain import number,index,status,calculate,digest,STEPS,CATEGORIES
import subcontract as sc,pmmodels,pakd
LABEL='DỰ THẢO LOCAL — không phải hồ sơ phát hành chính thức / không phải phê duyệt'
HEAD=Font(bold=True,color='FFFFFF');FILL=PatternFill('solid',fgColor='006D46')
def num(x):
    v=number(x);return float(v) if v is not None else None
def sheet(wb,title,cols,rows,first=False):
    ws=wb.active if first else wb.create_sheet(title)
    if first:ws.title=title
    ws.append(cols)
    for c in ws[1]:c.font=HEAD;c.fill=FILL;c.alignment=Alignment(wrap_text=True,vertical='top')
    for r in rows:ws.append([json.dumps(v,ensure_ascii=False) if isinstance(v,(list,dict)) else v for v in r])
    for row in ws.iter_rows(min_row=2):  # D77: text that looks like a formula stays text (formula injection)
        for c in row:
            if isinstance(c.value,str) and c.value.startswith(('=','+','-','@')):c.data_type='s'
    ws.freeze_panes='A2'
    for i,c in enumerate(cols,1):ws.column_dimensions[ws.cell(1,i).column_letter].width=max(10,min(48,len(str(c))+4))
    return ws
def wb_bytes(wb):
    bio=io.BytesIO();wb.save(bio);return bio.getvalue()
def _pkg(s,code):
    p=next((x for x in s.get('packages') or [] if x['code']==code),None)
    if not p:raise KeyError('Không có gói '+str(code))
    return p
def bb_workbook(s,code,internal=False,sub=None):
    p=_pkg(s,code);ab=index(s['boq']);d=p.get('delivery') or {};wb=Workbook()
    info=[['Hồ sơ',s['name']],['Mã dự án',s['context'].get('project_code')],['Gói',f"{p['code']} — {p.get('name') or ''}"],['Phạm vi',p.get('scope_note')],
          ['Hình thức giao thầu',d.get('mode')],['Vật liệu (VL)','NTP' if d.get('mode')=='GOP' else d.get('VL')],['Nhân công (NC)','NTP' if d.get('mode')=='GOP' else d.get('NC')],['Ca máy (M)','NTP' if d.get('mode')=='GOP' else d.get('M')],
          ['Hình thức giá',d.get('price_form')],['Nguồn lực dùng chung',json.dumps(d.get('shared') or [],ensure_ascii=False)],['Revision hồ sơ',s['revision']],['Trạng thái',LABEL],
          ['Hướng dẫn','Điền cột Đơn giá (trước VAT) cho từng Mã dòng BB; không đổi Mã dòng BB/Mã WBS; khối lượng khác thì ghi vào cột Khối lượng NTP đề xuất và giải thích ở Ghi chú.']]
    ws=sheet(wb,'THONG_TIN_GOI',['Mục','Nội dung'],info,first=True);ws.column_dimensions['B'].width=90
    cols=['Mã dòng BB','Mã WBS','Work Result','LBS','Mô tả','Đơn vị','Khối lượng','Đơn giá']+(['Khối lượng NTP đề xuất','Ghi chú'] if not internal else ['Giá gợi ý A (AB×K)','Căn cứ A','Giá gợi ý B (CC1+phí)','Căn cứ B','Giá gợi ý C (báo giá)','Căn cứ C','Chênh lệch %','Dòng AB (id)','Tỷ lệ phân bổ %'])
    basis=((sub or {}).get('basis') or {}) if internal else {}
    rows=[]
    for l in p.get('lines') or []:
        a=ab.get(l['ab_line_id']) or {};q=number(a.get('quantity'));al=number(l.get('alloc_pct'))
        bbq=float(q*al/100) if q is not None and al is not None else None;sn=l['snapshot']
        r=[l['bb_code'],sn.get('wbs_code'),', '.join(sn.get('work_results') or []),', '.join(sn.get('lbs') or []),a.get('description'),a.get('uom'),bbq,None]
        if internal:
            b=basis.get(l['bb_code'],{})
            for k in 'ABC':r+=[num(b.get(k,{}).get('value')),b.get(k,{}).get('reason')]
            r+=[num(b.get('spread_pct')),l['ab_line_id'],num(l.get('alloc_pct'))]
        else:r+=[None,None]
        rows.append(r)
    sheet(wb,'BOQ_BB',cols,rows)
    return wb_bytes(wb)
def _zip(files):
    bio=io.BytesIO();man=[]
    with zipfile.ZipFile(bio,'w',zipfile.ZIP_DEFLATED) as z:
        for name,data in files:
            z.writestr(name,data);man.append({'file':name,'sha256':hashlib.sha256(data).hexdigest(),'bytes':len(data)})
        z.writestr('MANIFEST.json',json.dumps({'status':LABEL,'created':datetime.now().isoformat(timespec='seconds'),'files':man},ensure_ascii=False,indent=1))
    return bio.getvalue()
def _model(s,store):
    ms=pmmodels.models(store);return ms.get(s['context'].get('pm_model')) or {}
def tender_pack(s,store,events,estimate_bytes=None):
    st=status(s);calc=calculate(s);m=_model(s,store);sub=sc.summary(s,pmmodels.models(store)) if s.get('packages') else {}
    wb=Workbook()
    sheet(wb,'TRANG_THAI',['Bước','Tên','Đã xác nhận','Số câu hỏi mở'],[[x['step']+1,x['title'],'CÓ' if x['confirmed'] else 'CHƯA',x['question_count']] for x in st],first=True)
    wb['TRANG_THAI'].append([]);wb['TRANG_THAI'].append(['Trạng thái hồ sơ','ĐỦ GATE 8 BƯỚC' if all(x['confirmed'] for x in st) else 'DỰ THẢO — còn bước chưa xác nhận']);wb['TRANG_THAI'].append([LABEL])
    sheet(wb,'BOQ_AB',['id','Mã WBS (ERP)','Mô tả','Đơn vị','Khối lượng','Đơn giá HĐ','Loại','Nguồn','Hash'],[[r['id'],r.get('source_code'),r.get('description'),r.get('uom'),num(r.get('quantity')),num(r.get('unit_price')),r.get('kind'),r.get('source_ref'),r.get('source_hash')] for r in s['boq']])
    sheet(wb,'WBS_OCCURRENCE',['id','Work Result','LBS','Vùng giá','Phạm vi','Nghiệm thu','CA','CAM'],[[w.get('id'),w.get('work_result_code'),w.get('lbs'),w.get('price_zone'),w.get('scope'),w.get('acceptance'),w.get('ca'),w.get('cam')] for w in s.get('wbs') or []])
    boq=index(s['boq']);occ=index(s.get('wbs') or [])
    sheet(wb,'LIEN_KET_WBS',['Mã WBS (ERP)','Mô tả','Work Result','LBS','Trọng số','Khối lượng','Đơn vị','Biện pháp'],[[(boq.get(x.get('source_line_id')) or {}).get('source_code'),(boq.get(x.get('source_line_id')) or {}).get('description'),(occ.get(x.get('wbs_id')) or {}).get('work_result_code'),(occ.get(x.get('wbs_id')) or {}).get('lbs'),x.get('weight'),num(x.get('quantity')),x.get('uom'),x.get('method')] for x in s.get('mapping') or []])
    sheet(wb,'CBS',['CBS','Tên','Giá trị đã biết'],[[k,v,num(calc['cbs'].get(k))] for k,v in CATEGORIES.items()]+[['TỔNG','Đã biết' if not calc['complete'] else 'Tổng dự toán',num(calc['total'] if calc['complete'] else calc['known_amount'])],['Trạng thái',calc['status'],None]])
    open_items=[[x['step']+1,x['title'],q['field'],q['question']] for x in st for q in x['questions']]+[['GT','Gói thầu',q['field'],q['question']] for q in (sub.get('questions') or [])]
    sheet(wb,'VIEC_CON_MO',['Bước','Tên','Trường','Câu hỏi'],open_items)
    sheet(wb,'CHECKLIST_QLDA',['Mã','Giai đoạn','Hồ sơ','Trạng thái'],[[c['code'],c['stage'],c['name'],c['state']] for c in pmmodels.checklist(s,m,st,sub)] if m else [['—','—','Chưa chọn mô hình QLDA cho hồ sơ','MISSING']])
    wb['CHECKLIST_QLDA'].append([]);wb['CHECKLIST_QLDA'].append(['Mô hình',m.get('name') if m else None,m.get('status') if m else None])
    files=[('01_HO_SO_LAP_DU_TOAN_TONG_HOP.xlsx',wb_bytes(wb))]
    if estimate_bytes:files.append(('02_DU_TOAN_DAY_DU_TPL03.xlsx',estimate_bytes))
    files.append(('03_SAO_LUU_HO_SO.json',json.dumps({'project':s,'audit_events':len(events or []),'digest':digest(s)},ensure_ascii=False,indent=1).encode()))
    return _zip(files)
def procurement_pack(s,store):
    if not s.get('packages'):raise ValueError('Hồ sơ chưa có gói thầu con.')
    models=pmmodels.models(store);m=models.get(s['context'].get('pm_model')) or {};sub=sc.summary(s,models);st=status(s)
    ab=index(s['boq']);bysub={x['code']:x for x in sub['packages']};wb=Workbook()
    plan=[]
    for p in s['packages']:
        d=p.get('delivery') or {};label=(m.get('procurement_methods') or {}).get(p.get('procurement_method'),p.get('procurement_method'))
        plan.append([p['code'],p.get('name'),'CC1 tự làm' if p.get('is_self') else f"{label or 'chưa chọn hình thức chọn thầu'}; giao {d.get('mode') or '?'}",None,f"{len(p.get('lines') or [])} dòng BB; {len([q for q in s.get('quotes') or [] if q['package_id']==p['id']])} báo giá",(p.get('award') or {}).get('proposed_by'),'DRAFT_AWARD' if p.get('award') else 'OPEN'])
    sheet(wb,'04_PROCUREMENT_PACKAGE_PLAN',['Procurement_ID','Package','Strategy_Hypothesis','Milestone','Evidence','Authority_or_Review','Status'],plan,first=True)
    rows=[]
    for p in s['packages']:
        d=p.get('delivery') or {};b=bysub[p['code']]['basis'];val=Decimal(0);tot={'A':Decimal(0),'B':Decimal(0),'C':Decimal(0)};full={'A':True,'B':True,'C':True}
        for l in p.get('lines') or []:
            a=ab.get(l['ab_line_id']) or {};q=number(a.get('quantity'));al=number(l.get('alloc_pct'));bbq=q*al/100 if q is not None and al is not None else None;up=number(a.get('unit_price'))
            if bbq is not None and up is not None:val+=bbq*up
            for k in 'ABC':
                v=number((b.get(l['bb_code']) or {}).get(k,{}).get('value'))
                if v is None or bbq is None:full[k]=False
                else:tot[k]+=v*bbq
        aw=p.get('award') or {};awt=sum((number(v)*(number((ab.get(l['ab_line_id']) or {}).get('quantity')) or 0)*(number(l.get('alloc_pct')) or 0)/100 for l in p.get('lines') or [] for k,v in (aw.get('prices') or {}).items() if k==l['bb_code'] and number(v) is not None),Decimal(0))
        rows.append([p['code'],p.get('name'),p.get('procurement_method'),d.get('mode'),d.get('VL'),d.get('NC'),d.get('M'),d.get('price_form'),num(d.get('fee_pct')),d.get('fee_ref'),len(p.get('lines') or []),float(val)]+[float(tot[k]) if full[k] else None for k in 'ABC']+[aw.get('vendor'),float(awt) if aw else None])
    sheet(wb,'GOI_THAU',['Gói','Tên','Hình thức chọn thầu','Giao thầu','VL','NC','M','Hình thức giá','Phí %','Căn cứ phí','Số dòng BB','Giá trị theo đơn giá AB','Tổng gợi ý A','Tổng gợi ý B','Tổng gợi ý C','NTP đề xuất','Tổng giá giao đề xuất'],rows)
    wb['GOI_THAU'].append([]);wb['GOI_THAU'].append(['Tổng gợi ý để trống khi có dòng thiếu căn cứ (không lấp số thiếu).'])
    sheet(wb,'NGUON_LUC_DUNG_CHUNG',['Gói','Mã nguồn lực','Bên cấp','Phân bổ','Căn cứ'],[[p['code'],r.get('resource_code'),r.get('provider'),r.get('allocation'),r.get('ref')] for p in s['packages'] for r in (p.get('delivery') or {}).get('shared') or []])
    cmp=[]
    for p in s['packages']:
        qs=[q for q in s.get('quotes') or [] if q['package_id']==p['id']];b=bysub[p['code']]['basis']
        for l in p.get('lines') or []:
            x=b.get(l['bb_code'],{});row=[p['code'],l['bb_code'],l['snapshot'].get('wbs_code'),(ab.get(l['ab_line_id']) or {}).get('description')]+[num(x.get(k,{}).get('value')) for k in 'ABC']
            row+=['; '.join(f"{q['vendor']}: {next((r.get('unit_price') for r in q['rows'] if r.get('bb_code')==l['bb_code']),'—')}" for q in qs),num((p.get('award') or {}).get('prices',{}).get(l['bb_code']))]
            cmp.append(row)
    sheet(wb,'SO_SANH_GIA',['Gói','Mã dòng BB','Mã WBS','Mô tả','Gợi ý A (AB×K)','Gợi ý B (CC1+phí)','Gợi ý C (báo giá)','Báo giá các NTP','Giá giao đề xuất'],cmp)
    sheet(wb,'THAM_DINH_BAO_GIA',['Gói','Nhà thầu','Hợp lệ','Tổng theo KL BB','Mã dòng BB','Loại','Chặn','Nội dung'],[[x['code'],r['vendor'],'CÓ' if r['valid'] else 'KHÔNG',num(r['total']),f.get('bb_code'),f['type'],'CÓ' if f['blocking'] else '',f['text']] for x in sub['packages'] for r in x['reviews'] for f in (r['findings'] or [{'bb_code':None,'type':'OK','blocking':False,'text':'Không có phát hiện'}])])
    sheet(wb,'DE_XUAT_GIAO_THAU',['Gói','Nhà thầu','Lý do','Người đề xuất','Ngày','Trạng thái','Ghi chú'],[[p['code'],p['award'].get('vendor'),p['award'].get('reason'),p['award'].get('proposed_by'),p['award'].get('date'),p['award'].get('status'),p['award'].get('note')] for p in s['packages'] if p.get('award')])
    sheet(wb,'PHU_KHOI_LUONG',['Dòng AB','Mã WBS','Mô tả','Đơn vị','KL AB','KL đã phân bổ','% phân bổ','Trạng thái','Phân bổ'],[[r['ab_line_id'],r['wbs_code'],r['description'],r['uom'],num(r['ab_quantity']),num(r['bb_quantity']),num(r['allocated_pct']),r['state'],'; '.join(f"{x['bb_code']} {x['alloc_pct']}%" for x in r['parts'])] for r in sub['reconcile']])
    groups={}
    for r in sub['reconcile']:
        g=groups.setdefault(r['wbs_code'] or '(không mã)',{'ab':Decimal(0),'bb':Decimal(0),'n':0,'ok':0});g['n']+=1;g['ok']+=r['state']=='OK'
        g['ab']+=number(r['ab_quantity']) or 0;g['bb']+=number(r['bb_quantity']) or 0
    sheet(wb,'DOI_CHIEU_WBS_AB_BB',['Mã WBS','Số dòng AB','Dòng phủ đủ','Σ KL AB','Σ KL BB + CC1','Lệch'],[[k,g['n'],g['ok'],float(g['ab']),float(g['bb']),float(g['bb']-g['ab'])] for k,g in sorted(groups.items())])
    sheet(wb,'HOAN_THANH_THEO_WBS',['Mã WBS','Số dòng','% hoàn thành','Phương pháp','Trạng thái'],[[g['wbs'],g['lines'],num(g['pct']),g['method'],g['state']] for g in sub['progress']['wbs']])
    sheet(wb,'CHECKLIST_QLDA',['Mã','Giai đoạn','Hồ sơ','Trạng thái'],[[c['code'],c['stage'],c['name'],c['state']] for c in pmmodels.checklist(s,m,st,sub)] if m else [['—','—','Chưa chọn mô hình QLDA','MISSING']])
    sheet(wb,'VIEC_CON_MO',['Trường','Câu hỏi'],[[q['field'],q['question']] for q in sub['questions']])
    files=[('01_HO_SO_CHON_THAU_GIAO_THAU.xlsx',wb_bytes(wb))]
    for p in s['packages']:
        if p.get('is_self'):continue
        files.append((f"BOQ_BB/{p['code']}_GUI_NHA_THAU_PHU.xlsx",bb_workbook(s,p['code'])))
        files.append((f"BOQ_BB/{p['code']}_NOI_BO_GIA_GOI_Y.xlsx",bb_workbook(s,p['code'],True,bysub[p['code']])))
    files+=pakd_files(s)
    import contracts
    for c in s.get('contracts') or []:files.append((f"HOP_DONG/HO_SO_HD_NTP_{c['package_code']}_{c['status']}.xlsx",contracts.workbook(s,c['id'])))
    return _zip(files)
def _safe(name):return re.sub(r'[^0-9A-Za-zÀ-ỹ._-]+','_',str(name)).strip('_')[:60] or 'x'
def pakd_files(s):
    """PAKD BM01 for the procurement department (BL100): the scenario being worked on (live, labelled DANG_LAP) and every frozen
    version exactly as frozen (CHOT). Nothing is attached when PAKD was never started, so an empty BM01 is not sent out."""
    if not s.get('pakd'):return []
    st=pakd.state(s);cur=next(c for c in st['scenarios'] if c['id']==st['current'])
    out=[(f"PAKD/PAKD_BM01_DANG_LAP_{_safe(cur['name'])}.xlsx",pakd.workbook(s,cur['id']))]
    for v in st['versions']:out.append((f"PAKD/PAKD_BM01_CHOT_{_safe(v['name'])}_{v['inputs_sha256'][:8]}.xlsx",pakd.workbook(s,None,v['id'])))
    return out
