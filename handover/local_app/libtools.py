"""E13 (D59): CC1 data library tools.

- derive(): a CC1 dataset version made from a state dataset by coefficients (general, per part VL/NC/M, per code prefix)
  and/or manual row overrides with a reason. The source version is never changed; every row keeps its lineage.
- Code dictionaries: one workbook with every code set (MASTER work results, CBS schemes, LBS, families, resource master,
  BOQ→TT38 rules, construction methods, PM models, and the open project's WBS/mapping/packages) for people to check and
  edit by hand; an edited sheet comes back as a new REFERENCE version after a dry-run diff. MASTER V1.20 stays read-only.
"""
import base64,csv,hashlib,io,json,re
from decimal import Decimal
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from domain import number,catalog,CATEGORIES,digest
METHOD_NAME='Từ điển biện pháp thi công CC1';CBS_NAME='CBS CC1';LBS_NAME='LBS CC1';PM_NAME='Mô hình QLDA CC1'
DICT_NAME='Từ điển BOQ→TT38 CC1 (AI-QS proxy draft)';RES_NAME='Danh mục nguồn lực CC1 từ TT38 + mapping (NC/MTC/VL/PCT)'
TARGET={'PRICES':'unit_price','RESOURCES':'coefficient'}
def _s(d):
    if d is None:return None
    t=format(d,'f')
    return t.rstrip('0').rstrip('.') if '.' in t else t
def part_of(r):
    c=str(r.get('resource_code') or r.get('code') or '').upper()
    if c.startswith('VL'):return 'VL'
    if c.startswith('NC'):return 'NC'
    if c.startswith('MTC') or re.match(r'^M\d',c):return 'M'
    return {'C10':'VL','C20':'NC','C30':'M'}.get(str(r.get('cbs') or ''))
def _active(store,kind,name):
    return next((d for d in store.library_list() if d['kind']==kind and d['name']==name and d['status']=='ACTIVE'),None)
# ---------- derived datasets (FR-13, FR-14) ----------
def derive(store,did,body):
    src=store.library_get(did,False);kind=src['kind']
    if kind not in TARGET:raise ValueError('Chỉ tạo bộ CC1 từ bộ PRICES (đơn giá) hoặc RESOURCES (hao phí); NORMS không có số để nhân hệ số.')
    name=str(body.get('name') or '').strip();basis=str(body.get('basis_ref') or '').strip()
    if not name or len(name)>200:raise ValueError('Nhập tên bộ CC1 mới (1–200 ký tự).')
    if name==src['name']:raise ValueError('Tên bộ CC1 phải khác tên bộ nguồn (bộ nguồn giữ nguyên).')
    if not basis:raise ValueError('Nhập căn cứ áp dụng hệ số/chỉnh tay.')
    f=body.get('factors') or {};fac={}
    for k in ('all','VL','NC','M'):
        if f.get(k) not in (None,''):
            v=number(f.get(k))
            if v is None or v<=0:raise ValueError(f'Hệ số {k} phải là số > 0 (để trống nếu không áp dụng).')
            fac[k]=v
    pre={}
    for p,k in (f.get('prefixes') or {}).items():
        v=number(k)
        if not str(p).strip() or v is None or v<=0:raise ValueError('Hệ số theo tiền tố mã phải có tiền tố và số > 0.')
        pre[str(p).strip().upper()]=v
    ov=body.get('overrides') or []
    if not isinstance(ov,list):raise ValueError('Các dòng chỉnh tay phải là danh sách.')
    if not fac and not pre and not ov:raise ValueError('Nhập ít nhất một hệ số hoặc một dòng chỉnh tay.')
    field=TARGET[kind];rows=[dict(r) for r in store.library_rows(did)];byid={r.get('id'):r for r in rows}
    changed=nulls=0;unmatched=set(pre)
    for r in rows:
        x=number(r.get(field));code=str(r.get('resource_code') or r.get('code') or '').upper()
        k=fac.get('all',Decimal(1))*fac.get(part_of(r) or '-',Decimal(1))
        hit=max((p for p in pre if code.startswith(p)),key=len,default=None)
        if hit:k*=pre[hit];unmatched.discard(hit)
        r['derived']={'from_dataset':did,'from_sha256':src['sha256'],'factor':_s(k),'original':r.get(field)}
        if k==1:continue
        if x is None:nulls+=1;continue  # a missing value stays missing; never 0
        r[field]=_s(x*k);changed+=1
    done=[]
    for o in ov:
        r=byid.get(o.get('row_id'));fld=o.get('field') or field;reason=str(o.get('reason') or '').strip()
        if r is None:raise ValueError('Dòng chỉnh tay không có trong bộ nguồn: '+str(o.get('row_id')))
        if not reason:raise ValueError('Mỗi dòng chỉnh tay cần lý do.')
        if fld=='cbs':
            if o.get('value') not in CATEGORIES:raise ValueError('CBS phải là một trong '+', '.join(CATEGORIES))
            val=o['value']
        elif fld==field:
            v=number(o.get('value'))
            if v is None or v<=0:raise ValueError('Giá trị chỉnh tay phải là số > 0.')
            val=_s(v)
        else:raise ValueError(f'Chỉ chỉnh tay {field} hoặc cbs.')
        r['derived']=dict(r['derived'],override_field=fld,override_original=r.get(fld),override_reason=reason[:500]);r[fld]=val;done.append(o.get('row_id'))
    lineage={'source':{'id':did,'name':src['name'],'version':src['version'],'sha256':src['sha256']},'factors':{k:_s(v) for k,v in fac.items()},'prefixes':{k:_s(v) for k,v in pre.items()},'overrides':len(done),'basis':basis}
    note=('CC1 DERIVED · '+json.dumps(lineage,ensure_ascii=False))[:1000]
    d=store.library_add(kind,name,digest(rows),'derived:'+did,note,rows)
    return {'dataset':d,'stats':{'rows':len(rows),'changed':changed,'null_kept':nulls,'overrides':len(done),'unmatched_prefixes':sorted(unmatched)}}
# ---------- dictionaries workbook (FR-18) ----------
METHOD_TEMPLATE=[{'code':'THU_CONG','name':'Đổ bê tông thủ công (vữa trộn tại chỗ)','work_group':'Bê tông','norm_param':'METHOD=THU_CONG → TT38 AF.1','applies_to':'AF.1','note':'Mẫu DRAFT từ cách chia chương VI TT38 đã dùng trong engine gợi ý','status':'DRAFT'},
                 {'code':'CAN_CAU','name':'Đổ bê tông bằng cần cẩu','work_group':'Bê tông','norm_param':'METHOD=CAN_CAU → TT38 AF.2','applies_to':'AF.2','note':'Mẫu DRAFT','status':'DRAFT'},
                 {'code':'BOM','name':'Đổ bê tông bằng máy bơm','work_group':'Bê tông','norm_param':'METHOD=BOM → TT38 AF.3','applies_to':'AF.3','note':'Mẫu DRAFT','status':'DRAFT'}]
# sheet → (dataset name, key column, parent column, editable)
SHEETS={'BIEN_PHAP':(METHOD_NAME,'code',None,True),'CBS_CC1':(CBS_NAME,'code','parent',True),'LBS':(LBS_NAME,'code','parent',True),
        'BOQ_TT38':(DICT_NAME,'id',None,True),'NGUON_LUC':(RES_NAME,'id',None,True),'MO_HINH_QLDA':(PM_NAME,'code',None,True),
        'WORK_RESULT':(None,'business_code',None,False),'CBS_SCHEMES':(None,'business_code',None,False),'FAMILIES':(None,'family_id',None,False),
        'P_WBS':(None,'id',None,False),'P_MAPPING':(None,'id',None,False),'P_GOI_THAU':(None,'bb_code',None,False)}
def _cell(v):return json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v
def _rows_of(store,name):
    d=next((x for x in store.library_list() if x['kind']=='REFERENCE' and x['name']==name and x['status']=='ACTIVE'),None)
    return (store.library_rows(d['id']),d) if d else (None,None)
def dictionary_sheets(store,s=None):
    import pmmodels
    c=catalog();out=[]
    def add(sheet,rows,src):out.append((sheet,rows,src))
    add('WORK_RESULT',[{k:r.get(k) for k in ('business_code','name','level','parent_code','node_kind','readiness','lifecycle_status','definition_available','source_id')} for r in c['WORK_RESULTS']],'dictionary.json WORK_RESULTS (MASTER V1.20 IN-REVIEW, chỉ đọc)')
    add('CBS_SCHEMES',c['CBS_SCHEMES'],'dictionary.json CBS_SCHEMES (4 lược đồ, chưa có crosswalk tài chính)')
    rows,d=_rows_of(store,CBS_NAME)
    add('CBS_CC1',rows if rows is not None else [{'code':r['business_code'],'name':r['name'],'parent':'','scheme':r['scheme_id'],'source_ref':r.get('source_refs'),'status':'DRAFT'} for r in c['CBS_SCHEMES'] if r['scheme_id']=='MD_20261009'],
        f'thư viện {CBS_NAME} v{d["version"]}' if d else 'mẫu từ CBS_SCHEMES MD_20261009 (chưa nạp thư viện)')
    rows,d=_rows_of(store,LBS_NAME)
    add('LBS',rows if rows is not None else [{'code':r['LBS_Code'],'name':r.get('Tên phạm vi'),'parent':next((x['LBS_Code'] for x in c['SOURCE_LBS'] if x['LBS_ID']==r.get('Parent_LBS_ID')),''),'level':r.get('LBS_Level'),'type':r.get('LBS_Type'),'source_ref':f"{r.get('source_id')} row {r.get('source_row')}",'status':r.get('Status')} for r in c['SOURCE_LBS']],
        f'thư viện {LBS_NAME} v{d["version"]}' if d else 'dictionary.json SOURCE_LBS (TPL02/TPL03, 44 mã)')
    add('FAMILIES',c['FAMILIES'],'dictionary.json FAMILIES (họ mã, chỉ đọc)')
    for sheet,name in (('NGUON_LUC',RES_NAME),('BOQ_TT38',DICT_NAME)):
        rows,d=_rows_of(store,name);add(sheet,rows or [],f'thư viện {name} v{d["version"]}' if d else 'chưa nạp thư viện')
    rows,d=_rows_of(store,METHOD_NAME);add('BIEN_PHAP',rows if rows is not None else METHOD_TEMPLATE,f'thư viện {METHOD_NAME} v{d["version"]}' if d else 'mẫu DRAFT 3 biện pháp đổ bê tông (chưa nạp thư viện)')
    add('MO_HINH_QLDA',[{'code':k,'name':m['name'],'status':m['status'],'json':{x:y for x,y in m.items() if x not in ('dataset',)}} for k,m in pmmodels.models(store).items()],'mô hình QLDA (DRAFT_NEEDS_LEGAL_REVIEW)')
    if s:
        add('P_WBS',s.get('wbs') or [],'hồ sơ '+s['name']+' — WBS occurrence (sửa trong app, bước 3)')
        boq={r['id']:r for r in s['boq']};occ={w.get('id'):w for w in s.get('wbs') or []}
        add('P_MAPPING',[{'id':m.get('id'),'wbs_code':(boq.get(m.get('source_line_id')) or {}).get('source_code'),'description':(boq.get(m.get('source_line_id')) or {}).get('description'),'work_result_code':(occ.get(m.get('wbs_id')) or {}).get('work_result_code'),'lbs':(occ.get(m.get('wbs_id')) or {}).get('lbs'),'weight':m.get('weight'),'quantity':m.get('quantity'),'uom':m.get('uom')} for m in s.get('mapping') or []],'hồ sơ — BOQ AB → WBS (sửa trong app, bước 3)')
        add('P_GOI_THAU',[{'bb_code':l['bb_code'],'package':p['code'],'package_name':p.get('name'),'ab_line_id':l['ab_line_id'],'wbs_code':l['snapshot'].get('wbs_code'),'work_results':', '.join(l['snapshot'].get('work_results') or []),'lbs':', '.join(l['snapshot'].get('lbs') or []),'alloc_pct':l.get('alloc_pct')} for p in s.get('packages') or [] for l in p.get('lines') or []],'hồ sơ — dòng BB (sửa trong tab Gói thầu)')
    return out
def dictionaries_workbook(store,s=None):
    sheets=dictionary_sheets(store,s);wb=Workbook();ws=wb.active;ws.title='00_HUONG_DAN'
    head=Font(bold=True,color='FFFFFF');fill=PatternFill('solid',fgColor='006D46')
    ws.append(['Sheet','Nguồn','Số dòng','Nhập lại được','Bộ thư viện đích','Khóa'])
    for c in ws[1]:c.font=head;c.fill=fill
    for sheet,rows,src in sheets:
        name,key,_,ed=SHEETS[sheet];ws.append([sheet,src,len(rows),'CÓ' if ed else 'KHÔNG (chỉ đối chiếu)',name or '',key])
    ws.append([]);ws.append(['Cách dùng: sửa trực tiếp các sheet "Nhập lại được", giữ nguyên dòng tiêu đề và cột khóa; nạp lại một sheet ở Thư viện dữ liệu → "Nhập từ điển" → xem khác biệt → lưu thành phiên bản mới (bản cũ được lưu trữ, không xóa).'])
    ws.append(['MASTER V1.20 (WORK_RESULT) đang IN-REVIEW theo governance: chỉ đối chiếu, sửa qua quy trình MASTER. Cột dạng [..]/{..} là JSON.'])
    ws.append(['DỰ THẢO LOCAL — không phải bộ mã ban hành chính thức.']);ws.column_dimensions['A'].width=16;ws.column_dimensions['B'].width=70
    for sheet,rows,src in sheets:
        w=wb.create_sheet(sheet);cols=[]
        for r in rows:
            for k in r:
                if k not in cols:cols.append(k)
        if not cols:cols=list(METHOD_TEMPLATE[0]) if sheet=='BIEN_PHAP' else [SHEETS[sheet][1]]
        w.append(cols)
        for c in w[1]:c.font=head;c.fill=fill;c.alignment=Alignment(wrap_text=True)
        for r in rows:w.append([_cell(r.get(k)) for k in cols])
        w.freeze_panes='A2'
    bio=io.BytesIO();wb.save(bio);return bio.getvalue()
# ---------- import an edited dictionary sheet ----------
def _table(filename,data,sheet):
    if str(filename).lower().endswith('.csv'):
        t=list(csv.reader(io.StringIO(data.decode('utf-8-sig',errors='replace'))))
    else:
        wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
        if sheet not in wb.sheetnames:raise ValueError(f'File không có sheet {sheet}.')
        t=[list(r) for r in wb[sheet].iter_rows(values_only=True)]
    if not t:raise ValueError('Sheet rỗng.')
    cols=[str(c).strip() if c is not None else '' for c in t[0]];rows=[]
    for r in t[1:]:
        if all(v in (None,'') for v in r):continue
        rows.append({c:(None if i>=len(r) or r[i] in (None,'') else (r[i] if isinstance(r[i],str) else _s(Decimal(str(r[i]))) if isinstance(r[i],(int,float)) and not isinstance(r[i],bool) else str(r[i]))) for i,c in enumerate(cols) if c})
    return cols,rows
def import_dictionary(store,body):
    sheet=str(body.get('sheet') or '').strip()
    if sheet not in SHEETS:raise ValueError('Sheet không thuộc bộ từ điển: '+sheet)
    name,key,parent,ed=SHEETS[sheet]
    if not ed:raise ValueError(f'Sheet {sheet} chỉ để đối chiếu (MASTER V1.20 / họ mã / dữ liệu hồ sơ): sửa qua quy trình MASTER hoặc trong app, không nhập ở đây.')
    data=base64.b64decode(body.get('data') or '',validate=True)
    if not data or len(data)>35*1024*1024:raise ValueError('File rỗng hoặc vượt 35 MB.')
    cols,rows=_table(body.get('filename') or 'x.xlsx',data,sheet);errors=[]
    if key not in cols:raise ValueError(f'Thiếu cột khóa {key}.')
    cur,curd=_rows_of(store,name);cur=cur or []
    listish={k for r in cur for k,v in r.items() if isinstance(v,(list,dict))}|({'json'} if sheet=='MO_HINH_QLDA' else set())
    seen=set()
    for i,r in enumerate(rows,2):
        for k in listish:
            if isinstance(r.get(k),str) and r[k][:1] in '[{':
                try:r[k]=json.loads(r[k])
                except ValueError:errors.append(f'Dòng {i}: cột {k} không phải JSON hợp lệ.')
        k=r.get(key)
        if not k:errors.append(f'Dòng {i}: thiếu {key}.');continue
        if k in seen:errors.append(f'Dòng {i}: {key} trùng "{k}".')
        seen.add(k)
        if sheet=='BOQ_TT38':
            for f in ('match','ctx'):
                try:re.compile(r.get(f) or '')
                except re.error:errors.append(f'Dòng {i}: {f} không phải biểu thức chính quy (regex) hợp lệ.')
    if parent and parent in cols:
        for i,r in enumerate(rows,2):
            if r.get(parent) and r[parent] not in seen:errors.append(f'Dòng {i}: {parent} "{r[parent]}" không có trong sheet.')
    if sheet=='CBS_CC1':
        for i,r in enumerate(rows,2):
            if not re.fullmatch(r'[A-Z0-9][A-Z0-9.\-]{0,30}',str(r.get('code') or '')):errors.append(f'Dòng {i}: mã CBS "{r.get("code")}" sai định dạng.')
    old={r.get(key):r for r in cur};new={r.get(key):r for r in rows if r.get(key)}
    norm=lambda r:{k:v for k,v in r.items() if v not in (None,'')}
    diff={'added':sorted(set(new)-set(old))[:200],'removed':sorted(set(old)-set(new))[:200],'changed':sorted(k for k in set(new)&set(old) if norm(new[k])!=norm(old[k]))[:200]}
    counts={'added':len(set(new)-set(old)),'removed':len(set(old)-set(new)),'changed':sum(1 for k in set(new)&set(old) if norm(new[k])!=norm(old[k])),'rows':len(rows)}
    res={'sheet':sheet,'dataset':name,'current_version':curd['version'] if curd else None,'diff':diff,'counts':counts,'errors':errors[:200],'error_count':len(errors)}
    if body.get('dry_run',True):return res
    if errors:raise ValueError(f'{len(errors)} lỗi, chưa lưu: '+'; '.join(errors[:5]))
    if not any(counts[k] for k in ('added','removed','changed')):raise ValueError('Không có thay đổi so với phiên bản đang dùng.')
    sha=hashlib.sha256(data).hexdigest();store.upload(sha,str(body.get('filename') or sheet)[:200],data)
    d=store.library_add('REFERENCE',name,sha,str(body.get('filename') or '')[:200],('Nhập từ điển sheet '+sheet+' · '+str(body.get('note') or ''))[:1000],rows)
    if curd:store.library_archive(curd['id'])
    return dict(res,saved=d)
# ---------- new price datasets (FR-16) ----------
PRICE_REQUIRED=('resource_code','uom','unit_price','location','period','currency')
def validate_prices(records):
    """Errors for a PRICES upload: required columns, positive numeric price, YYYY-MM period, unique (code, spec, uom, location, period)."""
    if not records:return ['File không có dòng giá.']
    cols=set().union(*(r.keys() for r in records));miss=[c for c in PRICE_REQUIRED if c not in cols]
    if miss:return ['Thiếu cột bắt buộc: '+', '.join(miss)+' (cần resource_code, spec (tùy chọn), uom, unit_price, location, period YYYY-MM, currency).']
    errors=[];seen={}
    for i,r in enumerate(records,2):
        v=number(r.get('unit_price'))
        if v is None or v<=0:errors.append(f'Dòng {i}: unit_price "{r.get("unit_price")}" không phải số dương (để trống không được thay bằng 0).')
        if not re.fullmatch(r'\d{4}-\d{2}',str(r.get('period') or '')):errors.append(f'Dòng {i}: period phải dạng YYYY-MM.')
        for c in ('resource_code','uom','location','currency'):
            if not str(r.get(c) or '').strip():errors.append(f'Dòng {i}: thiếu {c}.')
        k=tuple(str(r.get(c) or '').strip() for c in ('resource_code','spec','uom','location','period'))
        if k in seen:errors.append(f'Dòng {i}: trùng khóa giá với dòng {seen[k]} ({", ".join(x for x in k if x)}).')
        else:seen[k]=i
    return errors
