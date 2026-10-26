import copy,json
from decimal import Decimal
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from domain import ROOT,calculate,digest,number,CATEGORIES
def source_template():
    directory=ROOT.parent/'MAU WBS MINH & TONG HOP'
    register=json.loads((ROOT/'runs_brian/enterprise_dictionary_20261017_01/source-register-templates.json').read_text(encoding='utf-8'))
    source=next(r for r in register if r['source_id']=='TPL03');p=directory/source['filename']
    if not p.exists():return None
    import hashlib
    if hashlib.sha256(p.read_bytes()).hexdigest()!=source['sha256']:raise ValueError('Mẫu nguồn đã đổi hash. Cần tiếp nhận successor và xác nhận bố cục trước khi xuất.')
    return p
def scalar(v):return json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v
def table(wb,name,rows,headers=None):
    ws=wb.create_sheet(name);keys=headers or list(dict.fromkeys(k for r in rows for k in r));ws.append(keys)
    for i,r in enumerate(rows,2):
        ws.append([scalar(r.get(k)) for k in keys])
        for j in range(1,len(keys)+1):
            c=ws.cell(i,j)
            if isinstance(c.value,str) and c.value.startswith(('=','+','-','@')):c.data_type='s'
    ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
    for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17365D')
    for j,key in enumerate(keys,1):ws.column_dimensions[ws.cell(1,j).column_letter].width=min(65,max(18,len(str(key))+3))
    return ws
def workbook(s,events,out):
    result=calculate(s);assert result['complete'];wb=Workbook();wb.remove(wb.active);raw={r['id']:r for r in s['boq']};parts=s['mapping']; template=source_template()
    if not template:raise ValueError('Thiếu file mẫu CC1 đã chỉ định; không tự thay mẫu.')
    from hashlib import sha256
    template_hash=sha256(template.read_bytes()).hexdigest()
    source=load_workbook(template,read_only=False,data_only=False);src=source['10_WP_BOQ_COMPOSITION'];ws=wb.create_sheet('10_WP_BOQ_COMPOSITION')
    for row in src.iter_rows(min_row=1,max_row=4,max_col=11):
        for c in row:
            target=ws.cell(c.row,c.column,c.value)
            for attribute in ['font','fill','border','alignment','protection']:setattr(target,attribute,copy.copy(getattr(c,attribute)))
            target.number_format=c.number_format
    for merged in src.merged_cells.ranges:
        if merged.max_row<=4:ws.merge_cells(str(merged))
    for key,value in src.column_dimensions.items():ws.column_dimensions[key].width=value.width
    source.close()
    ws['A2']='DỰ THẢO LOCAL · '+s['name']+' · normalized item có truy vết; chi phí mới ở DU_TOAN/TONG_HOP'
    wbs={w['id']:w for w in s['wbs']}
    for i,p in enumerate(parts,5):
        r=raw[p['source_line_id']];w=wbs[p['wbs_id']];amount=None
        if number(r.get('unit_price')) is not None and number(r.get('original_quantity',r.get('quantity'))) is not None:amount=str(number(r.get('original_quantity',r['quantity']))*number(r['unit_price'])*number(p['weight']))
        values=[w['lbs']+'.'+w['work_result_code'],w['scope'],r.get('source_code') or r['id'],r['description'],p['uom'],str(p['quantity']),r.get('unit_price'),amount,w['work_result_code'],p['mapping_ref'],'DỰ THẢO LOCAL; đơn giá/giá trị BOQ nguồn khác giá dự toán CC1; xem DU_TOAN']
        for j,v in enumerate(values,1):
            c=ws.cell(i,j,v)
            if isinstance(v,str) and v.startswith(('=','+','-','@')):c.data_type='s'
    ws.freeze_panes='E5';ws.auto_filter.ref=f'A4:K{4+len(parts)}'
    estimates=[];compare={r['item_id']:r for r in result['items']}
    for p in parts:
        lines=[r for r in result['lines'] if r['item_id']==p['id']];totals={c:sum((Decimal(r['amount']) for r in lines if r['cbs']==c),Decimal(0)) for c in CATEGORIES};total=sum(totals.values(),Decimal(0));w=wbs[p['wbs_id']]
        estimates.append(dict(item_id=p['id'],source_line_id=p['source_line_id'],WBS=p['wbs_id'],Work_Result=w['work_result_code'],LBS=w['lbs'],price_zone=w.get('price_zone') or s['context']['location'],description=raw[p['source_line_id']]['description'],quantity=str(p['quantity']),uom=p['uom'],unit_direct_cost=str(total/number(p['quantity'])),**{k:str(v) for k,v in totals.items()},direct_total=str(total),allocated_indirect=compare[p['id']]['allocated_indirect'],internal_cost=compare[p['id']]['internal_cost']))
    table(wb,'DU_TOAN',estimates);table(wb,'SO_SANH_HOP_DONG',[dict(r,description=raw[r['source_line_id']]['description'],note='Tham khảo: đơn giá BOQ nguồn ('+str(s['context'].get('contract_ref'))+') so với giá thành nội bộ; không tự kết luận lãi/lỗ') for r in result['items']]);table(wb,'HAO_PHI_DON_GIA',result['lines']);table(wb,'DINH_MUC',s['norms']);table(wb,'DINH_MUC_AP_DUNG',s['applications']);table(wb,'HAO_PHI_GOC',s['resources']);table(wb,'DON_GIA',s['prices']);table(wb,'CBS_CC1',[{'code':k,'name':CATEGORIES[k],'amount':v,'currency':s['context']['currency'],'basis':s['context']['cost_basis'],'scope':'Nguồn lực + rule COST; TAX/PROFIT/CONTINGENCY tách ở CHI_PHI_BO_SUNG; phân bổ xuống item chỉ ở SO_SANH_HOP_DONG'} for k,v in result['cbs'].items()]);table(wb,'CHI_PHI_BO_SUNG',result['rules'],['id','name','nature','cbs','method','basis','value','amount','source_ref'])
    table(wb,'TONG_HOP',[{'item':'Mục đích dự toán','value':s['context'].get('estimate_purpose')},{'item':'Chi phí theo hao phí','value':result['direct_cost']},{'item':'Tổng trước làm tròn','value':result['total_unrounded']},{'item':'Tổng sau làm tròn','value':result['total']},{'item':'Tiền tệ','value':s['context']['currency']},{'item':'Phủ hao phí','value':f"{result['coverage']['priced']}/{result['coverage']['needed']}"},{'item':'Phân bổ gián tiếp','value':result['allocation']['driver']},{'item':'Trạng thái','value':'DỰ THẢO LOCAL - CHƯA PHÁT HÀNH CC1'},{'item':'Quy tắc làm tròn','value':result['rounding']},{'item':'Snapshot SHA256','value':digest(s)},{'item':'Nguồn bố cục','value':template.name+' / 10_WP_BOQ_COMPOSITION'},{'item':'Template SHA256','value':template_hash}])
    table(wb,'WBS_OCCURRENCE',s['wbs']);table(wb,'LIEN_KET_MA',parts);table(wb,'BOQ_NGUON',s['boq']);table(wb,'HO_SO',[{'field':k,'value':v} for k,v in s['context'].items()]);table(wb,'KIEM_TRA',[{'field':k,'value':v} for k,v in s['export'].items()]);table(wb,'AUDIT',events)
    table(wb,'DOI_SOAT',[{'source_id':r['id'],'kind':r['kind'],'raw_quantity':r.get('quantity'),'uom':r.get('uom'),'allocated_weight':str(sum((number(p['weight']) for p in parts if p['source_line_id']==r['id']),Decimal(0))) if r['kind']=='MEASURED' else None,'disposition_ref':r.get('disposition_ref'),'source_ref':r.get('source_ref'),'source_hash':r.get('source_hash')} for r in s['boq']])
    table(wb,'NGUON_VA_GIOI_HAN',[{'source':'Tài liệu triển khai bộ mã v1.0 09/10/2026','meaning':'Work Result/occurrence, relations, CBS7, CA, identity, basis, traceability; nguồn PROPOSED'},{'source':'MASTER V1.20','meaning':'481 mã; trạng thái nguồn không thay phê duyệt hiện tại'},{'source':'TT36/37/38 và QĐ1538','meaning':'Áp dụng theo legal_ref của hồ sơ; không tự định lại giá hợp đồng cũ'},{'source':'Mẫu CC1 '+template.name,'meaning':'Kế thừa header/style 10_WP_BOQ_COMPOSITION; các bảng dự toán thêm là snapshot local có truy vết; không chứng nhận mẫu ban hành'},{'source':'Snapshot values','meaning':'Tiền và hệ số lưu chuỗi Decimal chính xác; không có công thức Excel cần recalculation; chỉ tổng cuối làm tròn, các bảng chi tiết giữ độ chính xác'}])
    out.parent.mkdir(parents=True,exist_ok=True);wb.save(out)
    return result
