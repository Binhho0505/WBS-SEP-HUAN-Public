"""E18 (D78): hồ sơ hợp đồng NTP trên BB′ (= BOQ CC1 ký gửi nhà thầu phụ, D74) — PRD AB→BB′ FR29–FR32, LOCAL DRAFT.

What this is: a local staging record of the subcontract dossier, NOT a contract generator and NOT a signature.
  FR29  a draft dossier per package: the clause headings PRD FR29 lists (each entered with a source), the pinned form file
        (name + SHA-256; the form itself is OPEN — OQ07) and a back-to-back checklist of AB clauses, each decided with a rationale.
        No contract text is generated: there is no pinned template (PRD §8).
  FR30  states CONTRACT_DRAFT → IN_REVIEW → READY_FOR_AUTHORIZED_SIGNING → SIGNED_EVIDENCE_RECORDED → AMENDED → CLOSED.
        READY needs every gate (server side); at READY the BB′ lines are snapshotted with a digest. Signing is recorded from
        evidence the user supplies (file hash, parties, contract ref/date, recorder, verification note); a scan does not prove
        a valid signature and authority verification stays separate (OQ02). Once signed, the package lines and award prices
        are locked: changes go through an amendment; the signed revision is never edited or deleted.
  FR31  change requests with impact on quantity/cost/time, funding source, decision maker, effective date and lineage;
        AB variation and NTP variation have separate states; an NTP cost agreed while AB has not approved is unfunded exposure.
  FR32  facts NTP_ACCEPTED (per BB′ line), AB_ACCEPTED (per AB line) and NTP_PAYMENT (per contract), each with period,
        evidence, project revision and a declared DELTA/CUMULATIVE basis (one basis per series, so nothing is counted twice),
        reversal instead of edit. NTP progress (done) stays in E14; nothing here turns NTP completion into AB acceptance or revenue.
Odoo: APPLICABLE, field-level design HOLD (R1 inventory not done); this is local staging only, no ERP write.
"""
import copy,hashlib,json,re,uuid
from datetime import datetime
from decimal import Decimal
from domain import number,index,meaningful
import subcontract as sc

STATES=('CONTRACT_DRAFT','IN_REVIEW','READY_FOR_AUTHORIZED_SIGNING','SIGNED_EVIDENCE_RECORDED','AMENDED','CLOSED')
STATE_NAMES={'CONTRACT_DRAFT':'Nháp','IN_REVIEW':'Đang xem xét','READY_FOR_AUTHORIZED_SIGNING':'Sẵn sàng trình ký có thẩm quyền',
             'SIGNED_EVIDENCE_RECORDED':'Đã ghi bằng chứng ký','AMENDED':'Đã có phụ lục','CLOSED':'Đã đóng'}
SIGNED=('SIGNED_EVIDENCE_RECORDED','AMENDED','CLOSED')
# FR29 clause headings, verbatim order of the PRD; the content of each is the user's, with a source
TERMS={'parties':'Chủ thể','scope':'Phạm vi / BOQ','price_type':'Loại giá','taxes':'Thuế','schedule':'Tiến độ / mốc',
       'payment':'Thanh toán / tạm ứng / giữ lại / bảo lãnh','acceptance':'Điều kiện nghiệm thu','variation':'Thay đổi / điều chỉnh giá',
       'cc1_supplied':'Vật tư CC1 cấp','risk':'Rủi ro','insurance':'Bảo hiểm','warranty':'Bảo hành','breach':'Vi phạm',
       'disputes':'Giải quyết tranh chấp','annexes':'Phụ lục áp dụng'}
B2B=('APPLY','ADAPT','NOT_APPLICABLE');B2B_NAMES={'APPLY':'Áp dụng nguyên','ADAPT':'Áp dụng có điều chỉnh','NOT_APPLICABLE':'Không áp dụng'}
AB_STATES=('NOT_SUBMITTED','SUBMITTED','APPROVED','REJECTED');NTP_STATES=('NOT_ISSUED','ISSUED','AGREED','REJECTED')
SOURCES=('DESIGN','AB','NTP','CC1')
FACTS={'NTP_ACCEPTED':'KL nghiệm thu với NTP','AB_ACCEPTED':'KL được CĐT (AB) nghiệm thu','NTP_PAYMENT':'Tiền đã thanh toán NTP'}
BASES=('CUMULATIVE','DELTA');BASIS_NAMES={'CUMULATIVE':'lũy kế','DELTA':'phát sinh trong kỳ'}
IMPACT_NAMES={'ntp_cost':'chi phí NTP','ab_value':'giá trị AB','time_days':'thời gian (ngày)'}
NOTE='DỰ THẢO TRÊN MÁY — hồ sơ hợp đồng NTP để xem xét; không phải hợp đồng, không phải chữ ký; mẫu hợp đồng và thẩm quyền ký ĐANG MỞ (OQ07, OQ02).'
PERIOD=re.compile(r'^\d{4}-(0[1-9]|1[0-2])$');DAY=re.compile(r'^\d{4}-\d{2}-\d{2}$')
def now():return datetime.now().isoformat(timespec='seconds')
def s_(d):return sc.s_(d)
def _txt(v,n=2000):return str(v or '').strip()[:n]
def contracts(s):return s.get('contracts') or []
def _get(s,cid):
    c=next((x for x in contracts(s) if x['id']==cid),None)
    if not c:raise ValueError('Không có hồ sơ hợp đồng này.')
    return c
def _pkg(s,pid):
    p=next((x for x in sc.packages(s) if x['id']==pid),None)
    if not p or p.get('is_self'):raise ValueError('Chọn gói thầu con (không phải gói CC1 tự làm).')
    return p
# ---------- BB′ lines ----------
def live_lines(s,p):
    """BB′ lines of a package as they would be signed now: AB keys + allocated quantity + award unit price (null if missing)."""
    ab=index(s['boq']);aw=(p.get('award') or {});out=[]
    for l in p.get('lines') or []:
        a=ab.get(l['ab_line_id']) or {};q=number(a.get('quantity'));al=number(l.get('alloc_pct'));qty=q*al/sc.HUNDRED if q is not None and al is not None else None
        up=number((aw.get('prices') or {}).get(l['bb_code']));k=sc.wbs_keys(s,l['ab_line_id'])
        out.append({'bb_code':l['bb_code'],'ab_line_id':l['ab_line_id'],'wbs_code':k['wbs_code'],'work_results':k['work_results'],'lbs':k['lbs'],
                    'description':a.get('description'),'uom':a.get('uom'),'alloc_pct':l.get('alloc_pct'),'quantity':s_(qty),'unit_price':s_(up),
                    'amount':s_(qty*up) if qty is not None and up is not None else None,'price_ref':(aw.get('price_refs') or {}).get(l['bb_code'])})
    return out
def digest(obj):return hashlib.sha256(json.dumps(obj,sort_keys=True,ensure_ascii=False,default=str).encode()).hexdigest()
def total(lines):
    vals=[number(l['amount']) for l in lines]
    return None if any(v is None for v in vals) else s_(sum(vals,Decimal(0)))
def current_lines(c):
    """Signed lines in force: the last recorded amendment, else the signed snapshot."""
    rec=[a for a in c.get('amendments') or [] if a.get('signed')]
    return rec[-1]['lines'] if rec else (c.get('snapshot') or {}).get('lines')
# ---------- gates ----------
def gates(s,c):
    """Every missing item for READY_FOR_AUTHORIZED_SIGNING (empty = may become ready). Server-side; the UI only shows them."""
    out=[];p=next((x for x in sc.packages(s) if x['id']==c['package_id']),None)
    if not p:return ['Gói thầu của hồ sơ không còn.']
    aw=p.get('award') or {}
    if not aw:out.append('Gói chưa có đề xuất giao thầu (giá giao từng dòng BB′).')
    elif not meaningful(aw.get('vendor')):out.append('Đề xuất giao thầu chưa ghi nhà thầu phụ.')
    elif sc.award_stale(s,p):out.append('Đề xuất giao thầu đã cũ so với dòng gói / AB / báo giá — lưu lại đề xuất (FR28).')
    lines=live_lines(s,p)
    if not lines:out.append('Gói chưa có dòng BB′.')
    for l in lines:
        if l['unit_price'] is None:out.append(f'Dòng {l["bb_code"]}: chưa có đơn giá giao.')
        if l['quantity'] is None:out.append(f'Dòng {l["bb_code"]}: chưa có khối lượng.')
    cov=sc.coverage(s)
    for l in p.get('lines') or []:
        st=(cov.get(l['ab_line_id']) or {}).get('state')
        if st in ('OVER','INVALID'):out.append(f'Dòng {l["bb_code"]}: phân bổ dòng AB '+('vượt 100%' if st=='OVER' else 'không hợp lệ')+' (AC03).')
    stale={x.get('bb_code') for x in sc.stale_lines(s)}
    for l in p.get('lines') or []:
        if l['bb_code'] in stale:out.append(f'Dòng {l["bb_code"]}: dòng AB đã đổi sau khi giao — xem lại gói (BR08).')
    for k,n in TERMS.items():
        t=(c.get('terms') or {}).get(k) or {}
        if not meaningful(t.get('text')) or not meaningful(t.get('ref')):out.append(f'Điều khoản "{n}": nhập nội dung và căn cứ (ghi rõ nếu không áp dụng).')
    f=c.get('form') or {}
    if not f.get('sha256'):out.append('Chưa chốt mẫu hợp đồng (tải file mẫu được dùng; mẫu chuẩn là câu hỏi mở OQ07).')
    b=c.get('back_to_back') or []
    if not b:out.append('Bảng đối chiếu tương ứng với HĐ AB (back-to-back) trống: liệt kê điều khoản AB đã xem xét.')
    for i,x in enumerate(b,1):
        if x.get('decision') not in B2B or not meaningful(x.get('rationale')):out.append(f'Đối chiếu tương ứng dòng {i}: chọn áp dụng / điều chỉnh / không áp dụng và ghi lý do.')
    a=c.get('authority') or {}
    if not meaningful(a.get('name')) or not meaningful(a.get('ref')):out.append('Ghi người có thẩm quyền ký và căn cứ ủy quyền (app không kiểm thẩm quyền — OQ02).')
    return out
def ready_stale(s,c):
    """True when the BB′ lines changed after READY: signing must go back to review."""
    snap=c.get('snapshot') or {}
    if c['status']!='READY_FOR_AUTHORIZED_SIGNING' or not snap:return False
    p=next((x for x in sc.packages(s) if x['id']==c['package_id']),None)
    return not p or digest(live_lines(s,p))!=snap.get('digest') or (p.get('award') or {}).get('vendor')!=snap.get('vendor')
def drift(s,c):
    """Signed lines whose AB quantity changed since signing (AC17): proposals for a change request; the signed lines stay."""
    lines=current_lines(c)
    if c['status'] not in SIGNED or not lines:return []
    ab=index(s['boq']);out=[]
    for l in lines:
        a=ab.get(l['ab_line_id']);q=number((a or {}).get('quantity'));al=number(l.get('alloc_pct'))
        now_q=q*al/sc.HUNDRED if q is not None and al is not None else None;was=number(l['quantity'])
        if a is None:out.append({'bb_code':l['bb_code'],'signed_qty':l['quantity'],'current_qty':None,'delta':None,'text':'Dòng AB không còn'})
        elif now_q!=was:out.append({'bb_code':l['bb_code'],'signed_qty':l['quantity'],'current_qty':s_(now_q),'delta':s_(now_q-was) if None not in (now_q,was) else None,'text':'KL AB đổi sau khi ký'})
    return out
# ---------- lock (called before saving packages / award) ----------
def guard(old,new):
    """Refuse changing lines or award prices of a package whose contract is signed (FR30: change through an amendment)."""
    for c in contracts(old):
        if c['status'] not in SIGNED:continue
        po=next((x for x in sc.packages(old) if x['id']==c['package_id']),None);pn=next((x for x in sc.packages(new) if x['id']==c['package_id']),None)
        if po is None:continue
        if pn is None:raise ValueError(f'Gói {po["code"]} có hợp đồng đã ghi bằng chứng ký: không xóa được gói.')
        key=lambda p:sorted((l['bb_code'],l['ab_line_id'],str(l.get('alloc_pct'))) for l in p.get('lines') or [])
        if key(po)!=key(pn):raise ValueError(f'Gói {po["code"]} có hợp đồng đã ký: không sửa/thêm/bớt dòng BB′ — lập yêu cầu thay đổi và phụ lục.')
        if ((po.get('award') or {}).get('prices'),(po.get('award') or {}).get('vendor'))!=((pn.get('award') or {}).get('prices'),(pn.get('award') or {}).get('vendor')):
            raise ValueError(f'Gói {po["code"]} có hợp đồng đã ký: không sửa giá giao / nhà thầu — lập yêu cầu thay đổi và phụ lục.')
    progress_refs(new)
def progress_refs(s):
    """No orphan facts: a BB′ line with recorded progress/acceptance cannot disappear."""
    codes={l['bb_code'] for p in sc.packages(s) for l in p.get('lines') or []}
    for e in (s.get('progress') or [])+[f for f in s.get('facts') or [] if f.get('bb_code')]:
        if e['bb_code'] not in codes:raise ValueError(f'Dòng {e["bb_code"]} đã có khối lượng hoàn thành/nghiệm thu: không xóa dòng (sẽ mồ côi dữ liệu).')
# ---------- actions ----------
def _hist(c,frm,to,by,note=''):c.setdefault('history',[]).append({'at':now(),'by':_txt(by,200),'from':frm,'to':to,'note':_txt(note,1000)})
def act(s,body,file=None):
    """Mutations: create | save | submit | return | ready | sign | amend_draft | amend_record | close | change | fact.
    `file` = (name, bytes, sha256) for form / signed evidence uploads. Returns the changed project."""
    a=body.get('action');s=copy.deepcopy(s);by=_txt(body.get('by'),200)
    if a=='create':
        p=_pkg(s,body.get('package_id'))
        if any(c['package_id']==p['id'] and c['status']!='CLOSED' for c in contracts(s)):raise ValueError(f'Gói {p["code"]} đã có hồ sơ hợp đồng đang mở.')
        c={'id':str(uuid.uuid4()),'package_id':p['id'],'package_code':p['code'],'status':'CONTRACT_DRAFT','created':now(),'created_by':by,
           'ref':_txt(body.get('ref'),120),'terms':{},'form':{},'back_to_back':[],'authority':{},'snapshot':None,'signed':None,'amendments':[],'history':[]}
        _hist(c,None,'CONTRACT_DRAFT',by,'Tạo hồ sơ');s.setdefault('contracts',[]).append(c);return s
    if a=='change':return _change(s,body)
    if a=='fact':return _fact(s,body)
    c=_get(s,body.get('contract_id'));st=c['status']
    if a=='save':
        if st!='CONTRACT_DRAFT':raise ValueError('Chỉ sửa hồ sơ ở trạng thái Nháp (trả về Nháp để sửa).')
        if 'ref' in body:c['ref']=_txt(body.get('ref'),120)
        if 'terms' in body:
            t=body.get('terms') or {}
            if not isinstance(t,dict) or any(k not in TERMS for k in t):raise ValueError('Điều khoản không thuộc danh mục FR29.')
            c['terms']={k:{'text':_txt(v.get('text'),4000),'ref':_txt(v.get('ref'),1000)} for k,v in t.items() if isinstance(v,dict)}
        if 'back_to_back' in body:
            b=body.get('back_to_back') or []
            if not isinstance(b,list) or len(b)>300:raise ValueError('Bảng đối chiếu tương ứng (back-to-back) là danh sách (tối đa 300).')
            for x in b:
                if not isinstance(x,dict) or not meaningful(x.get('clause')):raise ValueError('Mỗi dòng đối chiếu tương ứng cần điều khoản AB.')
                if x.get('decision') not in (None,'')+B2B:raise ValueError('Quyết định đối chiếu tương ứng phải là Áp dụng nguyên (APPLY), Áp dụng có điều chỉnh (ADAPT) hoặc Không áp dụng (NOT_APPLICABLE).')
            c['back_to_back']=[{'id':x.get('id') or str(uuid.uuid4()),'clause':_txt(x.get('clause'),500),'ab_ref':_txt(x.get('ab_ref'),300),'decision':x.get('decision') or None,'rationale':_txt(x.get('rationale'),2000)} for x in b]
        if 'authority' in body:
            x=body.get('authority') or {};c['authority']={'name':_txt(x.get('name'),200),'role':_txt(x.get('role'),200),'ref':_txt(x.get('ref'),500)}
        if file:c['form']={'name':file[0][:200],'sha256':file[2],'bytes':len(file[1]),'ref':_txt(body.get('form_ref'),500),'at':now(),'by':by}
        return s
    moves={'submit':('CONTRACT_DRAFT','IN_REVIEW'),'return':(('IN_REVIEW','READY_FOR_AUTHORIZED_SIGNING'),'CONTRACT_DRAFT'),'ready':('IN_REVIEW','READY_FOR_AUTHORIZED_SIGNING')}
    if a in moves:
        frm,to=moves[a]
        if st not in ((frm,) if isinstance(frm,str) else frm):raise ValueError(f'Không chuyển được từ "{STATE_NAMES[st]}" sang "{STATE_NAMES[to]}".')
        if not by:raise ValueError('Ghi người thực hiện.')
        if a=='ready':
            g=gates(s,c)
            if g:raise ValueError(f'Chưa đủ điều kiện trình ký ({len(g)} mục): '+g[0])
            p=_pkg(s,c['package_id']);lines=live_lines(s,p)
            c['snapshot']={'lines':lines,'digest':digest(lines),'total':total(lines),'vendor':(p.get('award') or {}).get('vendor'),'project_revision':s['revision'],
                           'terms_digest':digest(c['terms']),'form_sha256':c['form'].get('sha256'),'at':now()}
        if a=='return':c['snapshot']=None
        c['status']=to;_hist(c,st,to,by,body.get('note'));return s
    if a=='sign':
        if st!='READY_FOR_AUTHORIZED_SIGNING':raise ValueError('Chỉ ghi bằng chứng ký khi hồ sơ ở trạng thái Sẵn sàng trình ký.')
        if ready_stale(s,c):raise ValueError('Dòng BB′ đã đổi sau khi trình ký — trả về Nháp và xem lại.')
        ev=_evidence(body,file);c['signed']=ev;c['status']='SIGNED_EVIDENCE_RECORDED';c['ref']=ev['contract_ref']
        _hist(c,st,c['status'],by,'Ghi bằng chứng ký (không xác minh chữ ký/thẩm quyền)');return s
    if a=='amend_draft':
        if st not in ('SIGNED_EVIDENCE_RECORDED','AMENDED'):raise ValueError('Phụ lục chỉ lập cho hợp đồng đã ghi bằng chứng ký.')
        if any(not x.get('signed') for x in c['amendments']):raise ValueError('Đang có một phụ lục nháp chưa ghi bằng chứng ký.')
        ids=[str(x) for x in body.get('change_ids') or []];chs={x['id']:x for x in s.get('changes') or []}
        if not ids or any(i not in chs or chs[i].get('contract_id')!=c['id'] for i in ids):raise ValueError('Phụ lục phải dẫn chiếu yêu cầu thay đổi của hợp đồng này.')
        lines=_amend_lines(current_lines(c),body.get('lines'))
        c['amendments'].append({'id':str(uuid.uuid4()),'no':len(c['amendments'])+1,'reason':_txt(body.get('reason'),2000),'change_ids':ids,'lines':lines,
                                'digest':digest(lines),'total':total(lines),'created':now(),'by':by,'signed':None})
        if not meaningful(c['amendments'][-1]['reason']):raise ValueError('Ghi lý do phụ lục.')
        return s
    if a=='amend_record':
        am=next((x for x in c['amendments'] if not x.get('signed')),None)
        if not am:raise ValueError('Không có phụ lục nháp.')
        am['signed']=_evidence(body,file);c['status']='AMENDED';_hist(c,st,'AMENDED',by,f'Ghi bằng chứng ký phụ lục số {am["no"]}');return s
    if a=='close':
        if st not in ('SIGNED_EVIDENCE_RECORDED','AMENDED'):raise ValueError('Chỉ đóng hợp đồng đã ghi bằng chứng ký.')
        if not meaningful(body.get('note')):raise ValueError('Ghi căn cứ đóng hồ sơ (thanh lý / quyết toán).')
        c['status']='CLOSED';_hist(c,st,'CLOSED',by,body.get('note'));return s
    raise ValueError('Thao tác không hợp lệ.')
def _evidence(body,file):
    if not file:raise ValueError('Tải file bằng chứng ký (bản quét / file điện tử).')
    parties=[_txt(x,200) for x in body.get('parties') or [] if meaningful(x)]
    if len(parties)<2:raise ValueError('Ghi các bên ký (ít nhất hai bên).')
    ref=_txt(body.get('contract_ref'),120);d=_txt(body.get('date'),10)
    if not ref:raise ValueError('Ghi số hợp đồng / phụ lục.')
    if not DAY.match(d):raise ValueError('Ngày ký dạng YYYY-MM-DD.')
    if not meaningful(body.get('by')):raise ValueError('Ghi người ghi nhận bằng chứng.')
    if not meaningful(body.get('verification')):raise ValueError('Ghi cách đã kiểm bằng chứng (vd đối chiếu bản gốc) — app không xác minh chữ ký.')
    return {'file':file[0][:200],'sha256':file[2],'bytes':len(file[1]),'parties':parties,'contract_ref':ref,'date':d,'recorded_by':_txt(body.get('by'),200),
            'verification':_txt(body.get('verification'),1000),'at':now(),'authority_verified':False}
def _amend_lines(base,changes):
    """New line set of an amendment: start from the lines in force; change quantity / unit price per bb_code (with a ref)."""
    if not isinstance(changes,list) or not changes:raise ValueError('Phụ lục cần ít nhất một dòng thay đổi (bb_code, khối lượng và/hoặc đơn giá, căn cứ).')
    out={l['bb_code']:dict(l) for l in base}
    for x in changes:
        code=x.get('bb_code')
        if code not in out:raise ValueError(f'Dòng {code} không thuộc hợp đồng (thêm dòng mới cần OQ03 — chưa hỗ trợ).')
        if not meaningful(x.get('ref')):raise ValueError(f'Dòng {code}: ghi căn cứ thay đổi.')
        l=out[code]
        for k in ('quantity','unit_price'):
            if x.get(k) not in (None,''):
                v=number(x.get(k))
                if v is None or v<0:raise ValueError(f'Dòng {code}: {k} không hợp lệ.')
                l[k]=s_(v)
        q,u=number(l['quantity']),number(l['unit_price']);l['amount']=s_(q*u) if None not in (q,u) else None;l['change_ref']=_txt(x.get('ref'),500)
    return list(out.values())
# ---------- FR31 change requests ----------
def _change(s,body):
    ch=s.setdefault('changes',[]);cid=body.get('change_id');by=_txt(body.get('by'),200)
    if cid:
        x=next((c for c in ch if c['id']==cid),None)
        if not x:raise ValueError('Không có yêu cầu thay đổi này.')
        for k,states in (('ab_status',AB_STATES),('ntp_status',NTP_STATES)):
            if k in body:
                if body[k] not in states:raise ValueError(('Trạng thái phía AB' if k=='ab_status' else 'Trạng thái phía NTP')+' không hợp lệ.')
                if body[k]!=x[k]:x.setdefault('history',[]).append({'at':now(),'by':by,'field':k,'from':x[k],'to':body[k],'ref':_txt(body.get('ref'),500)});x[k]=body[k]
        if 'decided_by' in body:x['decided_by']=_txt(body.get('decided_by'),200)
        if 'effective' in body:
            e=_txt(body.get('effective'),10)
            if e and not DAY.match(e):raise ValueError('Ngày hiệu lực dạng YYYY-MM-DD.')
            x['effective']=e or None
        return s
    if body.get('source') not in SOURCES:raise ValueError('Nguồn thay đổi: Thiết kế (DESIGN), Chủ đầu tư (AB), NTP hoặc CC1.')
    title=_txt(body.get('title'),300)
    if not title or not meaningful(body.get('request_ref')):raise ValueError('Ghi nội dung và căn cứ yêu cầu thay đổi.')
    c=_get(s,body.get('contract_id')) if body.get('contract_id') else None
    imp=body.get('impact') or {};out={}
    for k in ('ntp_cost','ab_value','time_days'):
        v=imp.get(k);out[k]=None if v in (None,'') else s_(number(v)) if number(v) is not None else None
        if v not in (None,'') and number(v) is None:raise ValueError(f'Ảnh hưởng {IMPACT_NAMES.get(k,k)} không phải số.')
    qty=[]
    for q in imp.get('quantities') or []:
        if number(q.get('delta')) is None:raise ValueError('Ảnh hưởng khối lượng: mức thay đổi phải là số.')
        qty.append({'bb_code':_txt(q.get('bb_code'),60),'delta':s_(number(q['delta']))})
    out['quantities']=qty
    ch.append({'id':str(uuid.uuid4()),'no':len(ch)+1,'source':body['source'],'title':title,'request_ref':_txt(body.get('request_ref'),500),
               'contract_id':c['id'] if c else None,'package_id':c['package_id'] if c else body.get('package_id'),'impact':out,'funding':_txt(body.get('funding'),500),
               'ab_status':'NOT_SUBMITTED','ntp_status':'NOT_ISSUED','decided_by':None,'effective':None,'created':now(),'by':by,
               'lineage':{'project_revision':s['revision'],'contract_digest':((c or {}).get('snapshot') or {}).get('digest')},'history':[]})
    return s
def exposure(s):
    """FR31: NTP-side cost agreed/issued while the AB side is not approved = unfunded (no right to bill yet). Unknown stays unknown."""
    rows=[];tot=Decimal(0);unknown=0
    for x in s.get('changes') or []:
        if x['ntp_status'] in ('ISSUED','AGREED') and x['ab_status']!='APPROVED':
            v=number(x['impact'].get('ntp_cost'))
            if v is None:unknown+=1
            else:tot+=v
            rows.append({'no':x['no'],'title':x['title'],'ntp_cost':x['impact'].get('ntp_cost'),'ab_status':x['ab_status'],'ntp_status':x['ntp_status']})
    return {'rows':rows,'unfunded':s_(tot),'unknown':unknown,'note':'AB chưa chấp thuận: được ghi rủi ro chi phí nhưng chưa có quyền thu.'}
# ---------- FR32 facts ----------
def _fact(s,body):
    k=body.get('kind');by=_txt(body.get('by'),200);facts=s.setdefault('facts',[])
    if body.get('reverses'):
        x=next((f for f in facts if f['id']==body['reverses']),None)
        if not x or x.get('reversed_by'):raise ValueError('Không có bản ghi để đảo hoặc đã đảo.')
        if not meaningful(body.get('evidence')):raise ValueError('Ghi căn cứ đảo bút toán.')
        r={'id':str(uuid.uuid4()),'reversal_of':x['id'],'kind':x['kind'],'at':now(),'by':by,'evidence':_txt(body.get('evidence'),500),'revision':s['revision']}
        x['reversed_by']=r['id'];facts.append(r);return s
    if k not in FACTS:raise ValueError('Loại: KL nghiệm thu với NTP (NTP_ACCEPTED), KL được CĐT (AB) nghiệm thu (AB_ACCEPTED) hoặc Tiền đã thanh toán NTP (NTP_PAYMENT).')
    per=_txt(body.get('period'),7)
    if not PERIOD.match(per):raise ValueError('Kỳ dạng YYYY-MM.')
    if body.get('basis') not in BASES:raise ValueError('Khai báo Lũy kế (CUMULATIVE) hoặc Phát sinh trong kỳ (DELTA).')
    if not meaningful(body.get('evidence')):raise ValueError('Ghi căn cứ (biên bản nghiệm thu / chứng từ thanh toán).')
    v=number(body.get('value'))
    if v is None or v<0:raise ValueError('Giá trị phải là số ≥ 0 (trống không được coi là 0).')
    key={}
    if k=='NTP_ACCEPTED':
        if body.get('bb_code') not in {l['bb_code'] for p in sc.packages(s) for l in p.get('lines') or []}:raise ValueError('Chọn dòng BB′.')
        key={'bb_code':body['bb_code']}
    elif k=='AB_ACCEPTED':
        if body.get('ab_line_id') not in sc.measured(s):raise ValueError('Chọn dòng BOQ AB.')
        key={'ab_line_id':body['ab_line_id']}
    else:key={'contract_id':_get(s,body.get('contract_id'))['id']}
    same=[f for f in facts if f.get('kind')==k and not f.get('reversal_of') and not f.get('reversed_by') and all(f.get(a)==b for a,b in key.items())]
    if same and same[0]['basis']!=body['basis']:raise ValueError(f'Chuỗi này đã khai báo {BASIS_NAMES.get(same[0]["basis"],same[0]["basis"])}; không trộn lũy kế và phát sinh (tránh cộng hai lần).')
    if body['basis']=='CUMULATIVE':
        if any(f['period']==per for f in same):raise ValueError('Kỳ này đã có số lũy kế — đảo bản ghi cũ rồi nhập lại.')
        prev=[f for f in same if f['period']<per]
        if prev and v<number(max(prev,key=lambda f:f['period'])['value']):raise ValueError('Lũy kế nhỏ hơn kỳ trước — kiểm lại hoặc đảo bản ghi sai.')
        if any(f['period']>per for f in same):raise ValueError('Đã có số lũy kế của kỳ sau; nhập theo thứ tự kỳ.')
    facts.append({'id':str(uuid.uuid4()),'kind':k,**key,'period':per,'basis':body['basis'],'value':s_(v),'evidence':_txt(body.get('evidence'),500),'by':by,'at':now(),'revision':s['revision']})
    return s
def fact_totals(s):
    """Latest value per series: CUMULATIVE → last period, DELTA → sum. Reversed facts are excluded. No series → null (unknown, not 0)."""
    series={}
    for f in s.get('facts') or []:
        if f.get('reversal_of') or f.get('reversed_by'):continue
        k=(f['kind'],f.get('bb_code') or f.get('ab_line_id') or f.get('contract_id'));series.setdefault(k,[]).append(f)
    out={}
    for k,fs in series.items():
        out[k]=s_(number(max(fs,key=lambda f:f['period'])['value'])) if fs[0]['basis']=='CUMULATIVE' else s_(sum((number(f['value']) for f in fs),Decimal(0)))
    return out
def trace(s):
    """FR32 per BB′ line: NTP done (E14), NTP accepted; per AB line: AB accepted — kept apart, never derived from each other."""
    tot=fact_totals(s);latest={}
    for e in sorted(s.get('progress') or [],key=lambda e:(e['date'],e['id'])):latest[e['bb_code']]=e['cum_qty']
    lines=[];warn=[]
    for p in sc.packages(s):
        for l in p.get('lines') or []:
            done=latest.get(l['bb_code']);acc=tot.get(('NTP_ACCEPTED',l['bb_code']))
            if acc is not None and done is not None and number(acc)>number(done):warn.append(f'{l["bb_code"]}: KL nghiệm thu NTP lớn hơn KL thực hiện.')
            lines.append({'package':p['code'],'bb_code':l['bb_code'],'ab_line_id':l['ab_line_id'],'ntp_done':done,'ntp_accepted':acc})
    ab=[{'ab_line_id':i,'source_code':a.get('source_code'),'quantity':a.get('quantity'),'ab_accepted':tot.get(('AB_ACCEPTED',i))} for i,a in sc.measured(s).items()]
    pay=[{'contract_id':c['id'],'package':c['package_code'],'paid':tot.get(('NTP_PAYMENT',c['id'])),'signed_total':total(current_lines(c)) if current_lines(c) else None} for c in contracts(s)]
    return {'lines':lines,'ab':ab,'payments':pay,'warnings':warn}
# ---------- view, restore, Excel ----------
def view(s):
    out=[]
    for c in contracts(s):
        p=next((x for x in sc.packages(s) if x['id']==c['package_id']),None)
        out.append(dict(c,state_name=STATE_NAMES[c['status']],gates=gates(s,c) if c['status'] in ('CONTRACT_DRAFT','IN_REVIEW') else [],ready_stale=ready_stale(s,c),
                        drift=drift(s,c),live_lines=live_lines(s,p) if p and c['status'] not in SIGNED else None,lines_in_force=current_lines(c) if c['status'] in SIGNED else None))
    return {'contracts':out,'changes':s.get('changes') or [],'exposure':exposure(s),'trace':trace(s),'facts':s.get('facts') or [],
            'terms':TERMS,'states':STATE_NAMES,'b2b':B2B_NAMES,'ab_states':AB_STATES,'ntp_states':NTP_STATES,'sources':SOURCES,'fact_kinds':FACTS,'note':NOTE}
def restore(old):
    """Restored dossiers are evidence, not authority: READY goes back to review (readiness is a review decision); signed
    records keep their evidence, marked restored."""
    cs=[]
    for c in old.get('contracts') or []:
        c=copy.deepcopy(c);c['restored']=True
        if c['status']=='READY_FOR_AUTHORIZED_SIGNING':c['status']='IN_REVIEW';c['snapshot']=None;_hist(c,'READY_FOR_AUTHORIZED_SIGNING','IN_REVIEW','RESTORE','Khôi phục: xem lại trước khi trình ký')
        cs.append(c)
    return cs
def workbook(s,cid):
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    c=_get(s,cid);wb=Workbook();hdr=PatternFill('solid',fgColor='106B3C');white=Font(bold=True,color='FFFFFF')
    def sheet(ws,name,head,rows,widths):
        ws.title=name;ws.append([NOTE]);ws.append([f'Gói {c["package_code"]} · {STATE_NAMES[c["status"]]} · số HĐ {c.get("ref") or "—"}']);ws.append(head)
        for x in ws[3]:x.fill=hdr;x.font=white;x.alignment=Alignment(wrap_text=True)
        for r in rows:ws.append(r)
        for row in ws.iter_rows(min_row=4):
            for x in row:
                if isinstance(x.value,str) and x.value.startswith(('=','+','-','@')):x.data_type='s'
        for i,w in enumerate(widths):ws.column_dimensions[chr(65+i)].width=w
        ws.freeze_panes='A4'
    num=lambda v:float(v) if v not in (None,'') else None
    lines=current_lines(c) if c['status'] in SIGNED else live_lines(s,_pkg(s,c['package_id']))
    sheet(wb.active,'BOQ_BB_PRIME',['Mã dòng BB′','Dòng AB','Mã WBS (ERP)','Work Result','LBS','Mô tả','ĐVT','% phân bổ','Khối lượng','Đơn giá giao','Thành tiền','Căn cứ giá'],
          [[l['bb_code'],l['ab_line_id'],l['wbs_code'],', '.join(l.get('work_results') or []),', '.join(l.get('lbs') or []),l['description'],l['uom'],num(l['alloc_pct']),num(l['quantity']),num(l['unit_price']),num(l['amount']),l.get('price_ref') or l.get('change_ref')] for l in lines],
          [12,14,14,16,16,40,8,10,12,14,16,30])
    sheet(wb.create_sheet(),'DIEU_KHOAN',['Mục (FR29)','Nội dung','Căn cứ'],[[n,((c.get('terms') or {}).get(k) or {}).get('text'),((c.get('terms') or {}).get(k) or {}).get('ref')] for k,n in TERMS.items()],[30,80,40])
    sheet(wb.create_sheet(),'BACK_TO_BACK_AB',['Điều khoản AB','Vị trí AB','Quyết định','Lý do'],[[x['clause'],x['ab_ref'],B2B_NAMES.get(x['decision'],'chưa chọn'),x['rationale']] for x in c.get('back_to_back') or []],[40,20,22,60])
    ev=[['Form',(c.get('form') or {}).get('name'),(c.get('form') or {}).get('sha256'),(c.get('form') or {}).get('ref')]]
    for nm,e in [('Bằng chứng ký',c.get('signed'))]+[(f'Phụ lục {a["no"]}',a.get('signed')) for a in c.get('amendments') or []]:
        if e:ev.append([nm,e['file'],e['sha256'],f'{e["contract_ref"]} · {e["date"]} · {", ".join(e["parties"])} · ghi bởi {e["recorded_by"]} · kiểm: {e["verification"]} · thẩm quyền chưa xác minh'])
    sheet(wb.create_sheet(),'BANG_CHUNG',['Loại','File','SHA-256','Chi tiết'],ev,[18,30,66,80])
    sheet(wb.create_sheet(),'LICH_SU',['Lúc','Người','Từ','Sang','Ghi chú'],[[h['at'],h['by'],STATE_NAMES.get(h['from'],h['from']),STATE_NAMES.get(h['to'],h['to']),h['note']] for h in c.get('history') or []],[20,20,26,26,60])
    b=io.BytesIO();wb.save(b);return b.getvalue()
