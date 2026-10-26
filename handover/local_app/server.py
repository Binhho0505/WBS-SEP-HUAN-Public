"""Loopback-only local application. No ERP writes, cloud release or implied approval."""
import argparse,base64,csv,hashlib,io,json,mimetypes,secrets,sys,uuid,zipfile,threading,os
from http.server import BaseHTTPRequestHandler,ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse,parse_qs
from openpyxl import load_workbook
from domain import ROOT,STEPS,KEYS,SCHEMA,CATEGORIES,PURPOSES,PROFILES,status,calculate,catalog,digest,upgrade,needed_price_keys
from store import Store,Conflict,LIBRARY_KINDS
from exporter import workbook
import ai,codelib,extensions,qto_bridge
CODE_LIBRARY_XLSX=ROOT/'runs_brian/norm_library_20261017/export/THU_VIEN_MA_CC1_TT38.xlsx'
HERE=Path(__file__).resolve().parent
MAX=35*1024*1024
def imported(data):
    name=Path(data.get('filename','BOQ.csv')).name
    binary=base64.b64decode(data['data'],validate=True)
    if len(binary)>25*1024*1024:raise ValueError('File lớn hơn 25 MB. Chọn phần dữ liệu của một gói thầu/cơ sở trước khi nạp.')
    hash=hashlib.sha256(binary).hexdigest();suffix=Path(name).suffix.lower(); sheet=data.get('sheet'); header=int(data.get('header_row',1))
    if not 1<=header<=100:raise ValueError('Dòng tiêu đề phải từ 1 đến 100.')
    if suffix=='.csv':
        decoded=binary.decode('utf-8-sig');sample=decoded[:4096]
        try:dialect=csv.Sniffer().sniff(sample,delimiters=',;\t')
        except csv.Error:dialect=csv.excel
        rows=list(csv.reader(io.StringIO(decoded),dialect));sheets=['CSV'];sheet='CSV';formulas=None
    elif suffix=='.xlsx':
        with zipfile.ZipFile(io.BytesIO(binary)) as z:
            if sum(x.file_size for x in z.infolist())>150*1024*1024:raise ValueError('File Excel vượt dung lượng giải nén cho phép.')
        wb=load_workbook(io.BytesIO(binary),read_only=True,data_only=True);sheets=wb.sheetnames
        if sheet not in sheets:sheet=sheets[0]
        rows=[]
        for r in wb[sheet].iter_rows(values_only=True):
            rows.append([v.isoformat() if hasattr(v,'isoformat') else v for v in r])
            if len(rows)>20000:wb.close();raise ValueError('Chọn file con tối đa 20.000 dòng cho một hồ sơ trên máy.')
        wb.close()
        # D79 FR02: keep the source formula next to the cached value (the value used stays the cached one; formulas never run)
        wf=load_workbook(io.BytesIO(binary),read_only=True,data_only=False);formulas=[]
        for r in wf[sheet].iter_rows(values_only=True):
            formulas.append([v if isinstance(v,str) and v.startswith('=') else None for v in r])
            if len(formulas)>len(rows):break
        wf.close()
    else:raise ValueError('Chỉ nhận CSV UTF-8 và XLSX; không chạy macro.')
    if len(rows)>20000:raise ValueError('Tối đa 20.000 dòng; chia theo gói/cơ sở.')
    if len(rows)<header:raise ValueError('Không có dòng tiêu đề đã chọn.')
    headers=[str(v or f'COLUMN_{i+1}') for i,v in enumerate(rows[header-1])]
    if len(set(headers))!=len(headers):raise ValueError('Dòng tiêu đề có tên cột trùng; chọn đúng dòng tiêu đề hoặc chuẩn hóa tên cột.')
    result={'filename':name,'sha256':hash,'sheets':sheets,'sheet':sheet,'headers':headers,'preview':rows[header:header+8],'count':len(rows)-header,'rows':[],'warnings':[]}
    columns=data.get('columns')
    if data.get('table') in ['wbs','mapping','applications','norms','resources','prices','reference']:
        records=[]
        for values in rows[header:]:
            if not any(v is not None and str(v).strip() for v in values):continue
            record={k:(str(v) if isinstance(v,(int,float)) and not isinstance(v,bool) else v) for k,v in zip(headers,values)}
            if not record.get('id'):record['id']=str(uuid.uuid4())
            for key in ['table_verified']:
                if key in record:record[key]=str(record[key]).lower() in ['true','1','yes']
            records.append(record)
        result['records']=records
    if columns:
        for k in ['description','quantity','uom']:
            if columns.get(k) not in headers:raise ValueError('Chọn cột '+k+' trước khi nạp.')
        for n,values in enumerate(rows[header:],header+1):
            if not any(v is not None and str(v).strip() for v in values):continue
            raw=dict(zip(headers,values));get=lambda k:raw.get(columns.get(k))
            display=str(get('kind') or '').lower();kind='HEADING' if display in ['section','line_section','heading'] else 'NOTE' if display in ['note','line_note'] else 'MEASURED'
            result['rows'].append({'id':str(uuid.uuid5(uuid.NAMESPACE_URL,hash+':'+sheet+':'+str(n))),'source_code':str(get('source_code') or ''),'description':str(get('description') or ''),'quantity':get('quantity'),'uom':get('uom'),'original_quantity':get('quantity'),'original_uom':get('uom'),'unit_price':get('unit_price'),'kind':kind,'source_ref':name+' / '+sheet+' / row '+str(n),'source_hash':hash,'raw':raw})
            r=result['rows'][-1];r['original_source_code']=r['source_code'];r['original_description']=r['description'];r['original_unit_price']=r.get('unit_price')  # D73 lineage for manual edits
            fr=formulas[n-1] if formulas and n-1<len(formulas) else None
            if fr and any(fr):
                r['raw_formulas']={h:f for h,f in zip(headers,fr) if f}
                for col in ('quantity','unit_price'):
                    f=r['raw_formulas'].get(columns.get(col) or '')
                    if f and any(k in f.upper() for k in ('SUM(','SUBTOTAL(')):result['warnings'].append(f'Dòng {n} ({r["source_code"]}): cột {col} là công thức {f[:60]} — có thể là dòng tổng, kiểm tra để không cộng hai lần.')
        seen={}
        for r in result['rows']:
            if r['kind']=='MEASURED' and r['source_code']:seen.setdefault(r['source_code'],[]).append(r['source_ref'].rsplit('row ',1)[-1])
        for code,at in seen.items():
            if len(at)>1:result['warnings'].append(f'Mã {code} xuất hiện {len(at)} dòng (dòng {", ".join(at[:5])}) — kiểm tra trùng.')
        result['warnings']=result['warnings'][:200]
    return result,binary
def builtin(source):
    if source not in ['AB','BB','PAKD_GT','PAKD_TT']:raise ValueError('Chọn một cơ sở AB, BB, PAKD_GT hoặc PAKD_TT.')
    rows=[]
    for line in (ROOT/'runs_brian/meso_real_01'/f'{source}_source_rows.jsonl').read_text(encoding='utf-8').splitlines():
        s=json.loads(line);raw=s['raw_cells']
        def find(ending):
            return next((v for k,v in raw.items() if k.endswith(ending)),None)
        kind=str(find('/Display Type') or '').lower();kind='HEADING' if kind in ['section','line_section'] else 'NOTE' if kind in ['note','line_note'] else 'MEASURED'
        rows.append({'id':str(uuid.uuid5(uuid.NAMESPACE_URL,source+s['source_sha256']+str(s['row']))),'source_code':str(find('/WBS Code') or ''),'description':str(find('/Description') or ''),'quantity':find('/Quantity'),'uom':find('/Origin UOM Name') or find('/Unit of Measure'),'unit_price':find('/Unit Price'),'kind':kind,'source_ref':source+' / '+s['sheet']+' / row '+str(s['row']),'source_hash':s['source_sha256'],'inference':'SOURCE_GROUP_UNVERIFIED' if s.get('group_inference_unverified') else 'SOURCE_REPORTED','raw':raw})
    for row in rows:
        row['original_quantity']=row['quantity'];row['original_uom']=row['uom'];row['original_source_code']=row['source_code'];row['original_description']=row['description'];row['original_unit_price']=row.get('unit_price')
    return rows

BOQ_TRACKED=('source_code','description','unit_price','quantity','uom')
def sync_mapping(s,old_boq,data):
    """D73: keep the WBS mapping and norm applications of BOQ lines that still exist (before: any added/removed line wiped them all).
    Removed or no-longer-measured lines lose their items and applications; new measured lines get an empty item; a line whose
    quantity/UOM changed updates its single unsplit item when that item still carried the old values."""
    meas={r['id']:r for r in data if r.get('kind')=='MEASURED'};old={r['id']:r for r in old_boq}
    keep=[m for m in s['mapping'] if m.get('source_line_id') in meas];gone={m.get('id') for m in s['mapping']}-{m.get('id') for m in keep}
    s['applications']=[a for a in s['applications'] if a.get('item_id') not in gone]
    by={}
    for m in keep:by.setdefault(m['source_line_id'],[]).append(m)
    for sid,rows in by.items():
        o=old.get(sid);n=meas[sid]
        if o and len(rows)==1 and str(rows[0].get('quantity'))==str(o.get('quantity')) and rows[0].get('uom')==o.get('uom'):rows[0]['quantity']=n.get('quantity');rows[0]['uom']=n.get('uom')
    for r in data:
        if r.get('kind')=='MEASURED' and r['id'] not in by:keep.append({'id':str(uuid.uuid4()),'source_line_id':r['id'],'quantity':r.get('quantity'),'uom':r.get('uom'),'weight':'1','wbs_id':''})
    s['mapping']=keep
class App(ThreadingHTTPServer):
    daemon_threads=True
    def __init__(self,address,db):
        super().__init__(address,Handler);self.store=Store(db);self.token=secrets.token_urlsafe(32);self.export_dir=Path(db).parent/'exports';self.export_lock=threading.Lock();self.ai=ai.Settings(Path(db).parent/'ai_settings.json');self.qto=qto_bridge.Settings(Path(db).parent/'qto_settings.json')  # E21: CERP-QTO bridge
def summary(result):
    """Browser view of a calculation: totals, coverage and first gaps; full lines stay in the Excel snapshot."""
    return {k:v for k,v in result.items() if k not in ['lines','items']}|{'missing':result['missing'][:25],'line_count':len(result['lines']),'items_preview':result['items'][:200]}
def library_norms(store,body):
    """Copy recipes (by code) and their resource lines from library datasets into a project draft, keeping lineage."""
    norms=store.library_get(body.get('norm_dataset'));res=store.library_get(body['resource_dataset']) if body.get('resource_dataset') else {'rows':[],'id':None,'sha256':None}
    if norms['kind']!='NORMS' or (res['id'] and res['kind']!='RESOURCES'):raise ValueError('Chọn đúng bộ NORMS và RESOURCES.')
    codes={str(c).strip() for c in body.get('codes',[]) if str(c).strip()}
    chosen=[dict(r,library_dataset=norms['id'],library_sha256=norms['sha256']) for r in norms['rows'] if not codes or str(r.get('code')) in codes]
    ids={r.get('id') for r in chosen}
    lines=[dict(r,library_dataset=res['id'],library_sha256=res['sha256']) for r in res['rows'] if r.get('norm_id') in ids]
    for l in lines:
        if l.get('master_code'):l['resource_code']=l['master_code']  # CC1 resource master code (D29/D30) is the price key
    return {'norms':chosen,'resources':lines}
def library_prices(store,s,body):
    """Pick price rows matching the keys the current applications need, in the project period/currency."""
    data=store.library_get(body.get('dataset'))
    if data['kind']!='PRICES':raise ValueError('Chọn bộ dữ liệu PRICES.')
    need=needed_price_keys(s);ctx=s['context']
    rows=[dict(r,library_dataset=data['id'],library_sha256=data['sha256']) for r in data['rows'] if (r.get('resource_code'),r.get('spec'),r.get('uom'),r.get('location')) in need and r.get('period')==ctx.get('period') and r.get('currency')==ctx.get('currency')]
    return {'rows':rows,'needed':len(need),'matched_keys':len({(r.get('resource_code'),r.get('spec'),r.get('uom'),r.get('location')) for r in rows})}
def price_stubs(s):
    """Empty price rows (unit_price null) for every needed key without a price; values must still be entered."""
    have={(p.get('resource_code'),p.get('spec'),p.get('uom'),p.get('location')) for p in s['prices']};ctx=s['context']
    return [{'id':str(uuid.uuid4()),'resource_code':k[0],'spec':k[1],'uom':k[2],'location':k[3],'period':ctx.get('period'),'currency':ctx.get('currency'),'unit_price':None,'tax_basis':'EX_VAT','delivery':'AT_SITE'} for k in sorted(needed_price_keys(s)-have,key=lambda x:tuple(str(v) for v in x))]
class Handler(BaseHTTPRequestHandler):
    server_version='CC1Local/0.2'
    def log_message(self,fmt,*args):pass
    def reply(self,obj,code=200):
        data=json.dumps(obj,ensure_ascii=False).encode();self.send(data,'application/json; charset=utf-8',code)
    def send(self,data,typ,code=200,filename=None):
        self.send_response(code);self.send_header('Content-Type',typ);self.send_header('Content-Length',str(len(data)));self.send_header('Cache-Control','no-store');self.send_header('X-Content-Type-Options','nosniff');self.send_header('Content-Security-Policy',"default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' data:; connect-src 'self'; frame-ancestors 'none'; base-uri 'none'")
        if filename:self.send_header('Content-Disposition','attachment; filename="'+filename+'"')
        self.end_headers();self.wfile.write(data)
    def guard(self,mutation=False):
        if self.headers.get('Host') not in [f'127.0.0.1:{self.server.server_port}',f'localhost:{self.server.server_port}']:raise PermissionError('Host không hợp lệ.')
        origin=self.headers.get('Origin')
        if origin and origin not in [f'http://127.0.0.1:{self.server.server_port}',f'http://localhost:{self.server.server_port}']:raise PermissionError('Origin không hợp lệ.')
        if mutation and self.headers.get('X-CC1-Session')!=self.server.token:raise PermissionError('Thiếu mã phiên làm việc trên máy; tải lại trang.')
    def proxied(self,method):
        """E21: /v1/*, /healthz and /qto/ui/* go to the CERP-QTO service with the stored token. Same-origin only (the
        embedded CERP-QTO page cannot send the 8765 session header); bodies are streamed as-is (multipart uploads)."""
        path=urlparse(self.path).path
        if not qto_bridge.wants_proxy(path):return False
        try:self.guard(False)
        except PermissionError as e:self.reply({'detail':str(e)},403);return True
        if method!='GET':
            ref=self.headers.get('Origin') or self.headers.get('Referer') or ''
            if not ref.startswith((f'http://127.0.0.1:{self.server.server_port}/',f'http://localhost:{self.server.server_port}/',f'http://127.0.0.1:{self.server.server_port}',f'http://localhost:{self.server.server_port}')):
                self.reply({'detail':'Yêu cầu tới CERP-QTO phải xuất phát từ giao diện 8765.'},403);return True
        qto_bridge.proxy(self,self.server.qto,method);return True
    def do_PUT(self):
        if not self.proxied('PUT'):self.reply({'error':'Không có đường dẫn này.'},404)
    def do_DELETE(self):
        if not self.proxied('DELETE'):self.reply({'error':'Không có đường dẫn này.'},404)
    def do_GET(self):
        if self.proxied('GET'):return
        try:
            self.guard();path=urlparse(self.path).path;params=parse_qs(urlparse(self.path).query)
            if path=='/api/qto/status':return self.reply(qto_bridge.status(self.server.qto))
            if extensions.get(self,path,params):return  # D59: E10–E16 routes
            if path=='/api/bootstrap':return self.reply({'token':self.server.token,'steps':STEPS,'projects':self.server.store.list(),'categories':CATEGORIES,'purposes':PURPOSES,'profiles':PROFILES,'template':{'id':'TPL03_COMPOSITION','name':'CC1-LBS-WBS-BOQ-XL-NS-02-02-CAU-ME-SO-V0.2.0-DRAFT.xlsx','sheet':'10_WP_BOQ_COMPOSITION','status':'SOURCE_DRAFT; các bảng dự toán kế tiếp lập trên máy'},'version':SCHEMA+'-local'})
            if path=='/api/library':return self.reply({'datasets':self.server.store.library_list(),'kinds':list(LIBRARY_KINDS)})
            if path=='/api/library-search':
                kind=params.get('kind',['PRICES'])[0]
                if kind not in LIBRARY_KINDS:raise ValueError('Loại thư viện không hợp lệ.')
                return self.reply(self.server.store.library_search(kind,params.get('q',[''])[0][:200],min(int(params.get('limit',['50'])[0]),200),params.get('dataset',[None])[0]))
            if path.startswith('/api/library/'):return self.reply(self.server.store.library_get(path.split('/')[3]))
            if path=='/api/ai/providers':return self.reply(self.server.ai.public())
            if path=='/api/code-library/search':return self.reply(codelib.search_norms(self.server.store,params.get('q',[''])[0][:200],min(int(params.get('limit',['50'])[0]),200)))
            if path=='/api/code-library/norm':return self.reply(codelib.norm_detail(self.server.store,params.get('id',[''])[0]))
            if path=='/api/code-library.xlsx':
                if not CODE_LIBRARY_XLSX.exists():raise FileNotFoundError('Chưa có file thư viện mã; chạy build_code_library.py.')
                return self.send(CODE_LIBRARY_XLSX.read_bytes(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',filename=CODE_LIBRARY_XLSX.name)
            if path=='/api/catalog':
                query=params.get('q',[''])[0].lower();return self.reply({'work_results':[dict(code=r['business_code'],name=r['name'],status=r['readiness']) for r in catalog()['WORK_RESULTS'] if query in (r['business_code']+' '+r['name']).lower()]})
            if path=='/api/norm-search':
                query=params.get('q',[''])[0].lower()
                rows=json.loads((ROOT/'runs_brian/tt38_boq_20261017_01/tt38-index.json').read_text(encoding='utf-8'))
                return self.reply({'rows':[r for r in rows if query and query in (r['code']+' '+r['title']).lower()][:40],'note':'Chỉ mục có mã nhóm/tham chiếu; không tự chọn mã chi tiết.'})
            if path=='/api/structures':return self.reply({'families':catalog()['FAMILIES'],'catalog':catalog()['LEGACY_CATALOG'],'relations':catalog()['LEGACY_RELATIONS']})
            if path=='/api/builtin':return self.reply({'rows':builtin(params.get('source',['AB'])[0])})
            if path.startswith('/api/projects/'):
                bits=path.split('/');pid=bits[3];s=self.server.store.get(pid)
                if len(bits)>4 and bits[4]=='backup':return self.send(json.dumps({'schema_version':SCHEMA,'project':s,'audit':self.server.store.events(pid)},ensure_ascii=False,indent=2).encode(),'application/json; charset=utf-8',filename='CC1_PROJECT_'+pid+'.json')
                if len(bits)>4 and bits[4]=='ai':return self.reply({'proposals':self.server.store.proposals(pid)})
                return self.reply({'project':s,'stages':status(s),'audit':self.server.store.events(pid),'calculation':summary(calculate(s))})
            if path.startswith('/api/download/'):
                name=path.split('/')[-1]
                if not __import__('re').fullmatch(r'[a-f0-9-]+\.xlsx',name):raise ValueError('Tên bản số liệu không hợp lệ.')
                return self.send((self.server.export_dir/name).read_bytes(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',filename=name)
            asset=__import__('re').fullmatch(r'/(brand|fonts)/([a-z0-9-]+)\.(png|woff2)',path)  # CC1 brand logo + vendored Inter (D49)
            if asset and (HERE/'static'/asset.group(1)/(asset.group(2)+'.'+asset.group(3))).is_file():
                return self.send((HERE/'static'/asset.group(1)/(asset.group(2)+'.'+asset.group(3))).read_bytes(),{'png':'image/png','woff2':'font/woff2'}[asset.group(3)])
            if path=='/favicon.ico':return self.send((HERE/'static/brand/cc1-mark.png').read_bytes(),'image/png')
            allowed={'/':'index.html','/app.js':'app.js','/validation.js':'validation.js','/subcontract.js':'subcontract.js','/flows.js':'flows.js','/theme-init.js':'theme-init.js','/style.css':'style.css'}
            if path not in allowed:return self.reply({'error':'Không có đường dẫn này.'},404)
            p=HERE/'static'/allowed[path];return self.send(p.read_bytes(),mimetypes.guess_type(p.name)[0]+('; charset=utf-8' if p.suffix in ['.html','.js','.css'] else ''))
        except PermissionError as e:self.reply({'error':str(e)},403)
        except (ValueError,KeyError,FileNotFoundError) as e:self.reply({'error':str(e)},404)
        except Exception as e:self.reply({'error':'Không đọc được hồ sơ: '+str(e)},400)
    def do_POST(self):
        if self.proxied('POST'):return
        try:
            self.guard(True);length=int(self.headers.get('Content-Length','0'))
            if not 0<length<=MAX:raise ValueError('Dữ liệu gửi lên rỗng hoặc vượt 35 MB.')
            body=json.loads(self.rfile.read(length));path=urlparse(self.path).path
            if extensions.post(self,path,body):return  # D59: E10–E16 routes
            if path=='/api/shutdown':
                self.reply({'status':'STOPPING_LOCAL_APP'});threading.Thread(target=self.server.shutdown,daemon=True).start();return
            if path=='/api/qto/settings':  # E21: token is write-only; the reply is the status (user / roles / projects), never the token
                self.server.qto.set(body.get('url'),body.get('token') if 'token' in body else None);return self.reply(qto_bridge.status(self.server.qto))
            if path=='/api/projects':
                name=str(body.get('name','')).strip()
                if not name or len(name)>200:raise ValueError('Nhập tên hồ sơ từ 1 đến 200 ký tự.')
                return self.reply({'project':self.server.store.create(name)},201)
            if path=='/api/import':
                result,binary=imported(body);self.server.store.upload(result['sha256'],result['filename'],binary);return self.reply(result)
            if path=='/api/restore':
                if body.get('schema_version') not in ['0.1.0',SCHEMA] or not isinstance(body.get('project'),dict):raise ValueError('Chỉ nhận bản sao lưu cấu trúc 0.1.0 hoặc '+SCHEMA+'.')
                old=body['project'];original=digest(old)
                for i,key in enumerate(KEYS):
                    expected=dict if i in [0,6,7] else list
                    if not isinstance(old.get(key),expected):raise ValueError('Bản sao lưu thiếu/sai kiểu '+key)
                for key in ['wbs','applications']:
                    if not isinstance(old.get(key,[]),list):raise ValueError('Bản sao lưu thiếu/sai kiểu '+key)
                old=upgrade(dict(old,schema_version=body['schema_version']))
                s=self.server.store.create(str(old.get('name','Hồ sơ khôi phục'))[:200]+' · khôi phục')
                for key in KEYS+['wbs','applications','migrations','packages','quotes','progress','drawings','takeoffs','erp_reference','meta']:
                    if key in old:s[key]=old[key]
                if isinstance(old.get('contracts'),list):  # E18: dossiers restored as evidence; READY goes back to review
                    import contracts;s['contracts']=contracts.restore(old);s['changes']=old.get('changes') or [];s['facts']=old.get('facts') or []
                if isinstance(old.get('pakd'),dict):  # D77: PAKD scenarios + frozen versions are evidence; mark them restored, never new approvals
                    s['pakd']=dict(old['pakd'],versions=[dict(v,restored=True) for v in old['pakd'].get('versions') or []])
                s['resource_completeness_ref']=old.get('resource_completeness_ref');s['restored_from']={'project_id':old.get('id'),'schema_version':body['schema_version'],'digest':original,'prior_audit':body.get('audit',[]),'note':'Bằng chứng nạp từ bản sao lưu trên máy; các xác nhận đã đặt lại, cần kiểm tra lại.'}
                s=self.server.store.put(s,0,'RESTORE_LOCAL_CONTENT_REVIEW_REQUIRED');return self.reply({'project':s},201)
            if path=='/api/library':
                kind=body.get('kind');name=str(body.get('name','')).strip()
                if kind not in LIBRARY_KINDS:raise ValueError('Chọn loại NORMS, RESOURCES hoặc PRICES.')
                if not name or len(name)>200:raise ValueError('Nhập tên bộ dữ liệu 1–200 ký tự.')
                result,binary=imported(dict(body,table=LIBRARY_KINDS[kind]))
                if not result.get('records'):raise ValueError('File không có dòng dữ liệu.')
                if kind=='PRICES':  # D59/FR-16: new price sets are checked, never zero-filled
                    import libtools;errs=libtools.validate_prices(result['records'])
                    if errs:raise ValueError(f'{len(errs)} lỗi bảng giá, chưa nạp: '+'; '.join(errs[:6]))
                self.server.store.upload(result['sha256'],result['filename'],binary)
                return self.reply(self.server.store.library_add(kind,name,result['sha256'],result['filename'],str(body.get('note',''))[:1000],result['records']),201)
            if path.startswith('/api/library/') and path.endswith('/archive'):
                self.server.store.library_archive(path.split('/')[3]);return self.reply({'datasets':self.server.store.library_list()})
            if path=='/api/ai/settings':
                try:return self.reply(self.server.ai.update(body))
                except ai.AIError as e:return self.reply({'error':str(e)},400)
            if path=='/api/ai/models':
                try:return self.reply({'models':ai.list_models(self.server.ai,body.get('provider')),'providers':self.server.ai.public()})
                except ai.AIError as e:return self.reply({'error':str(e)},502)
            if path=='/api/ai/ask':
                s=self.server.store.get(str(body.get('project_id')));step=body.get('step')
                if step not in range(8):raise ValueError('Bước không hợp lệ.')
                try:answer=ai.ask(self.server.ai,s,step,body.get('provider'),str(body.get('model') or ''),body.get('question'),body.get('include_rows'),body.get('consent'))
                except ai.AIError as e:return self.reply({'error':str(e)},502)
                proposal=self.server.store.proposal_add(s['id'],step,body.get('provider'),answer['model'],body['question'].strip(),answer['prompt_digest'],answer['text'])
                return self.reply({'proposal':proposal,'external':answer['external']})
            if path.startswith('/api/projects/'):
                bits=path.split('/');pid=bits[3];s=self.server.store.get(pid)
                if bits[4]=='qto-project':
                    try:s=qto_bridge.link_project(self.server.store,s,body)
                    except PermissionError as e:return self.reply({'error':str(e)},409)
                    return self.reply({'project':s})
                if bits[4]=='qto-import':
                    try:s,res=qto_bridge.import_takeoff(self.server.store,self.server.qto,s,body)
                    except PermissionError as e:return self.reply({'error':str(e)},409 if 'cửa sổ' in str(e) else 401)
                    except ConnectionError as e:return self.reply({'error':str(e)},503)
                    return self.reply({'project':s,'result':res,'review':qto_bridge.review(s,body.get('threshold_pct'))})
                if bits[4]=='qto-review':return self.reply({'review':qto_bridge.review(s,body.get('threshold_pct'))})
                if bits[4]=='library-norms':return self.reply(library_norms(self.server.store,body))
                if bits[4]=='library-prices':return self.reply(library_prices(self.server.store,s,body))
                if bits[4]=='price-stubs':return self.reply({'rows':price_stubs(s)})
                if bits[4]=='apply-memory':return self.reply(codelib.apply_memory(self.server.store,s))
                if bits[4]=='auto-prices':return self.reply(codelib.auto_prices(self.server.store,s))
                if bits[4]=='suggest-norms':
                    ids=body.get('item_ids');mem=codelib.memory_for(self.server.store,s,set(ids) if isinstance(ids,list) else None);items=codelib.suggest_items(self.server.store,s,set(ids) if isinstance(ids,list) else None,max(1,min(int(body.get('limit') or 5),10)),str(body.get('method') or '')[:12])
                    for it in items:it['memory']=mem.get(it['item_id'])
                    return self.reply({'items':items})
                if bits[4]=='code-norm':
                    d=codelib.norm_detail(self.server.store,str(body.get('id','')));norm,lines=codelib.for_project(d['norm'],d['lines']);return self.reply({'norm':norm,'resources':lines})
                if bits[4]=='stage':
                    step=int(bits[5]);expected=body.get('revision')
                    if expected!=s['revision']:raise Conflict('Hồ sơ đã đổi; tải lại để tránh ghi đè.')
                    if step not in range(8):raise ValueError('Bước không hợp lệ.')
                    if not all(r['confirmed'] for r in status(s)[:step]):return self.reply({'error':'Hoàn tất các bước trước, không được bỏ qua.','stages':status(s)},409)
                    data=body.get('data');want=dict if step in [0,6,7] else list
                    if not isinstance(data,want):raise ValueError('Sai cấu trúc dữ liệu cho bước.')
                    if isinstance(data,list) and (len(data)>20000 or any(not isinstance(x,dict) for x in data)):raise ValueError('Mỗi dòng phải là một đối tượng; tối đa 20.000 dòng.')
                    if type(body.get('confirm',False)) is not bool:raise ValueError('Giá trị xác nhận (confirm) phải là đúng/sai.')
                    if step==1:
                        old_byid={r['id']:r for r in s['boq']};old_boq=s['boq']
                        for row in data:
                            old=old_byid.get(row.get('id'))
                            if old:
                                # D73: values may be edited by hand (code, description, BOQ price, quantity, UOM); raw/hash/source never change and the
                                # source value is kept server-side in original_<field> the first time it differs (lineage). Domain asks for disposition_ref.
                                for field in ['raw','source_hash','source_ref']:
                                    if old.get(field)!=row.get(field):raise ValueError('Không sửa dữ liệu gốc/mã băm/nguồn của dòng BOQ đã nạp; sửa giá trị thì ghi căn cứ, nguồn đổi thì nạp bản số liệu kế tiếp.')
                                for f in BOQ_TRACKED:
                                    o='original_'+f
                                    if o in old:row[o]=old[o]
                                    elif row.get(f)!=old.get(f):row[o]=old.get(f)
                                    else:row.pop(o,None)
                            elif str(row.get('source_ref') or '').startswith('MANUAL'):
                                # a line typed by the user: server-side hash of what was entered; no original_* until it is edited later
                                for f in ['raw','source_hash']+['original_'+x for x in BOQ_TRACKED]:row.pop(f,None)
                                entered={k:row.get(k) for k in ('source_code','description','kind','quantity','uom','unit_price')}
                                row['raw']={'MANUAL':True,'entered':entered,'by':s['context'].get('owner')}
                                row['source_hash']=hashlib.sha256(json.dumps(dict(entered,id=row.get('id')),ensure_ascii=False,sort_keys=True,default=str).encode()).hexdigest()
                            elif not (isinstance(row.get('source_hash'),str) and len(row['source_hash'])==64 and row.get('source_ref')):
                                raise ValueError('Dòng BOQ mới phải từ file đã nạp (có mã băm, nguồn) hoặc là dòng nhập tay (source_ref MANUAL…).')
                    s[KEYS[step]]=data;s['confirmed']=[i for i in s['confirmed'] if i<step]
                    for key,at in [('wbs',2),('applications',3),('resources',3)]:  # resources may arrive with recipes pulled from the library; later steps are reopened anyway
                        if step==at and body.get(key) is not None:
                            if not isinstance(body[key],list) or len(body[key])>20000 or any(not isinstance(x,dict) for x in body[key]):raise ValueError(key+' phải là danh sách đối tượng, tối đa 20.000 dòng.')
                            s[key]=body[key]
                    if step==3 and body.get('assignments') is not None:
                        assignments=body['assignments']
                        if not isinstance(assignments,list) or {a.get('id') for a in assignments}!={p['id'] for p in s['mapping']}:raise ValueError('Bảng gán tổ hợp định mức phải phủ đúng tập hạng mục chuẩn hóa hiện tại.')
                        byid={a['id']:a for a in assignments}
                        for item in s['mapping']:
                            for key in ['technical_spec','method','norm_match_ref']:item[key]=byid[item['id']].get(key)
                    if step==4:s['resource_completeness_ref']=body.get('resource_completeness_ref')
                    qs=status(s)[step]['questions']
                    if body.get('confirm') and not qs:s['confirmed'].append(step)
                    if step==1:sync_mapping(s,old_boq,data)
                    s=self.server.store.put(s,expected,'CONFIRM_LOCAL_INPUT' if step in s['confirmed'] else 'SAVE_DRAFT_STEP_'+str(step))
                    if step==3 and 3 in s['confirmed']:codelib.remember(self.server.store,s)  # BL63: remember gate-passed QS choices
                    return self.reply({'project':s,'stages':status(s),'questions':qs,'confirmed':step in s['confirmed']},200)
                if bits[4]=='export':
                    if body.get('revision')!=s['revision']:raise Conflict('Phiên bản hồ sơ đã đổi; tải lại trước khi xuất.')
                    if not all(r['confirmed'] for r in status(s)):return self.reply({'error':'Chưa đủ đầu vào để xuất dự toán đầy đủ.','stages':status(s)},409)
                    key=digest(s);name=key+'.xlsx';out=self.server.export_dir/name
                    with self.server.export_lock:
                        if not out.exists():
                            temp=out.with_suffix('.tmp');workbook(s,self.server.store.events(pid),temp);os.replace(temp,out)
                    return self.reply({'url':'/api/download/'+name,'snapshot_sha256':key,'file_sha256':hashlib.sha256(out.read_bytes()).hexdigest(),'status':'LOCAL_DRAFT_NOT_CORPORATE_RELEASE'})
            self.reply({'error':'Không có API này.'},404)
        except Conflict as e:self.reply({'error':str(e)},409)
        except PermissionError as e:self.reply({'error':str(e)},403)
        except (ValueError,KeyError,TypeError,IndexError,AttributeError) as e:self.reply({'error':'Kiểm tra đầu vào: '+str(e)},400)
        except Exception as e:self.reply({'error':'Không lưu được hồ sơ: '+str(e)},400)
def main():
    p=argparse.ArgumentParser();p.add_argument('--port',type=int,default=8765);p.add_argument('--db',default=str(HERE/'runtime/app.sqlite'));a=p.parse_args()
    app=App(('127.0.0.1',a.port),a.db);print(f'CC1 local: http://127.0.0.1:{app.server_port}',flush=True)
    try:app.serve_forever()
    except KeyboardInterrupt:pass
    finally:app.server_close()
if __name__=='__main__':main()
