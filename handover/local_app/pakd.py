"""E17 (D76): PAKD — phương án kinh doanh thi công (bảng BM01-TH) on top of the estimate and the subcontract packages.

Structure of BM01 follows the PAKD Giao thầu CC1 package the user supplied (B doanh thu, C.1.1–C.1.10, C.2–C.5, D lợi nhuận
gộp), computed from THIS app's data instead of a separate engine:
  B      = Σ AB quantity × AB unit price of measured lines (a line without price → B is unknown, listed as missing).
  C.1.4  VT CC1 cấp       = estimate VL of the shares CC1 supplies (TÁCH with VL=CC1, or the CC1 self package).
  C.1.5  Giao NTP trọn gói = subcontract unit price × allocated quantity of GỘP packages.
  C.1.6  Giao NTP theo phần = same for TÁCH packages (price covers only the parts the subcontractor carries).
  C.1.6b CC1 tự thực hiện  = estimate NC/M/other of the shares CC1 performs.
  C.1.7  Chưa giao thầu     = estimate cost of the unallocated share of each AB line.
  C.1.1–C.1.3, C.1.8–C.1.10, C.2, C.3, C.4, C.5 = lines the user enters (amount, or % × an allowed base) with a source.
  D = B − C; tỷ suất = D / B.

Rules (review findings F01/F03/F06 on the PAKD package): a missing price, cost or basis is never 0 — the affected rows are
null and the known subtotal plus a missing list are shown; the subcontract price comes from the basis chosen per package
(award draft, A, B or C of E11), never from AB price ratios; there are NO default percentages; a frozen version is
append-only with the digest of its inputs and is never overwritten. Estimate overheads of step 7 (rules C40–C70) are not
used here: PAKD indirect costs are entered as C.2/C.3 lines, so nothing is counted twice.
"""
import copy,hashlib,json,uuid
from datetime import datetime
from decimal import Decimal,ROUND_HALF_UP
from domain import number,index,meaningful
import subcontract as sc

HUNDRED=Decimal(100)
USER_CODES={'C.1.1':'Thiết kế bản vẽ thi công','C.1.2':'Thí nghiệm','C.1.3':'Điện phục vụ thi công / trạm biến áp',
            'C.1.8':'Phát sinh và tạm tính','C.1.9':'Chi phí khác','C.1.10':'Dự phòng rủi ro',
            'C.2':'Chi phí gián tiếp','C.3':'Chi phí tài chính','C.4':'Chi phí bảo hành','C.5':'Chi phí quyết toán'}
# which bases a percentage line may use (no cycles: C.1.8–10 cannot use C1, C.4/C.5 come last)
BASES={'C.1.1':(),'C.1.2':(),'C.1.3':(),'C.1.8':('DT','CPXD'),'C.1.9':('DT','CPXD'),'C.1.10':('DT','CPXD'),
       'C.2':('DT','CPXD','C1'),'C.3':('DT','CPXD','C1'),'C.4':('DT','C1'),'C.5':('DT','C1')}
BASE_NAMES={'DT':'Doanh thu (B)','CPXD':'Chi phí xây dựng (C.1.4–C.1.7)','C1':'Chi phí trực tiếp (C.1)'}
COMPUTED={'C.1.4':'Vật tư CC1 cấp','C.1.5':'Giao thầu NTP trọn gói','C.1.6':'Giao thầu NTP theo phần (NC / Máy / VT)',
          'C.1.6b':'CC1 tự thực hiện (NC, Máy, khác)','C.1.7':'Chưa giao thầu (theo dự toán CC1)'}
BASIS_OPTIONS={'AWARD':'Giá đề xuất giao thầu (NHÁP)','A':'Căn cứ A · đơn giá AB × K','B':'Căn cứ B · chi phí CC1 + phí','C':'Căn cứ C · báo giá hợp lệ thấp nhất'}
ORDER=['C.1.1','C.1.2','C.1.3','C.1.4','C.1.5','C.1.6','C.1.6b','C.1.7','C.1.8','C.1.9','C.1.10']
# D86: Quy trình PAKD Thi công v1.0 (11/10/2026) §5 — cấu trúc A–I. BM01 giữ nguyên; mỗi khoản được quy chiếu sang một nhóm.
# Khoản tính sẵn (C.1.4–C.1.7) thuộc B; G = giá giao thầu phụ là phần của B (H = A − B − C − D − E − F, không trừ G lần nữa).
# Khoản người lập nhập: mặc định chỉ khi tên khoản trùng nghĩa với nhóm; còn lại người lập chọn — không đoán.
PROCESS_GROUPS={'A':'Doanh thu AB','B':'Chi phí trực tiếp','C':'Chi phí gián tiếp','D':'Chuẩn bị thi công','E':'Tài chính, bảo hiểm, thuế phí','F':'Dự phòng',
                'G':'Giá giao thầu phụ (nằm trong B)','H':'Lãi gộp mục tiêu = A − B − C − D − E − F','I':'EAC và margin forecast'}
PROCESS_CHOICES=('B','C','D','E','F')
PROCESS_DEFAULT={'C.2':'C','C.3':'E','C.1.8':'F','C.1.10':'F'}
NOTE='DỰ THẢO TRÊN MÁY — PAKD tham khảo, không phải PAKD được duyệt; số liệu theo dữ liệu hồ sơ hiện tại.'

def s_(d):return sc.s_(d)
def _places(s):
    p=s['context'].get('money_places')
    try:p=int(p)
    except (TypeError,ValueError):p=0
    return p if p in (0,2) else 0
def _r(d,places):return None if d is None else s_(d.quantize(Decimal(1) if places==0 else Decimal('0.01'),ROUND_HALF_UP))
# ---------- state ----------
def state(s):
    """The PAKD part of the project (scenarios, current, versions); a first scenario exists virtually until saved."""
    p=s.get('pakd') or {}
    scen=p.get('scenarios') or [{'id':'scenario-1','name':'Kịch bản 1','basis':{},'lines':[]}]
    return {'scenarios':scen,'current':p.get('current') or scen[0]['id'],'versions':p.get('versions') or []}
def _scen(st,sid):
    x=next((c for c in st['scenarios'] if c['id']==sid),None)
    if not x:raise ValueError('Không có kịch bản này.')
    return x
def check_lines(lines):
    """Validate user lines; returns a clean copy. Raises ValueError on structural errors (unknown code/method/base)."""
    if not isinstance(lines,list):raise ValueError('Các dòng chi phí phải là danh sách.')
    out=[]
    for i,l in enumerate(lines):
        if not isinstance(l,dict):raise ValueError('Mỗi dòng chi phí là một đối tượng.')
        code=l.get('code')
        if code not in USER_CODES:raise ValueError(f'Dòng {i+1}: mã khoản mục {code!r} không thuộc C.1.1–C.1.3, C.1.8–C.1.10, C.2–C.5.')
        method=l.get('method') or 'AMOUNT'
        if method not in ('AMOUNT','PERCENT'):raise ValueError(f'Dòng {i+1}: phương pháp phải là Số tiền (AMOUNT) hoặc Tỷ lệ % (PERCENT).')
        if method=='PERCENT' and l.get('base') not in BASES[code]:
            raise ValueError(f'Dòng {i+1} ({code}): cơ sở % phải là '+(', '.join(BASES[code]) or 'không có — khoản này chỉ nhập số tiền')+'.')
        pg=l.get('process_group') or None
        if pg not in (None,*PROCESS_CHOICES):raise ValueError(f'Dòng {i+1} ({code}): nhóm quy trình phải là B, C, D, E hoặc F.')
        out.append({'id':l.get('id') or str(uuid.uuid4()),'code':code,'name':str(l.get('name') or USER_CODES[code])[:200],'method':method,'process_group':pg,
                    'base':l.get('base') if method=='PERCENT' else None,'value':None if l.get('value') in (None,'') else str(l.get('value')),
                    'source_ref':str(l.get('source_ref') or '')[:1000],'on':l.get('on') is not False})
    return out
def questions(s,scen):
    q=[]
    for i,l in enumerate(scen.get('lines') or []):
        if not l.get('on',True):continue
        f=f'pakd.lines[{i}]';v=number(l.get('value'))
        if v is None:q.append({'field':f+'.value','question':f'{l["code"]} {l["name"]}: nhập '+('tỷ lệ %' if l['method']=='PERCENT' else 'số tiền')+' (trống không được coi là 0).'})
        elif v<0 or (l['method']=='PERCENT' and v>HUNDRED):q.append({'field':f+'.value','question':f'{l["code"]} {l["name"]}: giá trị ngoài khoảng cho phép.'})
        if not meaningful(l.get('source_ref')):q.append({'field':f+'.source_ref','question':f'{l["code"]} {l["name"]}: ghi căn cứ (quy chế, báo giá, tính toán) — không có tỷ lệ mặc định.'})
    wg=work_groups(s)
    for p in sc.packages(s):
        if p.get('is_self') or not _uses_main(scen,p,wg):continue
        if (scen.get('basis') or {}).get(p['id']) not in BASIS_OPTIONS:
            q.append({'field':f'pakd.basis.{p["code"]}','question':f'Gói {p["code"]}: chọn căn cứ giá giao dùng cho PAKD (giá đề xuất / A / B / C).'})
    for g in plan_questions(s,scen):q.append(g)
    q+=cashflow_questions(scen)
    return q
# ---------- parts of a package (prompt 6.11) ----------
NOGROUP='—';MAIN='main';MAX_PARTS=20
def work_groups(s):
    """Work group of every measured AB line = the deepest BOQ heading whose code is a dotted prefix of the line code
    (prompt 6.7 reads headings deepest-first). No keyword taxonomy is invented; a line without such heading is in '—'."""
    heads={str(r['source_code']).strip():str(r.get('description') or '') for r in s['boq'] if r.get('kind')=='HEADING' and meaningful(r.get('source_code'))}
    out={}
    for i,a in sc.measured(s).items():
        bits=str(a.get('source_code') or '').strip().split('.');g=None
        for n in range(len(bits)-1,0,-1):
            k='.'.join(bits[:n])
            if k in heads:g=k;break
        out[i]=(g,heads[g]) if g else (NOGROUP,'Không có tiêu đề BOQ')
    return out
def package_groups(s,p,wg=None):
    """Groups of one package for the UI: code, name, its BB line codes and AB value of the package share (null if a price is missing)."""
    wg=wg or work_groups(s);ab=sc.measured(s);out={}
    for l in p.get('lines') or []:
        g=wg.get(l['ab_line_id'],(NOGROUP,'Không có tiêu đề BOQ'));x=out.setdefault(g[0],{'code':g[0],'name':g[1],'bb_codes':[],'value':Decimal(0),'ok':True})
        x['bb_codes'].append(l['bb_code']);a=ab.get(l['ab_line_id']) or {};q=number(a.get('quantity'));up=number(a.get('unit_price'));al=number(l.get('alloc_pct'))
        if None in (q,up,al):x['ok']=False
        else:x['value']+=q*up*al/HUNDRED
    return [dict(x,value=s_(x['value']) if x.pop('ok') else None) for x in out.values()]
def check_parts(s,parts):
    """Validate the per-package split of a scenario: {package_id: {parts:[…], by_group:{group: part}, by_line:{bb_code: part|'main'}}}.
    Parts that end up without any line are dropped (prompt 6.11: an empty part deletes itself)."""
    if parts in (None,''):return {}
    if not isinstance(parts,dict):raise ValueError('Chia phần phải là đối tượng theo gói.')
    pk={p['id']:p for p in sc.packages(s) if not p.get('is_self')};wg=work_groups(s);out={}
    for pid,plan in parts.items():
        if pid not in pk:raise ValueError('Chia phần: gói không tồn tại hoặc là gói CC1 tự làm.')
        if not isinstance(plan,dict):raise ValueError('Chia phần: mỗi gói là một đối tượng.')
        p=pk[pid];raw=plan.get('parts') or []
        if not isinstance(raw,list) or len(raw)>MAX_PARTS:raise ValueError(f'Gói {p["code"]}: tối đa {MAX_PARTS} phần.')
        clean=[];ids=set()
        for x in raw:
            if not isinstance(x,dict):raise ValueError(f'Gói {p["code"]}: mỗi phần là một đối tượng.')
            n=str(x.get('name') or '').strip()
            if not n:raise ValueError(f'Gói {p["code"]}: phần chưa có tên.')
            pid2=str(x.get('id') or uuid.uuid4())
            if pid2==MAIN or pid2 in ids:raise ValueError(f'Gói {p["code"]}: id phần trùng.')
            mode=x.get('mode') or None;basis=x.get('basis') or None
            if mode not in (None,)+sc.MODES:raise ValueError(f'Gói {p["code"]} phần "{n}": hình thức phải là Giao gộp (GOP) hoặc Giao tách (TACH).')
            if basis not in (None,*BASIS_OPTIONS):raise ValueError(f'Gói {p["code"]} phần "{n}": căn cứ giá giao phải là Giá đề xuất giao thầu (AWARD), Căn cứ A, B hoặc C.')
            prov={k:(x.get(k) or None) for k in sc.PARTS}
            if any(v not in (None,)+sc.PROVIDERS for v in prov.values()):raise ValueError(f'Gói {p["code"]} phần "{n}": bên chịu VL/NC/M phải là NTP hoặc CC1.')
            ids.add(pid2);clean.append({'id':pid2,'name':n[:120],'mode':mode,'basis':basis,**(prov if mode=='TACH' else {k:None for k in sc.PARTS})})
        lines={l['bb_code']:l for l in p.get('lines') or []};groups={wg.get(l['ab_line_id'],(NOGROUP,))[0] for l in lines.values()}
        bg={};bl={}
        for g,v in (plan.get('by_group') or {}).items():
            if v in (None,'',MAIN):continue
            if g not in groups:raise ValueError(f'Gói {p["code"]}: nhóm {g} không có dòng nào trong gói.')
            if v not in ids:raise ValueError(f'Gói {p["code"]}: nhóm {g} gán vào phần không tồn tại.')
            bg[g]=v
        for c,v in (plan.get('by_line') or {}).items():
            if v in (None,''):continue
            if c not in lines:raise ValueError(f'Gói {p["code"]}: dòng {c} không thuộc gói.')
            if v!=MAIN and v not in ids:raise ValueError(f'Gói {p["code"]}: dòng {c} gán vào phần không tồn tại.')
            bl[c]=v
        used={_part_id(bg,bl,c,wg.get(l['ab_line_id'],(NOGROUP,))[0]) for c,l in lines.items()}
        clean=[x for x in clean if x['id'] in used]
        keep={x['id'] for x in clean}
        if clean:out[pid]={'parts':clean,'by_group':{g:v for g,v in bg.items() if v in keep},'by_line':{c:v for c,v in bl.items() if v==MAIN or v in keep}}
    return out
def _part_id(bg,bl,bb_code,group):
    """Line → part: an explicit line choice first, then its work group, else the common part (prompt 6.11)."""
    return bl.get(bb_code) or bg.get(group) or MAIN
def _uses_main(scen,p,wg):
    plan=(scen.get('parts') or {}).get(p['id']) or {};ls=p.get('lines') or []
    return not ls or any(_part_id(plan.get('by_group') or {},plan.get('by_line') or {},l['bb_code'],wg.get(l['ab_line_id'],(NOGROUP,))[0])==MAIN for l in ls)
def plan_questions(s,scen):
    q=[];pk={p['id']:p for p in sc.packages(s)}
    for pid,plan in (scen.get('parts') or {}).items():
        p=pk.get(pid)
        if not p:continue
        for i,x in enumerate(plan.get('parts') or []):
            f=f'pakd.parts.{p["code"]}[{i}]'
            if x.get('mode') not in sc.MODES:q.append({'field':f+'.mode','question':f'Gói {p["code"]} phần "{x["name"]}": chọn hình thức giao thầu GỘP/TÁCH.'})
            elif x['mode']=='TACH' and any(x.get(k) not in sc.PROVIDERS for k in sc.PARTS):q.append({'field':f+'.parts','question':f'Gói {p["code"]} phần "{x["name"]}": chọn bên chịu VL, NC, M.'})
            if x.get('basis') not in BASIS_OPTIONS:q.append({'field':f+'.basis','question':f'Gói {p["code"]} phần "{x["name"]}": chọn căn cứ giá giao.'})
    return q
# ---------- cash flow (PRD FR22) ----------
TERM_KEYS=('advance_pct','advance_period','recovery_pct','retention_pct','retention_release','delay')
TERM_NAMES={'advance_pct':'tạm ứng %','advance_period':'kỳ tạm ứng','recovery_pct':'thu hồi tạm ứng % mỗi kỳ nghiệm thu','retention_pct':'giữ lại %',
            'retention_release':'kỳ hoàn trả tiền giữ lại','delay':'số kỳ chậm thanh toán sau nghiệm thu'}
NOT_MODELLED='Chưa mô hình hóa (chờ Tài chính quyết định): VAT đầu ra/đầu vào, bảo lãnh tạm ứng/thực hiện/bảo hành, điều kiện mua hàng. Dòng tiền trước VAT.'
MAX_PERIODS=120
def _pct_list(v,n,label):
    if v in (None,''):return [None]*n
    if not isinstance(v,list):raise ValueError(f'{label}: phải là danh sách % theo kỳ.')
    v=[None if x in (None,'') else str(x) for x in v[:n]]
    for x in v:
        if x is not None and (number(x) is None or number(x)<0):raise ValueError(f'{label}: giá trị {x!r} không hợp lệ.')
    return v+[None]*(n-len(v))
def check_cashflow(cf):
    """Validate cash-flow inputs; every term is entered by the user with a source — there are no default terms."""
    if cf in (None,''):return None
    if not isinstance(cf,dict):raise ValueError('Dòng tiền phải là đối tượng.')
    try:n=int(cf.get('periods'))
    except (TypeError,ValueError):raise ValueError('Dòng tiền: nhập số kỳ (tháng).')
    if not 1<=n<=MAX_PERIODS:raise ValueError(f'Dòng tiền: số kỳ 1–{MAX_PERIODS}.')
    start=str(cf.get('start') or '').strip()
    if start and not (len(start)==7 and start[4]=='-' and start[:4].isdigit() and start[5:].isdigit() and 1<=int(start[5:])<=12):raise ValueError('Dòng tiền: tháng bắt đầu dạng YYYY-MM.')
    def terms(t,label,keys=TERM_KEYS):
        t=t if isinstance(t,dict) else {};out={'ref':str(t.get('ref') or '')[:1000]}
        for k in keys:
            v=t.get(k);out[k]=None if v in (None,'') else str(v)
            if out[k] is not None and (number(out[k]) is None or number(out[k])<0):raise ValueError(f'Dòng tiền {label}: {TERM_NAMES[k]} không hợp lệ.')
        return out
    return {'periods':n,'start':start,'revenue_pct':_pct_list(cf.get('revenue_pct'),n,'Lịch nghiệm thu AB'),'ntp_pct':_pct_list(cf.get('ntp_pct'),n,'Lịch nghiệm thu NTP'),
            'own_pct':_pct_list(cf.get('own_pct'),n,'Lịch chi CC1'),'schedule_ref':str(cf.get('schedule_ref') or '')[:1000],
            'ab':terms(cf.get('ab'),'AB'),'ntp':terms(cf.get('ntp'),'NTP'),'own':terms(cf.get('own'),'CC1',('delay',)),
            'rate_pct':None if cf.get('rate_pct') in (None,'') else str(cf.get('rate_pct')),'rate_ref':str(cf.get('rate_ref') or '')[:1000]}
def cashflow_questions(scen):
    cf=scen.get('cashflow')
    if not cf:return []
    q=[];x=compute_cashflow_inputs(cf)
    for m in x:q.append({'field':'pakd.cashflow.'+m[0],'question':'Dòng tiền: '+m[1]})
    return q
def compute_cashflow_inputs(cf):
    """Missing / inconsistent cash-flow inputs as (field, text); amounts are not needed for this check."""
    out=[];n=cf['periods']
    if not meaningful(cf.get('schedule_ref')):out.append(('schedule_ref','ghi căn cứ lịch (tiến độ thi công / kế hoạch nghiệm thu).'))
    for key,label in (('revenue_pct','lịch nghiệm thu AB'),('ntp_pct','lịch nghiệm thu NTP'),('own_pct','lịch chi phí CC1')):
        v=cf.get(key) or []
        if any(number(x) is None for x in v):out.append((key,f'{label}: nhập % cho mọi kỳ (kỳ không có thì nhập 0).'))
        elif abs(sum((number(x) for x in v),Decimal(0))-HUNDRED)>Decimal('0.0001'):out.append((key,f'{label}: tổng % phải bằng 100 (đang {s_(sum((number(x) for x in v),Decimal(0)))}).'))
    for side,label in (('ab','điều khoản thanh toán HĐ AB'),('ntp','điều khoản thanh toán NTP')):
        t=cf.get(side) or {}
        need=['advance_pct','retention_pct','delay']
        if (number(t.get('advance_pct')) or 0)>0:need+=['advance_period','recovery_pct']
        if (number(t.get('retention_pct')) or 0)>0:need+=['retention_release']
        for k in need:
            if number(t.get(k)) is None:out.append((f'{side}.{k}',f'{label}: nhập {TERM_NAMES[k]} (trống không được coi là 0).'))
        for k in ('advance_period','retention_release'):
            v=number(t.get(k))
            if v is not None and (v!=v.to_integral() or not 1<=v<=MAX_PERIODS*2):out.append((f'{side}.{k}',f'{label}: {TERM_NAMES[k]} là số kỳ nguyên ≥ 1.'))
        if number(t.get('delay')) is not None and number(t['delay'])!=number(t['delay']).to_integral():out.append((f'{side}.delay',f'{label}: số kỳ chậm là số nguyên.'))
        for k in ('advance_pct','recovery_pct','retention_pct'):
            v=number(t.get(k))
            if v is not None and v>HUNDRED:out.append((f'{side}.{k}',f'{label}: {TERM_NAMES[k]} không quá 100.'))
        r=number(t.get('retention_release'))
        if r is not None and r<n and (number(t.get('retention_pct')) or 0)>0:out.append((f'{side}.retention_release',f'{label}: kỳ hoàn trả giữ lại phải ≥ kỳ cuối ({n}).'))
        if not meaningful(t.get('ref')):out.append((f'{side}.ref',f'{label}: ghi căn cứ (điều khoản hợp đồng / dự thảo HĐ NTP).'))
    d=number((cf.get('own') or {}).get('delay'))
    if d is None or d!=d.to_integral():out.append(('own.delay','chi phí CC1: nhập số kỳ chậm thanh toán (0 nếu trả ngay).'))
    if not meaningful((cf.get('own') or {}).get('ref')):out.append(('own.ref','chi phí CC1: ghi căn cứ thời hạn thanh toán.'))
    return out
def _label(start,i):
    if not start:return f'T{i}'
    y,m=int(start[:4]),int(start[5:])-1+i-1;return f'{y+m//12:04d}-{m%12+1:02d}'
def _stream(total,pcts,t):
    """Contract stream (AB receipts or NTP payments): earned per period, advance, recovery, retention, delay → cash per period."""
    cash={};earned={};adv=total*number(t['advance_pct'])/HUNDRED;left=adv;ret_all=Decimal(0)
    if adv:cash[int(number(t['advance_period']))]=cash.get(int(number(t['advance_period'])),Decimal(0))+adv
    rec_pct=number(t.get('recovery_pct')) or Decimal(0);ret_pct=number(t['retention_pct']);delay=int(number(t['delay']))
    for i,p in enumerate(pcts,1):
        e=total*number(p)/HUNDRED;earned[i]=e
        if not e:continue
        rec=min(e*rec_pct/HUNDRED,left);left-=rec;ret=e*ret_pct/HUNDRED;ret_all+=ret
        cash[i+delay]=cash.get(i+delay,Decimal(0))+e-rec-ret
    if left>0:  # advance not recovered by the schedule: settled with the last payment (shown as a warning)
        last=max(cash);cash[last]-=left
    if ret_all:r=int(number(t['retention_release']));cash[r]=cash.get(r,Decimal(0))+ret_all
    return cash,earned,left
def compute_cashflow(s,scen,res):
    """Cash flow draft of a scenario. Profit (BM01) and cash are separate: an advance is not revenue and its recovery is not a cost."""
    cf=scen.get('cashflow')
    if not cf:return None
    places=_places(s);miss=[m[1] for m in compute_cashflow_inputs(cf)];warn=[]
    tot=res['totals'];B=number(tot['B']);C=number(tot['C'])
    ntp_v=[r['value'] for r in res['rows'] if r['code'] in ('C.1.5','C.1.6')]
    NTP=None if any(v is None for v in ntp_v) else sum((number(v) for v in ntp_v),Decimal(0))
    if B is None:miss.append('doanh thu B chưa đủ (xem BM01).')
    if C is None or NTP is None:miss.append('tổng chi phí C hoặc chi phí NTP chưa đủ (xem BM01).')
    out={'inputs':cf,'complete':False,'missing':miss,'warnings':warn,'not_modelled':NOT_MODELLED,'periods':[],'totals':None}
    if miss:return out
    OWN=C-NTP
    ab,earned,left_ab=_stream(B,cf['revenue_pct'],cf['ab']);ntp,_,left_ntp=_stream(NTP,cf['ntp_pct'],cf['ntp'])
    if left_ab>0:warn.append(f'Tạm ứng AB chưa thu hồi hết theo tỷ lệ thu hồi ({_r(left_ab,places)}); trừ vào lần thanh toán cuối.')
    if left_ntp>0:warn.append(f'Tạm ứng NTP chưa thu hồi hết theo tỷ lệ thu hồi ({_r(left_ntp,places)}); trừ vào lần thanh toán cuối.')
    d=int(number(cf['own']['delay']));own={}
    for i,p in enumerate(cf['own_pct'],1):
        v=OWN*number(p)/HUNDRED
        if v:own[i+d]=own.get(i+d,Decimal(0))+v
    last=max([cf['periods'],*ab,*ntp,*own]);cum=Decimal(0);rows=[];low=Decimal(0);low_at=None;fin=Decimal(0);rate=number(cf.get('rate_pct'))
    for i in range(1,last+1):
        a=ab.get(i,Decimal(0));n=ntp.get(i,Decimal(0));o=own.get(i,Decimal(0));net=a-n-o;cum+=net
        if cum<low:low=cum;low_at=i
        if cum<0 and rate is not None:fin+=-cum*rate/HUNDRED/12
        rows.append({'period':i,'label':_label(cf['start'],i),'earned':_r(earned.get(i,Decimal(0)),places),'ab_in':_r(a,places),'ntp_out':_r(n,places),'own_out':_r(o,places),'net':_r(net,places),'cum':_r(cum,places)})
    if rate is not None and not meaningful(cf.get('rate_ref')):rate=None;miss.append('ghi căn cứ lãi suất vốn (quy tắc huy động vốn) để tính chi phí tài chính.')
    out.update(periods=rows,complete=not miss,totals={'in':_r(sum(ab.values(),Decimal(0)),places),'out':_r(sum(ntp.values(),Decimal(0))+sum(own.values(),Decimal(0)),places),
               'net':_r(cum,places),'peak_funding':_r(-low,places),'peak_period':_label(cf['start'],low_at) if low_at else None,
               'financing':_r(fin,places) if rate is not None else None,'financing_note':'Lãi trên số dư âm lũy kế theo tháng; chỉ để đối chiếu C.3, không tự cộng vào BM01.'})
    if rate is None and not any('lãi suất' in m for m in miss):out['missing'].append('nhập lãi suất vốn năm (%) kèm căn cứ để tính chi phí tài chính (không bắt buộc cho dòng tiền).')
    out['complete']=not [m for m in out['missing'] if 'không bắt buộc' not in m]
    return out
# ---------- calculation ----------
def compute(s,scen,summary=None,costs=None):
    """BM01 rows, per-package direct margin and missing list for one scenario. Decimal throughout; null when incomplete."""
    summary=summary or sc.summary(s);costs=costs if costs is not None else sc.line_costs(s)
    places=_places(s);ab=sc.measured(s);basis_vals={x['code']:x['basis'] for x in summary['packages']}
    groups={k:Decimal(0) for k in COMPUTED};missing=[];pk_rows=[]
    def miss(code,text):missing.append({'row':code,'text':text})
    # B — revenue from the AB contract lines
    B=Decimal(0);b_ok=True
    for i,a in ab.items():
        q=number(a.get('quantity'));up=number(a.get('unit_price'))
        if q is None or up is None:b_ok=False;miss('B',f'Dòng AB {a.get("source_code") or i}: thiếu '+('khối lượng' if q is None else 'đơn giá hợp đồng'));continue
        B+=q*up
    cov=sc.coverage(s);wg=work_groups(s);quotes=s.get('quotes') or [];sumpk={x['code']:x for x in summary['packages']}
    for p in sc.packages(s):
        d=p.get('delivery') or {};mode=d.get('mode');b=(scen.get('basis') or {}).get(p['id']);bv=basis_vals.get(p['code']) or {}
        plan=(scen.get('parts') or {}).get(p['id']) if not p.get('is_self') else None;plan=plan or {}
        row={'id':p['id'],'code':p['code'],'name':p.get('name'),'is_self':bool(p.get('is_self')),'mode':mode,'basis':None if p.get('is_self') else b,
             'vendor':(p.get('award') or {}).get('vendor'),'ab_value':Decimal(0),'ab_ok':True,'cost':Decimal(0),'complete':True,'lines':len(p.get('lines') or [])}
        def pmiss(text,code):row['complete']=False;miss(code,f'{p["code"]}: {text}')
        parts={MAIN:{'id':MAIN,'name':'Phần chung','mode':mode,'basis':b,'d':d,'bv':bv,'lines':0,'cost':Decimal(0),'complete':True,'checked':False}}
        for x in plan.get('parts') or []:
            dd=dict(d,mode=x.get('mode'),**{k:x.get(k) for k in sc.PARTS});bb=x.get('basis')
            if bb=='B':  # basis B depends on who carries VL/NC/M — recompute it for this part's arrangement
                qs=[q for q in quotes if q['package_id']==p['id']];rv=(sumpk.get(p['code']) or {}).get('reviews') or []
                bvv=sc.basis_values(s,dict(p,delivery=dd),costs,[q for q,r in zip(qs,rv) if r['valid']])
            else:bvv=bv
            parts[x['id']]={'id':x['id'],'name':x['name'],'mode':x.get('mode'),'basis':bb,'d':dd,'bv':bvv,'lines':0,'cost':Decimal(0),'complete':True,'checked':False}
        def check(pt):
            """Missing arrangement of a part, reported once; returns False when its lines cannot be priced."""
            tag='' if pt['id']==MAIN else f'phần "{pt["name"]}": '
            ok=True
            if pt['mode'] not in sc.MODES:ok=False;(not pt['checked']) and pmiss(tag+'chưa chọn hình thức giao thầu GỘP/TÁCH','C.1.5')
            elif pt['basis'] not in BASIS_OPTIONS:ok=False;(not pt['checked']) and pmiss(tag+'chưa chọn căn cứ giá giao cho PAKD','C.1.5' if pt['mode']=='GOP' else 'C.1.6')
            if not ok:pt['complete']=False
            pt['checked']=True;return ok
        if not p.get('is_self') and not p.get('lines'):check(parts[MAIN])
        for l in p.get('lines') or []:
            pt=parts[MAIN] if p.get('is_self') else parts.get(_part_id(plan.get('by_group') or {},plan.get('by_line') or {},l['bb_code'],wg.get(l['ab_line_id'],(NOGROUP,))[0]),parts[MAIN])
            pt['lines']+=1;tag='' if pt['id']==MAIN else f'phần "{pt["name"]}" '
            def lmiss(text,code,pt=pt):pt['complete']=False;pmiss(text,code)
            a=ab.get(l['ab_line_id']);al=number(l.get('alloc_pct'))
            if not a or al is None or al<=0:lmiss(f'dòng {l["bb_code"]} không còn dòng AB hoặc % phân bổ không hợp lệ','C.1.7');continue
            q=number(a.get('quantity'));up=number(a.get('unit_price'));share=al/HUNDRED;c=costs.get(l['ab_line_id'])
            if q is None:lmiss(f'dòng {l["bb_code"]}: dòng AB thiếu khối lượng','C.1.7');continue
            if up is not None:row['ab_value']+=q*up*share
            else:row['ab_ok']=False
            def est(part,code,pt=pt,l=l,c=c,share=share):
                if not c or not c.get('complete'):lmiss(f'dòng {l["bb_code"]}: dự toán CC1 của dòng AB chưa đủ (bước 4–6) cho phần {part}',code,pt);return None
                v=c[part]*share;groups[code]+=v;row['cost']+=v;pt['cost']+=v;return v
            if p.get('is_self'):
                est('VL','C.1.4')
                for x in ('NC','M','other'):est(x,'C.1.6b')
                continue
            if not check(pt):continue
            dd=pt['d'];md=pt['mode'];bb=pt['basis'];code='C.1.5' if md=='GOP' else 'C.1.6'
            if md=='TACH':
                undecided=[x for x in sc.PARTS if dd.get(x) not in sc.PROVIDERS]
                if undecided:lmiss(tag+'chưa chọn bên chịu '+', '.join(undecided),'C.1.6');continue
                ntp=[x for x in sc.PARTS if dd.get(x)=='NTP']
                for x in sc.PARTS:
                    if dd.get(x)=='CC1':est(x,'C.1.4' if x=='VL' else 'C.1.6b')
                est('other','C.1.6b')
                if not ntp:continue
            if bb=='AWARD':
                price=None if sc.award_stale(s,p) else number(((p.get('award') or {}).get('prices') or {}).get(l['bb_code']))
                why=('đề xuất giao thầu đã cũ — lưu lại đề xuất' if sc.award_stale(s,p) else 'chưa có giá đề xuất giao thầu cho dòng') if price is None else None
            else:
                x=(pt['bv'].get(l['bb_code']) or {}).get(bb) or {};price=number(x.get('value'));why=f'căn cứ {bb}: {x.get("reason") or "chưa có giá"}' if price is None else None
            if price is None:lmiss(f'dòng {l["bb_code"]} — {tag}{why}',code);continue
            v=price*q*share;groups[code]+=v;row['cost']+=v;pt['cost']+=v
        if plan.get('parts'):
            row['parts']=[{'id':pt['id'],'name':pt['name'],'mode':pt['mode'],'basis':pt['basis'],'lines':pt['lines'],
                           'cost':_r(pt['cost'],places) if pt['complete'] else None,'known_cost':_r(pt['cost'],places)} for pt in parts.values() if pt['lines'] or pt['id']!=MAIN]
        pk_rows.append(row)
    # C.1.7 — the unallocated share of every AB line, at CC1 estimate cost
    for i,cv in cov.items():
        a=ab[i]
        if cv['state'] in ('OVER','INVALID'):miss('C.1.7',f'Dòng AB {a.get("source_code") or i}: phân bổ gói '+('vượt 100%' if cv['state']=='OVER' else 'không hợp lệ'));continue
        rest=HUNDRED-cv['allocated']
        if rest<=0:continue
        c=costs.get(i)
        if not c or not c.get('complete'):miss('C.1.7',f'Dòng AB {a.get("source_code") or i}: còn {s_(rest)}% chưa giao thầu nhưng dự toán CC1 chưa đủ');continue
        groups['C.1.7']+=sum((c[x] for x in ('VL','NC','M','other')),Decimal(0))*rest/HUNDRED
    rowmiss={m['row'] for m in missing}
    val={k:(None if k in rowmiss else v) for k,v in groups.items()};known=dict(groups)
    DT=B if b_ok else None
    user={k:[l for l in scen.get('lines') or [] if l['code']==k and l.get('on',True)] for k in USER_CODES}
    detail=[];rawamt={}
    def user_value(code,bases):
        tot=Decimal(0);ok=True
        for l in user[code]:
            v=number(l.get('value'))
            if v is None or not meaningful(l.get('source_ref')):ok=False;miss(code,f'{code} {l["name"]}: thiếu giá trị hoặc căn cứ');detail.append(dict(l,amount=None));continue
            if l['method']=='AMOUNT':amt=v
            else:
                bv=bases.get(l['base'])
                if bv is None:ok=False;miss(code,f'{code} {l["name"]}: cơ sở {BASE_NAMES.get(l["base"],l["base"])} chưa xác định');detail.append(dict(l,amount=None));continue
                amt=bv*v/HUNDRED
            tot+=amt;rawamt[l['id']]=amt;detail.append(dict(l,amount=_r(amt,places)))
        return tot if ok else None,tot
    for k in ('C.1.1','C.1.2','C.1.3'):val[k],known[k]=user_value(k,{})
    def add(keys):
        vals=[val[k] for k in keys]
        return (None if any(v is None for v in vals) else sum(vals,Decimal(0))),sum((known[k] for k in keys),Decimal(0))
    CPXD,CPXDk=add(['C.1.4','C.1.5','C.1.6','C.1.6b','C.1.7'])
    for k in ('C.1.8','C.1.9','C.1.10'):val[k],known[k]=user_value(k,{'DT':DT,'CPXD':CPXD})
    C1,C1k=add(ORDER)
    for k in ('C.2','C.3'):val[k],known[k]=user_value(k,{'DT':DT,'CPXD':CPXD,'C1':C1})
    for k in ('C.4','C.5'):val[k],known[k]=user_value(k,{'DT':DT,'C1':C1})
    C,Ck=add(ORDER+['C.2','C.3','C.4','C.5'])
    D=DT-C if DT is not None and C is not None else None
    margin=s_((D/DT*HUNDRED).quantize(Decimal('0.01'),ROUND_HALF_UP)) if D is not None and DT else None
    def r(code,name,v,kn=None,level=1,param=None):
        return {'code':code,'name':name,'value':_r(v,places),'known':_r(kn,places) if kn is not None else None,'level':level,'param':param,
                'pct_dt':s_((v/DT*HUNDRED).quantize(Decimal('0.01'),ROUND_HALF_UP)) if v is not None and DT else None}
    def param(code):
        ls=user.get(code) or []
        return '; '.join((f'{s_(number(l["value"]))}% × {l["base"]}' if l['method']=='PERCENT' and number(l.get('value')) is not None else 'số tiền')+(f' ({l["name"]})' if l['name']!=USER_CODES[code] else '') for l in ls) or None
    rows=[r('B','Doanh thu (giá trị HĐ A-B trước VAT)',DT,B,0,'Cố định theo BOQ AB'),r('C','Tổng chi phí',C,Ck,0),r('C.1','Chi phí trực tiếp',C1,C1k,1)]
    for k in ORDER:
        rows.append(r(k,USER_CODES.get(k) or COMPUTED[k],val[k],known[k],2,param(k) if k in USER_CODES else None))
    for k in ('C.2','C.3','C.4','C.5'):rows.append(r(k,USER_CODES[k],val[k],known[k],1,param(k)))
    rows.append(r('D','Lợi nhuận gộp (B − C)',D,None,0));rows[-1]['margin_pct']=margin
    for p in pk_rows:
        abv=p['ab_value'] if p.pop('ab_ok') else None;cost=p['cost'] if p['complete'] else None
        p.update(ab_value=_r(abv,places),cost=_r(cost,places),known_cost=_r(p['cost'],places) if not p['complete'] else None,
                 direct_margin=_r(abv-cost,places) if cost is not None and abv is not None else None,
                 direct_margin_pct=s_(((abv-cost)/abv*HUNDRED).quantize(Decimal('0.1'),ROUND_HALF_UP)) if cost is not None and abv else None)
    process=process_view(DT,val,known,detail,places,rawamt)
    return {'rows':rows,'packages':pk_rows,'missing':missing,'complete':not missing,'margin_pct':margin,'lines':detail,'note':NOTE,'process':process,
            'totals':{'B':_r(DT,places),'C':_r(C,places),'D':_r(D,places),'CPXD':_r(CPXD,places),'C1':_r(C1,places)}}
def process_view(DT,val,known,detail,places,rawamt=None):
    """Quy chiếu BM01 → cấu trúc A–I của Quy trình PAKD v1.0. H chỉ có khi mọi khoản đã được phân nhóm và đủ số; I chưa có (theo dõi thực hiện)."""
    g={k:Decimal(0) for k in 'BCDEF'};gk={k:Decimal(0) for k in 'BCDEF'};ok={k:True for k in 'BCDEF'};untagged=[];unt=Decimal(0)
    for code in ('C.1.4','C.1.5','C.1.6','C.1.6b','C.1.7'):
        if val.get(code) is None:ok['B']=False
        else:g['B']+=val[code]
        gk['B']+=known.get(code) or Decimal(0)
    for l in detail:
        tag=l.get('process_group') or PROCESS_DEFAULT.get(l['code']);amt=(rawamt or {}).get(l.get('id'));amt=number(l.get('amount')) if amt is None else amt  # unrounded, so H equals D of BM01
        if tag is None:untagged.append({'code':l['code'],'name':l['name'],'amount':l.get('amount')});unt+=amt or Decimal(0);continue
        if amt is None:ok[tag]=False
        else:g[tag]+=amt;gk[tag]+=amt
    G=None if val.get('C.1.5') is None or val.get('C.1.6') is None else val['C.1.5']+val['C.1.6']
    allk=all(ok.values()) and DT is not None and not untagged
    H=DT-sum(g.values()) if allk else None
    rows=[{'group':'A','name':PROCESS_GROUPS['A'],'value':_r(DT,places),'known':None,'source':'B của BM01'}]
    src={'B':'C.1.4–C.1.7 + khoản gắn nhóm B','C':'C.2 (mặc định) + khoản gắn nhóm C','D':'khoản gắn nhóm D','E':'C.3 (mặc định) + khoản gắn nhóm E','F':'C.1.8, C.1.10 (mặc định) + khoản gắn nhóm F'}
    for k in 'BCDEF':rows.append({'group':k,'name':PROCESS_GROUPS[k],'value':_r(g[k],places) if ok[k] else None,'known':_r(gk[k],places) if not ok[k] else None,'source':src[k]})
    rows.append({'group':'G','name':PROCESS_GROUPS['G'],'value':_r(G,places),'known':None,'source':'C.1.5 + C.1.6 (thông tin, đã nằm trong B)'})
    rows.append({'group':'H','name':PROCESS_GROUPS['H'],'value':_r(H,places),'known':None,'source':'= D của BM01 khi mọi khoản đã phân nhóm','margin_pct':s_((H/DT*HUNDRED).quantize(Decimal('0.01'),ROUND_HALF_UP)) if H is not None and DT else None})
    rows.append({'group':'I','name':PROCESS_GROUPS['I'],'value':None,'known':None,'source':'Chưa có: cần actual / ETC từ theo dõi thực hiện (bước 18)'})
    return {'rows':rows,'untagged':untagged,'untagged_amount':_r(unt,places),'complete':allk,'groups':PROCESS_GROUPS,'choices':list(PROCESS_CHOICES),'defaults':PROCESS_DEFAULT}
# ---------- digest, versions ----------
def inputs_digest(s,scen):
    """sha256 of everything that feeds the PAKD of a scenario (AB lines, packages, quotes, estimate, scenario inputs)."""
    calc_lines=[{k:l.get(k) for k in ('source_line_id','cbs','amount','line_type','percent_of')} for l in sc.calculate(s)['lines']]
    payload={'boq':[{k:r.get(k) for k in ('id','kind','quantity','uom','unit_price')} for r in s['boq']],
             'packages':[{k:p.get(k) for k in ('id','code','delivery','pricing','lines')}|{'award':(p.get('award') or {}).get('prices')} for p in sc.packages(s)],
             'quotes':sorted(q.get('sha256') or q['id'] for q in s.get('quotes') or []),'estimate':calc_lines,'money_places':_places(s),
             'scenario':{'basis':scen.get('basis') or {},'lines':scen.get('lines') or [],'parts':scen.get('parts') or {},'cashflow':scen.get('cashflow')}}
    return hashlib.sha256(json.dumps(payload,sort_keys=True,default=str,ensure_ascii=False).encode()).hexdigest()
def view(s,summary=None):
    st=state(s);summary=summary or sc.summary(s);costs=sc.line_costs(s);out=[]
    for c in st['scenarios']:
        res=compute(s,c,summary,costs)
        out.append({'id':c['id'],'name':c['name'],'basis':c.get('basis') or {},'lines':c.get('lines') or [],'parts':c.get('parts') or {},'cashflow_inputs':c.get('cashflow'),
                    'result':res,'cashflow':compute_cashflow(s,c,res),'questions':questions(s,c),'digest':inputs_digest(s,c)})
    digests={c['id']:c['digest'] for c in out}
    vers=[dict(v,stale=digests.get(v['scenario_id'])!=v['inputs_sha256']) for v in st['versions']]
    wg=work_groups(s);grp={p['id']:package_groups(s,p,wg) for p in sc.packages(s) if not p.get('is_self')}
    return {'scenarios':out,'current':st['current'],'versions':vers,'codes':USER_CODES,'process_groups':PROCESS_GROUPS,'process_choices':list(PROCESS_CHOICES),'process_defaults':PROCESS_DEFAULT,'groups':grp,'modes':{'GOP':'Giao gộp','TACH':'Giao tách'},'not_modelled':NOT_MODELLED,'bases':{k:list(v) for k,v in BASES.items()},
            'base_names':BASE_NAMES,'basis_options':BASIS_OPTIONS,'note':NOTE}
def act(s,body):
    """Mutations: save | add | copy | delete | select | freeze. Returns the changed project (caller stores it)."""
    st=state(s);a=body.get('action');p=copy.deepcopy(st)
    if a=='save':
        c=_scen(p,body.get('scenario_id'))
        if 'name' in body:
            n=str(body.get('name') or '').strip()
            if not n:raise ValueError('Tên kịch bản không được trống.')
            c['name']=n[:120]
        if 'basis' in body:
            b=body.get('basis') or {}
            if not isinstance(b,dict) or any(v not in BASIS_OPTIONS and v not in (None,'') for v in b.values()):raise ValueError('Căn cứ giá giao phải là Giá đề xuất giao thầu (AWARD), Căn cứ A, B hoặc C.')
            ids={x['id'] for x in sc.packages(s) if not x.get('is_self')}
            c['basis']={k:v for k,v in b.items() if k in ids and v}
        if 'lines' in body:c['lines']=check_lines(body.get('lines'))
        if 'parts' in body:c['parts']=check_parts(s,body.get('parts'))
        if 'cashflow' in body:c['cashflow']=check_cashflow(body.get('cashflow'))
    elif a in ('add','copy'):
        n=str(body.get('name') or '').strip()
        if not n:raise ValueError('Nhập tên kịch bản.')
        if any(c['name'].lower()==n.lower() for c in p['scenarios']):raise ValueError('Đã có kịch bản cùng tên.')
        src=_scen(p,body.get('scenario_id')) if a=='copy' else {'basis':{},'lines':[]}
        new={'id':str(uuid.uuid4()),'name':n[:120],'basis':copy.deepcopy(src.get('basis') or {}),'lines':[dict(l,id=str(uuid.uuid4())) for l in src.get('lines') or []],
             'parts':copy.deepcopy(src.get('parts') or {}),'cashflow':copy.deepcopy(src.get('cashflow'))}
        p['scenarios'].append(new);p['current']=new['id']
    elif a=='delete':
        c=_scen(p,body.get('scenario_id'))
        if len(p['scenarios'])==1:raise ValueError('Cần giữ ít nhất một kịch bản.')
        p['scenarios']=[x for x in p['scenarios'] if x['id']!=c['id']]
        if p['current']==c['id']:p['current']=p['scenarios'][0]['id']
    elif a=='select':p['current']=_scen(p,body.get('scenario_id'))['id']
    elif a=='freeze':
        c=_scen(p,body.get('scenario_id'));name=str(body.get('name') or '').strip();by=str(body.get('by') or '').strip()
        if not name or not by:raise ValueError('Nhập tên phiên bản và người chốt.')
        if any(v['name'].lower()==name.lower() for v in p['versions']):raise ValueError('Đã có phiên bản cùng tên; phiên bản đã chốt không ghi đè — đặt tên mới.')
        qs=questions(s,c);res=compute(s,c)
        if qs or not res['complete']:
            raise ValueError('Chưa chốt được: PAKD còn thiếu '+str(len(res['missing'])+len(qs))+' mục (xem danh sách "Còn thiếu").')
        cfr=compute_cashflow(s,c,res)
        if not cfr or not cfr.get('complete'):  # D86: Quy trình PAKD v1.0 §5/§10 — baseline phải có dòng tiền và kế hoạch tháng
            raise ValueError('Chưa chốt được: PAKD baseline cần dòng tiền / kế hoạch tháng đủ dữ liệu (Quy trình PAKD Thi công v1.0, mục 5 và 10) — lập dòng tiền trong kịch bản rồi chốt.')
        body_rows=json.dumps({'rows':res['rows'],'packages':res['packages'],'lines':res['lines'],'cashflow':cfr},sort_keys=True,ensure_ascii=False,default=str)
        p['versions'].append({'id':str(uuid.uuid4()),'name':name[:120],'scenario_id':c['id'],'scenario_name':c['name'],'at':datetime.now().isoformat(timespec='seconds'),
                              'by':by[:200],'note':str(body.get('note') or '')[:1000],'project_revision':s['revision'],'inputs_sha256':inputs_digest(s,c),
                              'result_sha256':hashlib.sha256(body_rows.encode()).hexdigest(),'rows':res['rows'],'packages':res['packages'],'lines':res['lines'],'cashflow':cfr,'process':res.get('process'),
                              'status':'FROZEN_LOCAL','disclaimer':'Phiên bản chốt trên máy: bất biến, không phải phê duyệt PAKD.'})
    else:raise ValueError('Thao tác phải là lưu (save), thêm (add), sao chép (copy), xóa (delete), chọn (select) hoặc chốt (freeze).')
    s['pakd']=p;return s
# ---------- Excel ----------
def workbook(s,scenario_id=None,version_id=None):
    """xlsx bytes: BM01, packages, user lines, missing — of a scenario (live) or a frozen version (as frozen)."""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    st=state(s)
    if version_id:
        v=next((x for x in st['versions'] if x['id']==version_id),None)
        if not v:raise KeyError('Không có phiên bản này.')
        title=f'Phiên bản chốt: {v["name"]} · {v["at"]} · {v["by"]} · inputs {v["inputs_sha256"][:16]}';res={'rows':v['rows'],'packages':v['packages'],'lines':v['lines'],'missing':[],'process':v.get('process')};cfr=v.get('cashflow')
    else:
        c=_scen(st,scenario_id or st['current']);res=compute(s,c);cfr=compute_cashflow(s,c,res);title=f'Kịch bản: {c["name"]} · revision hồ sơ {s["revision"]} · '+('ĐỦ' if res['complete'] else f'THIẾU {len(res["missing"])} mục')
    wb=Workbook();hdr=PatternFill('solid',fgColor='106B3C');bold=Font(bold=True);white=Font(bold=True,color='FFFFFF');grp=PatternFill('solid',fgColor='E7F1EB')
    def sheet(ws,name,head,rows,widths):
        ws.title=name;ws.append([NOTE]);ws.append([title]);ws.append([s['context'].get('project_code') or '',s['context'].get('contract_ref') or ''])
        ws.append(head)
        for c in ws[4]:c.fill=hdr;c.font=white;c.alignment=Alignment(wrap_text=True,vertical='center')
        for r in rows:ws.append(r)
        for row in ws.iter_rows(min_row=5):
            for c in row:
                if isinstance(c.value,str) and c.value.startswith(('=','+','-','@')):c.data_type='s'  # formula injection
        for i,w in enumerate(widths):ws.column_dimensions[chr(65+i)].width=w
        ws.freeze_panes='A5'
    num=lambda v:float(v) if v not in (None,'') else None
    ws=wb.active;sheet(ws,'BM01_TH',['STT','Khoản mục','Thông số','Giá trị (VND)','Đã biết khi còn thiếu','% doanh thu'],
        [[r['code'],r['name'],r.get('param') or '',num(r['value']),num(r.get('known')) if r['value'] is None else None,num(r.get('pct_dt'))] for r in res['rows']],[9,44,40,20,20,12])
    for row in ws.iter_rows(min_row=5):
        if row[0].value in ('B','C','D','C.1'):
            for c in row:c.fill=grp;c.font=bold
        for c in row[3:5]:c.number_format='#,##0'
    d=next((r for r in res['rows'] if r['code']=='D'),{})
    ws.append([]);ws.append(['','Tỷ suất lợi nhuận gộp (%)','',num(d.get('margin_pct'))])
    sheet(wb.create_sheet(),'GOI_THAU',['Gói','Tên','Hình thức','Căn cứ giá giao','Nhà thầu đề xuất','Giá trị AB phần gói','Chi phí gói','LN trực tiếp','LN trực tiếp %'],
        [[p['code'],p['name'],'CC1 tự làm' if p['is_self'] else (p.get('mode') or ''),BASIS_OPTIONS.get(p.get('basis'),'') if not p['is_self'] else 'Dự toán CC1',p.get('vendor') or '',
          num(p['ab_value']),num(p['cost']),num(p['direct_margin']),num(p['direct_margin_pct'])] for p in res['packages']]+
        [[p['code'],'   Phần: '+x['name']+f' · {x["lines"]} dòng',{'GOP':'GỘP','TACH':'TÁCH'}.get(x.get('mode'),'chưa chọn'),BASIS_OPTIONS.get(x.get('basis'),''),'',None,num(x['cost']),None,None]
         for p in res['packages'] for x in p.get('parts') or []],[9,30,12,30,22,18,18,18,12])
    sheet(wb.create_sheet(),'DONG_NHAP',['Mã','Tên khoản','Phương pháp','Cơ sở','Giá trị nhập','Thành tiền','Căn cứ','Đang tính'],
        [[l['code'],l['name'],l['method'],l.get('base') or '',l.get('value'),num(l.get('amount')),l.get('source_ref'),'Có' if l.get('on',True) else 'Không'] for l in res['lines']],[9,30,12,10,14,18,50,10])
    if cfr and cfr.get('periods'):
        cfrows=[[r['label'],num(r['earned']),num(r['ab_in']),num(r['ntp_out']),num(r['own_out']),num(r['net']),num(r['cum'])] for r in cfr['periods']]
        tt=cfr['totals'];cfrows+=[[],['Tổng thu',None,num(tt['in'])],['Tổng chi',None,None,num(tt['out'])],['Nhu cầu vốn đỉnh',None,None,None,None,None,num(tt['peak_funding'])],
                                  ['Kỳ vốn đỉnh',tt['peak_period']],['Chi phí tài chính ước tính',num(tt['financing']),tt['financing_note']]]
    else:cfrows=[['—','Chưa lập dòng tiền' if not cfr else 'Dòng tiền chưa đủ dữ liệu: '+'; '.join(cfr['missing'])]]
    if cfr:cfrows+=[[],[cfr['not_modelled']]]+[['Cảnh báo',w] for w in cfr.get('warnings') or []]+([['Còn thiếu',m] for m in cfr['missing']] if cfr.get('periods') else [])
    sheet(wb.create_sheet(),'DONG_TIEN',['Kỳ','Giá trị nghiệm thu AB (doanh thu)','Thu từ CĐT','Chi NTP','Chi CC1','Dòng tiền ròng','Lũy kế'],cfrows,[12,22,18,18,18,18,18])
    for row in wb['DONG_TIEN'].iter_rows(min_row=5):
        for c in row[1:7]:c.number_format='#,##0'
    pv=res.get('process') or {}
    sheet(wb.create_sheet(),'QUY_TRINH_A_I',['Nhóm','Khoản mục (Quy trình PAKD v1.0 §5)','Giá trị (VND)','Đã biết khi còn thiếu','Nguồn quy chiếu từ BM01'],
          [[r['group'],r['name'],num(r['value']),num(r.get('known')),r['source']] for r in pv.get('rows') or []]+([[],['Chưa phân nhóm',f"{len(pv['untagged'])} khoản · {pv['untagged_amount']}",None,None,'; '.join(f"{u['code']} {u['name']}" for u in pv['untagged'])]] if pv.get('untagged') else []),[8,48,20,20,60])
    sheet(wb.create_sheet(),'CON_THIEU',['Dòng BM01','Nội dung còn thiếu'],[[m['row'],m['text']] for m in res['missing']] or [['—','Không còn mục thiếu']],[12,100])
    b=io.BytesIO();wb.save(b);return b.getvalue()
