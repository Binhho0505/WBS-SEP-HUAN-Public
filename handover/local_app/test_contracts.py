"""E18 (D78) native tests: subcontract dossier on BB′ — PRD FR29–FR32, AC17–AC19 (+AC03, AC22, AC23). TEST ONLY fixtures.

Fixture (test_pakd.base): GT-01 = AB 1.1 200 m3 × 100% at award 95 + AB 1.2 10 t × 100% at award 40 → BB′ total 19,400.
"""
import base64,copy,io,json,tempfile,threading,unittest,urllib.request,urllib.error,zipfile
from pathlib import Path
from openpyxl import load_workbook
import subcontract as sc,contracts as ct,pakd
from server import App
from test_pakd import base
FILE=('TEST_ONLY_form.pdf',b'%PDF TEST ONLY form',None);SCAN=('TEST_ONLY_signed.pdf',b'%PDF TEST ONLY signed scan',None)
def f(x):import hashlib;return (x[0],x[1],hashlib.sha256(x[1]).hexdigest())
TERMS={k:{'text':f'TEST {n}','ref':'TEST căn cứ'} for k,n in ct.TERMS.items()}
B2B=[{'clause':'TEST Điều 5 thanh toán','ab_ref':'HĐ AB Đ5','decision':'ADAPT','rationale':'TEST kỳ thanh toán NTP sau AB 15 ngày'}]
SIGN={'parties':['TEST CC1','TEST NTP'],'contract_ref':'TEST-HĐ-01','date':'2026-10-21','by':'TEST QS','verification':'TEST đối chiếu bản gốc'}
def ready(s,by='TEST'):
    s=ct.act(s,{'action':'create','package_id':'id-GT-01','by':by});cid=s['contracts'][0]['id']
    s=ct.act(s,{'action':'save','contract_id':cid,'terms':TERMS,'back_to_back':B2B,'authority':{'name':'TEST GĐ','role':'TEST','ref':'TEST UQ-01'},'form_ref':'TEST mẫu','by':by},f(FILE))
    s=ct.act(s,{'action':'submit','contract_id':cid,'by':by});s=ct.act(s,{'action':'ready','contract_id':cid,'by':by});return s,cid
class ContractDomain(unittest.TestCase):
    def setUp(self):self.s=base()
    def c(self,s,cid):return ct._get(s,cid)
    def test_gates_and_state_machine(self):
        s=ct.act(self.s,{'action':'create','package_id':'id-GT-01','by':'TEST'});cid=s['contracts'][0]['id'];g=ct.gates(s,self.c(s,cid))
        self.assertEqual(len([x for x in g if x.startswith('Điều khoản')]),15);self.assertTrue(any('mẫu hợp đồng' in x for x in g) and any('back-to-back' in x.lower() for x in g) and any('thẩm quyền' in x for x in g))
        with self.assertRaises(ValueError):ct.act(s,{'action':'ready','contract_id':cid,'by':'TEST'})          # no skipping a state
        with self.assertRaises(ValueError):ct.act(s,{'action':'create','package_id':'id-GT-01','by':'TEST'})   # one open dossier per package
        with self.assertRaises(ValueError):ct.act(s,{'action':'create','package_id':'id-CC1','by':'TEST'})     # CC1 self package has no subcontract
        s=ct.act(s,{'action':'submit','contract_id':cid,'by':'TEST'})
        with self.assertRaises(ValueError) as e:ct.act(s,{'action':'ready','contract_id':cid,'by':'TEST'})
        self.assertIn('Chưa đủ điều kiện',str(e.exception))
        with self.assertRaises(ValueError):ct.act(s,{'action':'save','contract_id':cid,'terms':TERMS})          # locked outside draft
        s,cid=ready(self.s);c=self.c(s,cid)
        self.assertEqual((c['status'],c['snapshot']['total'],len(c['snapshot']['digest'])),('READY_FOR_AUTHORIZED_SIGNING','19400',64))
        self.assertEqual(c['form']['sha256'],f(FILE)[2])
    def test_ready_blocked_by_over_allocation_and_missing_price(self):  # AC03: no READY while an AB line is over-allocated
        s=copy.deepcopy(self.s);s['packages'][1]['lines'].append({'id':'lx','bb_code':'CC1.2','ab_line_id':'ab2','alloc_pct':'10'})
        s2,cid=None,None
        with self.assertRaises(ValueError) as e:ready(s)
        self.assertIn('vượt 100%',str(e.exception))
        s=copy.deepcopy(self.s);s['packages'][0]['award']['prices'].pop('GT-01.2')
        with self.assertRaises(ValueError) as e:ready(s)
        self.assertIn('GT-01.2: chưa có đơn giá giao',str(e.exception))
    def test_sign_evidence_and_lock(self):
        s,cid=ready(self.s)
        for bad in ({},dict(SIGN,parties=['TEST CC1']),dict(SIGN,date='21/10/2026'),dict(SIGN,verification='')):
            with self.assertRaises(ValueError):ct.act(s,dict(bad,action='sign',contract_id=cid),f(SCAN) if bad else None)
        s2=copy.deepcopy(s);s2=sc.save_award(s2,{'package_id':'id-GT-01','vendor':'TEST NTP','prices':{'GT-01.1':'96','GT-01.2':'40'}})
        self.assertTrue(ct.ready_stale(s2,self.c(s2,cid)))
        with self.assertRaises(ValueError):ct.act(s2,dict(SIGN,action='sign',contract_id=cid),f(SCAN))  # price changed after READY → back to review
        s=ct.act(s,dict(SIGN,action='sign',contract_id=cid),f(SCAN));c=self.c(s,cid)
        self.assertEqual((c['status'],c['signed']['sha256'],c['signed']['authority_verified']),('SIGNED_EVIDENCE_RECORDED',f(SCAN)[2],False))
        # FR30: the signed package cannot be edited through the package / award APIs
        new=sc.save_award(copy.deepcopy(s),{'package_id':'id-GT-01','vendor':'TEST NTP','prices':{'GT-01.1':'90','GT-01.2':'40'}})
        with self.assertRaises(ValueError):ct.guard(s,new)
        new=copy.deepcopy(s);new['packages'][0]['lines'][0]['alloc_pct']='50'
        with self.assertRaises(ValueError):ct.guard(s,new)
        new=copy.deepcopy(s);new['packages']=new['packages'][1:]
        with self.assertRaises(ValueError):ct.guard(s,new)
        with self.assertRaises(ValueError):ct.act(s,{'action':'return','contract_id':cid,'by':'TEST'})
    def test_ac17_signed_v1_unchanged_change_and_amendment(self):
        s,cid=ready(self.s);s=ct.act(s,dict(SIGN,action='sign',contract_id=cid),f(SCAN));v1=copy.deepcopy(self.c(s,cid)['snapshot'])
        next(r for r in s['boq'] if r['id']=='test-source')['quantity']='210'      # AB quantity changes after signing
        c=self.c(s,cid);self.assertEqual(ct.current_lines(c),v1['lines'])          # signed v1 untouched
        d=ct.drift(s,c);self.assertEqual([(x['bb_code'],x['signed_qty'],x['current_qty'],x['delta']) for x in d],[('GT-01.1','200','210','10')])
        s=ct.act(s,{'action':'change','contract_id':cid,'source':'AB','title':'TEST tăng KL 1.1','request_ref':'TEST CV CĐT','impact':{'quantities':[{'bb_code':'GT-01.1','delta':'10'}],'ntp_cost':'950','ab_value':'1000'},'by':'TEST'})
        ch=s['changes'][0];self.assertEqual((ch['ab_status'],ch['ntp_status'],ch['lineage']['contract_digest']),('NOT_SUBMITTED','NOT_ISSUED',v1['digest']))
        with self.assertRaises(ValueError):ct.act(s,{'action':'amend_draft','contract_id':cid,'change_ids':[],'reason':'x','lines':[{'bb_code':'GT-01.1','quantity':'210','ref':'x'}]})
        s=ct.act(s,{'action':'amend_draft','contract_id':cid,'change_ids':[ch['id']],'reason':'TEST phụ lục KL','lines':[{'bb_code':'GT-01.1','quantity':'210','ref':'TEST CV CĐT'}],'by':'TEST'})
        self.assertEqual(self.c(s,cid)['status'],'SIGNED_EVIDENCE_RECORDED')       # a draft amendment changes nothing in force
        s=ct.act(s,dict(SIGN,action='amend_record',contract_id=cid,contract_ref='TEST-PL-01'),f(('TEST_ONLY_pl1.pdf',b'%PDF TEST pl1',None)));c=self.c(s,cid)
        self.assertEqual((c['status'],ct.total(ct.current_lines(c)),c['snapshot']['total'],c['snapshot']['digest']),('AMENDED','20350','19400',v1['digest']))  # 95×210 + 400
        self.assertEqual(ct.drift(s,c),[])
        with self.assertRaises(ValueError):ct.act(s,{'action':'close','contract_id':cid,'by':'TEST'})  # closing needs a basis
        s=ct.act(s,{'action':'close','contract_id':cid,'by':'TEST','note':'TEST thanh lý'});self.assertEqual(self.c(s,cid)['status'],'CLOSED')
    def test_fr31_unfunded_exposure(self):
        s,cid=ready(self.s)
        for t,cost in (('TEST A','500'),('TEST B',None)):
            s=ct.act(s,{'action':'change','contract_id':cid,'source':'NTP','title':t,'request_ref':'TEST','impact':{'ntp_cost':cost},'by':'TEST'})
        a,b=s['changes']
        s=ct.act(s,{'action':'change','change_id':a['id'],'ntp_status':'AGREED','ab_status':'SUBMITTED','by':'TEST'})
        s=ct.act(s,{'action':'change','change_id':b['id'],'ntp_status':'ISSUED','by':'TEST'})
        e=ct.exposure(s);self.assertEqual((e['unfunded'],e['unknown'],len(e['rows'])),('500',1,2))   # unknown cost stays unknown, not 0
        s=ct.act(s,{'action':'change','change_id':a['id'],'ab_status':'APPROVED','by':'TEST'});self.assertEqual(ct.exposure(s)['unfunded'],'0')
        with self.assertRaises(ValueError):ct.act(s,{'action':'change','change_id':a['id'],'ab_status':'MAYBE'})
        self.assertEqual(len(s['changes'][0]['history']),3)
    def test_fr32_facts_ac18_ntp_done_is_not_ab_acceptance(self):
        s=sc.add_progress(copy.deepcopy(self.s),{'bb_code':'GT-01.1','cum_qty':'200','date':'2026-10-20','ref':'TEST NK'})
        r=pakd.compute(s,pakd._scen(pakd.state(s),pakd.state(s)['current']))
        t=ct.trace(s);l=next(x for x in t['lines'] if x['bb_code']=='GT-01.1')
        self.assertEqual((l['ntp_done'],l['ntp_accepted']),('200',None))
        self.assertIsNone(next(x for x in t['ab'] if x['ab_line_id']=='test-source')['ab_accepted'])   # AC18: AB acceptance stays unknown
        self.assertEqual(r['totals']['B'],pakd.compute(self.s,pakd._scen(pakd.state(self.s),pakd.state(self.s)['current']))['totals']['B'])  # revenue unchanged
        F=lambda **k:dict({'action':'fact','kind':'NTP_ACCEPTED','bb_code':'GT-01.1','basis':'CUMULATIVE','evidence':'TEST BB NT','by':'TEST'},**k)
        s=ct.act(s,F(period='2026-09',value='80'));s=ct.act(s,F(period='2026-10',value='150'))
        for bad in (F(period='2026-10',value='160'),F(period='2026-11',value='100'),F(period='2026-11',value='170',basis='DELTA'),F(period='2026-11',value=''),F(period='2026-13',value='1')):
            with self.assertRaises(ValueError):ct.act(s,bad)   # duplicate period, decreasing cumulative, mixed basis, empty, bad period
        self.assertEqual(ct.fact_totals(s)[('NTP_ACCEPTED','GT-01.1')],'150')
        last=s['facts'][-1]['id'];s=ct.act(s,{'action':'fact','reverses':last,'evidence':'TEST nhập sai','by':'TEST'})
        self.assertEqual(ct.fact_totals(s)[('NTP_ACCEPTED','GT-01.1')],'80')                       # reversal, not edit
        A=lambda **k:dict({'action':'fact','kind':'AB_ACCEPTED','ab_line_id':'test-source','basis':'DELTA','evidence':'TEST BB CĐT','by':'TEST'},**k)
        s=ct.act(s,A(period='2026-09',value='50'));s=ct.act(s,A(period='2026-10',value='30'))
        self.assertEqual(ct.fact_totals(s)[('AB_ACCEPTED','test-source')],'80')                     # deltas add up once
        self.assertEqual(next(x for x in ct.trace(s)['lines'] if x['bb_code']=='GT-01.1')['ntp_done'],'200')   # facts never touch NTP progress
        new=copy.deepcopy(s);new['packages'][0]['lines']=new['packages'][0]['lines'][1:]
        with self.assertRaises(ValueError):ct.progress_refs(new)                                     # no orphan progress
class DeleteDraftPackages(unittest.TestCase):  # D81
    def setUp(self):self.s=base()
    def test_delete_removes_quotes_award_and_pakd_refs(self):
        s=sc.add_quote(copy.deepcopy(self.s),{'package_id':'id-GT-01','vendor':'TEST NTP','filename':'q.csv'},'0'*64,[{'bb_code':'GT-01.1','unit_price':'90'}])
        sid=pakd.state(s)['current'];s=pakd.act(s,{'action':'save','scenario_id':sid,'basis':{'id-GT-01':'AWARD'},'parts':{'id-GT-01':{'parts':[{'id':'p1','name':'TEST','mode':'GOP','basis':'A'}],'by_line':{'GT-01.2':'p1'}}}})
        u=sc.package_usage(s)['id-GT-01'];self.assertEqual((u['lines'],u['quotes'],u['award'],u['pakd'],u['blocked_by']),(2,1,True,1,None))
        s,rep=sc.delete_packages(s,{'package_ids':['id-GT-01']})
        self.assertEqual([p['code'] for p in s['packages']],['CC1']);self.assertEqual(s['quotes'],[]);self.assertEqual(rep,[{'code':'GT-01','lines':2,'quotes':1,'award':True}])
        c=pakd.state(s)['scenarios'][0];self.assertEqual((c['basis'],c['parts']),({},{}))
        pakd.act(s,{'action':'save','scenario_id':sid,'parts':c['parts'],'basis':c['basis']})  # PAKD still saves (no dangling package id)
    def test_delete_refused_with_contract_or_progress(self):
        s=ct.act(self.s,{'action':'create','package_id':'id-GT-01','by':'TEST'})
        self.assertEqual(sc.package_usage(s)['id-GT-01']['blocked_by'],'hồ sơ hợp đồng')
        with self.assertRaises(ValueError):sc.delete_packages(s,{'package_id':'id-GT-01'})
        s=sc.add_progress(copy.deepcopy(self.s),{'bb_code':'CC1.1','cum_qty':'1','date':'2026-10-21','ref':'TEST'})
        with self.assertRaises(ValueError):sc.delete_packages(s,{'package_id':'id-CC1'})
        with self.assertRaises(ValueError):sc.delete_packages(self.s,{'package_ids':[]})
class ContractHttp(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.app=App(('127.0.0.1',0),Path(self.tmp.name)/'db.sqlite');threading.Thread(target=self.app.serve_forever,daemon=True).start()
        self.base=f'http://127.0.0.1:{self.app.server_port}';s=base();p=self.app.store.create('TEST HĐ');s.update(id=p['id'],revision=p['revision'],name=p['name'])
        self.s=self.app.store.put(s,p['revision'],'TEST_SEED');self.pid=self.s['id']
    def tearDown(self):self.app.shutdown();self.app.server_close();self.tmp.cleanup()
    def call(self,path,body=None,raw=False):
        req=urllib.request.Request(self.base+path,data=None if body is None else json.dumps(body).encode(),headers={} if body is None else {'Content-Type':'application/json','X-CC1-Session':self.app.token})
        with urllib.request.urlopen(req) as r:return r.read() if raw else json.load(r)
    def post(self,v,body,file=None):
        if file:body=dict(body,filename=file[0],data=base64.b64encode(file[1]).decode())
        return self.call(f'/api/projects/{self.pid}/contracts',dict(body,revision=v['revision']))
    def test_api_flow_lock_export_restore(self):
        v=self.call(f'/api/projects/{self.pid}/subcontract');v=self.post(v,{'action':'create','package_id':'id-GT-01','by':'TEST'});cid=v['contracts']['contracts'][0]['id']
        v=self.post(v,{'action':'save','contract_id':cid,'terms':TERMS,'back_to_back':B2B,'authority':{'name':'TEST','ref':'TEST UQ'},'by':'TEST'},FILE)
        for a in ('submit','ready'):v=self.post(v,{'action':a,'contract_id':cid,'by':'TEST'})
        c=v['contracts']['contracts'][0];self.assertEqual((c['status'],c['gates']),('READY_FOR_AUTHORIZED_SIGNING',[]))
        with self.assertRaises(urllib.error.HTTPError) as e:self.call(f'/api/projects/{self.pid}/contracts',{'revision':0,'action':'submit','contract_id':cid})
        self.assertEqual(e.exception.code,409)
        p=self.app.store.get(self.pid);r=self.call('/api/restore',{'schema_version':p.get('schema_version') or '0.2.0','project':p,'audit':[]})['project']
        self.assertEqual((r['contracts'][0]['status'],r['contracts'][0]['restored']),('IN_REVIEW',True))   # AC23: readiness is not restored
        v=self.post(v,dict(SIGN,action='sign',contract_id=cid),SCAN);self.assertEqual(v['contracts']['contracts'][0]['status'],'SIGNED_EVIDENCE_RECORDED')
        with self.assertRaises(urllib.error.HTTPError) as e:   # bypassing the UI: the award API refuses to change a signed package
            self.call(f'/api/projects/{self.pid}/award',{'revision':v['revision'],'package_id':'id-GT-01','vendor':'TEST NTP','prices':{'GT-01.1':'1','GT-01.2':'40'}})
        self.assertEqual(e.exception.code,400)
        wb=load_workbook(io.BytesIO(self.call(f'/api/projects/{self.pid}/contracts/{cid}.xlsx',raw=True)),read_only=True)
        self.assertEqual(wb.sheetnames,['BOQ_BB_PRIME','DIEU_KHOAN','BACK_TO_BACK_AB','BANG_CHUNG','LICH_SU'])
        rows=list(wb['BOQ_BB_PRIME'].iter_rows(values_only=True));self.assertIn('DỰ THẢO TRÊN MÁY',rows[0][0]);self.assertEqual(sum(r[10] for r in rows[3:]),19400)
        z=zipfile.ZipFile(io.BytesIO(self.call(f'/api/projects/{self.pid}/procurement-pack.zip',raw=True)))
        self.assertTrue(any(n.startswith('HOP_DONG/') for n in z.namelist()))
        ev=[e['action'] for e in self.app.store.events(self.pid)];self.assertIn('CONTRACT_SIGN',ev)
    def test_api_delete_draft_package(self):
        v=self.call(f'/api/projects/{self.pid}/subcontract');self.assertIn('id-GT-01',v['package_usage'])
        v=self.call(f'/api/projects/{self.pid}/packages/delete',{'revision':v['revision'],'package_ids':['id-GT-01']})
        self.assertEqual([p['code'] for p in v['packages']],['CC1']);self.assertEqual(v['result']['deleted'][0]['code'],'GT-01')
        self.assertIn('DELETE_DRAFT_PACKAGES',[e['action'] for e in self.app.store.events(self.pid)])
if __name__=='__main__':unittest.main(verbosity=1)
