"""D59 native tests for E10–E16 (TEST_DESIGN R1–R14). TEST ONLY fixtures: never CC1 prices, never approvals."""
import base64,copy,hashlib,io,json,tempfile,threading,unittest,urllib.request,urllib.error,zipfile
from decimal import Decimal
from pathlib import Path
from openpyxl import load_workbook
from domain import status
from store import Store
from server import App
import subcontract as sc,libtools,pmmodels,drawings,packs
from test_app import fixture

def multi():
    """TEST ONLY: 3 measured AB lines (1.1 priced in the CC1 estimate, 1.2 with AB price only, 2.1 without AB price) + a heading."""
    s=fixture()
    s['boq']=[dict(id='h1',source_code='1',description='TEST ONLY heading',kind='HEADING',source_ref='T:0',source_hash='b'*64)]+s['boq']+[
        dict(id='ab2',source_code='1.2',description='TEST ONLY rebar',quantity='10',uom='t',unit_price='50',kind='MEASURED',source_ref='T:2',source_hash='c'*64),
        dict(id='ab3',source_code='2.1',description='TEST ONLY excavation',quantity='100',uom='m3',kind='MEASURED',source_ref='T:3',source_hash='d'*64)]
    for i,ab in (('item2','ab2'),('item3','ab3')):
        b=next(r for r in s['boq'] if r['id']==ab)
        s['mapping'].append(dict(s['mapping'][0],id=i,source_line_id=ab,quantity=b['quantity'],uom=b['uom']))
    s['confirmed']=[0,1,2];return s
def pk(code,name='TEST gói',**kw):
    d=dict(id='id-'+code,code=code,name=name,lines=[],procurement_method=None,delivery={},pricing={});d.update(kw);return d
GOP=dict(mode='GOP',price_form='DON_GIA_CO_DINH',fee_pct='5',fee_ref='TEST phí')
class PackageDomain(unittest.TestCase):
    def setUp(self):
        self.s=multi();self.assertTrue(status(self.s)[2]['confirmed'],status(self.s)[2]['questions'])
        self.s=sc.save(self.s,{'packages':[pk('GT-01'),pk('GT-02'),pk('CC1','CC1 tự làm')],'pm_model':'VN_ND175'})
    def add(self,code,**kw):
        pid='id-'+code;self.s,res=sc.add_lines(self.s,dict(package_id=pid,**kw));return res
    def test_lines_keep_wbs_and_coverage(self):
        r=self.add('GT-01',by='WBS',wbs_prefix='1',alloc_pct='100')
        self.assertEqual(r['added'],['GT-01.1','GT-01.2'])
        l=self.s['packages'][0]['lines'][0];self.assertEqual(l['ab_line_id'],'test-source');self.assertEqual(l['snapshot']['wbs_code'],'1.1')
        self.assertEqual(l['snapshot']['lbs'],['TEST-LBS']);self.assertTrue(l['snapshot']['work_results'])
        self.add('GT-02',by='IDS',ids=['ab3','h1'],alloc_pct='60')
        q=[x['field'] for x in sc.questions(self.s,pmmodels.models())];self.assertIn('coverage.UNDER',q)
        r=self.add('CC1',by='IDS',ids=['ab3'],alloc='REMAINING');self.assertEqual(self.s['packages'][2]['lines'][0]['alloc_pct'],'40')
        cov=sc.coverage(self.s);self.assertTrue(all(c['state']=='OK' for c in cov.values()),cov)
        rec={r['ab_line_id']:r for r in sc.reconcile(self.s)};self.assertEqual(Decimal(rec['ab3']['bb_quantity']),Decimal(100))
        r2=self.add('GT-02',by='IDS',ids=['ab3'],alloc_pct='10');self.assertEqual(r2['skipped'][0]['reason'],'Đã có trong gói')
    def test_over_invalid_and_orphans(self):
        self.add('GT-01',by='IDS',ids=['ab2'],alloc_pct='80');self.add('GT-02',by='IDS',ids=['ab2'],alloc_pct='30')
        self.assertIn('coverage.OVER',[x['field'] for x in sc.questions(self.s)])
        with self.assertRaises(ValueError):sc.add_lines(self.s,dict(package_id='id-GT-01',by='IDS',ids=['ab3'],alloc_pct='0'))
        bad=copy.deepcopy(self.s['packages']);bad[0]['lines'][0]['alloc_pct']=None;s2=sc.save(copy.deepcopy(self.s),{'packages':bad})
        self.assertIn('coverage.INVALID',[x['field'] for x in sc.questions(s2)])
        for mutate,msg in [(lambda p:p[0]['lines'].append(dict(id='x',bb_code='GT-01.9',ab_line_id='nope',alloc_pct='1')),'mồ côi'),(lambda p:p[0].update(code='GT1'),'sai mẫu'),
                           (lambda p:p[1].update(code='GT-01'),'trùng'),(lambda p:p[0]['lines'][0].update(bb_code='GT-02.1'),'bắt đầu')]:
            p=copy.deepcopy(self.s['packages']);mutate(p)
            with self.assertRaises(ValueError) as e:sc.save(copy.deepcopy(self.s),{'packages':p})
            self.assertIn(msg,str(e.exception))
    def test_requires_wbs_step(self):
        s=multi();s['confirmed']=[0,1]
        with self.assertRaises(ValueError):sc.save(s,{'packages':[]})
    def test_delivery_gate_no_defaults(self):
        self.add('GT-01',by='IDS',ids=['test-source'],alloc_pct='100')
        q=' '.join(x['question'] for x in sc.questions(self.s,pmmodels.models()))
        for t in ('GỘP hay TÁCH','hình thức giá','phí','hình thức chọn nhà thầu phụ'):self.assertIn(t,q)
        p=self.s['packages'];p[0]['delivery']=dict(mode='TACH',VL='CC1',NC='CC1',M='CC1',price_form='DON_GIA_CO_DINH',fee_pct='0',fee_ref='TEST',shared=[{'resource_code':'X'}])
        self.s=sc.save(self.s,{'packages':p});q=' '.join(x['question'] for x in sc.questions(self.s))
        self.assertIn('không còn phạm vi',q);self.assertIn('nguồn lực dùng chung',q)
    def test_basis_a_b_c_never_fill(self):
        self.add('GT-01',by='WBS',wbs_prefix='1',alloc_pct='100');p=self.s['packages'][0]
        p['delivery']=dict(GOP);p['pricing']={'k_VL':'1.1'};self.s=sc.save(self.s,{'packages':self.s['packages']})
        b=sc.basis_values(self.s,self.s['packages'][0],sc.line_costs(self.s),[])
        self.assertIsNone(b['GT-01.1']['A']['value']);self.assertIn('D58',b['GT-01.1']['A']['reason'])
        self.assertIsNone(b['GT-01.1']['B']['value']);self.assertIn('NC, M',b['GT-01.1']['B']['reason'])
        self.assertIsNone(b['GT-01.2']['B']['value']);self.assertIn('chưa đủ',b['GT-01.2']['B']['reason'])
        self.s['packages'][0]['pricing']={'k_all':'0.9','k_VL':'1.1','k_NC':'1','k_M':'1'};b=sc.basis_values(self.s,self.s['packages'][0],sc.line_costs(self.s),[])
        self.assertEqual(Decimal(b['GT-01.1']['A']['value']),Decimal(90));self.assertEqual(Decimal(b['GT-01.2']['A']['value']),Decimal(45))
        self.assertEqual(Decimal(b['GT-01.1']['B']['value']),Decimal('0.17'))  # 30.03/200 × 1.1 × 1.05
        self.assertIsNone(b['GT-01.1']['C']['value'])
        self.s['packages'][0]['delivery']=dict(GOP,mode='TACH',VL='CC1',NC='NTP',M='NTP');b=sc.basis_values(self.s,self.s['packages'][0],sc.line_costs(self.s),[])
        self.assertIsNone(b['GT-01.1']['B']['value'],'NTP parts have no CC1 cost: never price them at 0')
    def quote(self,rows):
        out=io.StringIO();out.write('Mã dòng BB,Mã WBS,Mô tả,Đơn vị,Khối lượng,Đơn giá\n')
        for r in rows:out.write(','.join(r)+'\n')
        return sc.parse_quote('q.csv',out.getvalue().encode('utf-8-sig'))
    def test_review_finds_seeded_errors(self):
        self.add('GT-01',by='WBS',wbs_prefix='1',alloc_pct='100');p=self.s['packages'][0];p['delivery']=dict(GOP);p['pricing']={'k_all':'0.9'}
        self.s=sc.save(self.s,{'packages':self.s['packages']})
        for v,rows in (('NTP-1',[['GT-01.1','1.1','x','m3','200','95'],['GT-99.1','9','extra','m','1','1']]),
                       ('NTP-2',[['GT-01.1','1.1','x','m2','200','95'],['GT-01.2','1.2','x','t','12','']]),
                       ('NTP-3',[['GT-01.1','1.1','x','m³','200','200'],['GT-01.2','1.2','x','t','10','44']])):
            self.s=sc.add_quote(self.s,{'package_id':p['id'],'vendor':v,'filename':'q.csv'},'0'*64,self.quote(rows))
        sm={x['code']:x for x in sc.summary(self.s)['packages']}['GT-01'];rv={r['vendor']:r for r in sm['reviews']}
        self.assertEqual((rv['NTP-1']['counts']['MISSING'],rv['NTP-1']['counts']['EXTRA'],rv['NTP-1']['valid']),(1,1,False))
        self.assertEqual((rv['NTP-2']['counts']['UOM_MISMATCH'],rv['NTP-2']['counts']['PRICE_EMPTY'],rv['NTP-2']['counts']['QTY_DIFF']),(1,1,1))
        self.assertTrue(rv['NTP-3']['valid']);self.assertEqual(rv['NTP-3']['counts']['PRICE_HIGH'],1);self.assertEqual(len(rv['NTP-3']['findings']),1,'correct line raises nothing')
        self.assertEqual(sm['basis']['GT-01.1']['C']['value'],'200');self.assertEqual(sm['basis']['GT-01.2']['C']['value'],'44')
        # D79 AC16: blank ≠ zero; a row without its BB line id is not matched by description
        for v,rows in (('NTP-4',[['GT-01.1','1.1','x','m3','200','0'],['','1.2','TEST ONLY rebar','t','10','44']]),):
            self.s=sc.add_quote(self.s,{'package_id':p['id'],'vendor':v,'filename':'q.csv'},'1'*64,self.quote(rows))
        r4=next(r for r in next(x for x in sc.summary(self.s)['packages'] if x['code']=='GT-01')['reviews'] if r['vendor']=='NTP-4')
        self.assertEqual((r4['counts']['PRICE_NOT_POSITIVE'],r4['counts']['PRICE_EMPTY'],r4['counts']['MISSING'],r4['counts']['EXTRA']),(1,0,1,1))
        q3=next(q for q in self.s['quotes'] if q['vendor']=='NTP-3')
        self.s=sc.save_award(self.s,{'package_id':p['id'],'quote_id':q3['id'],'proposed_by':'TEST'})
        self.assertIn('award',' '.join(x['field'] for x in sc.questions(self.s)));self.assertEqual(self.s['packages'][0]['award']['prices']['GT-01.2'],'44')
        self.s=sc.save_award(self.s,{'package_id':p['id'],'quote_id':q3['id'],'proposed_by':'TEST','reason':'TEST thấp nhất hợp lệ'})
        self.assertFalse([x for x in sc.questions(self.s) if 'award' in x['field']]);self.assertEqual(self.s['packages'][0]['award']['status'],'DRAFT')
        with self.assertRaises(ValueError):sc.save(self.s,{'packages':self.s['packages'][1:]})
    def test_progress_rolls_up_to_owner_wbs(self):
        self.add('GT-01',by='IDS',ids=['test-source'],alloc_pct='100');self.add('GT-02',by='IDS',ids=['ab3'],alloc_pct='60');self.add('CC1',by='IDS',ids=['ab3'],alloc='REMAINING')
        for code,q,d in (('GT-01.1','50','2026-10-01'),('GT-01.1','100','2026-10-20'),('GT-02.1','70','2026-10-20')):
            self.s=sc.add_progress(self.s,{'bb_code':code,'cum_qty':q,'date':d,'ref':'TEST biên bản'})
        with self.assertRaises(ValueError):sc.add_progress(self.s,{'bb_code':'GT-01.1','cum_qty':'1','date':'2026-10-21'})
        r=sc.rollup(self.s);ln={x['ab_line_id']:x for x in r['lines']};w={x['wbs']:x for x in r['wbs']}
        self.assertEqual(ln['test-source']['pct'],'50');self.assertEqual(ln['ab3']['pct'],'60');self.assertEqual(w['1.1']['pct'],'50');self.assertEqual(w['1.1']['method'],'Theo giá trị hợp đồng AB')
        self.assertIsNone(w['2.1']['pct']);self.assertIn('PRD §5.3',w['2.1']['method']);self.assertTrue(r['warnings'],'70 > 60 allocated is capped and warned')  # D79: no average of line %
    def test_award_pins_digest_and_goes_stale(self):  # D79 FR28
        self.add('GT-01',by='IDS',ids=['ab2'],alloc_pct='100')
        self.s=sc.save_award(self.s,{'package_id':'id-GT-01','vendor':'TEST NTP','reason':'TEST','proposed_by':'TEST','prices':{'GT-01.1':'45'}})
        p=self.s['packages'][0];self.assertEqual(len(p['award']['inputs_sha256']),64);self.assertFalse(sc.award_stale(self.s,p))
        next(r for r in self.s['boq'] if r['id']=='ab2')['unit_price']='55'
        self.assertTrue(sc.award_stale(self.s,p));self.assertIn('award.GT-01.stale',[x['field'] for x in sc.questions(self.s)])
        self.s=sc.save_award(self.s,{'package_id':'id-GT-01','vendor':'TEST NTP','reason':'TEST','proposed_by':'TEST','prices':{'GT-01.1':'45'}})
        self.assertFalse(sc.award_stale(self.s,self.s['packages'][0]))
    def test_input_change_reopens_packages(self):
        self.add('GT-01',by='IDS',ids=['ab2'],alloc_pct='100');self.assertFalse(sc.stale_lines(self.s))
        next(r for r in self.s['boq'] if r['id']=='ab2')['quantity']='12'
        st=sc.stale_lines(self.s);self.assertEqual(st[0]['changed'],['quantity']);self.assertIn('packages.stale',[x['field'] for x in sc.questions(self.s)])
class LibraryDomain(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.st=Store(Path(self.tmp.name)/'db.sqlite')
        rows=[dict(id='p1',resource_code='VL.CAT',spec='',uom='m3',unit_price='100',location='DN',period='2026-09',currency='VND'),dict(id='p2',resource_code='NC.3/7',uom='công',unit_price='300000',location='DN',period='2026-09',currency='VND'),
              dict(id='p3',resource_code='M101.0101',uom='ca',unit_price=None,location='DN',period='2026-09',currency='VND'),dict(id='p4',resource_code='M102.0001',uom='ca',unit_price='1000',location='DN',period='2026-09',currency='VND')]
        self.src=self.st.library_add('PRICES','TEST ONLY nhà nước','e'*64,'t.csv','',rows)
    def tearDown(self):self.tmp.cleanup()
    def test_derive_factors_overrides_lineage(self):
        before=copy.deepcopy(self.st.library_rows(self.src['id']))
        r=libtools.derive(self.st,self.src['id'],{'name':'TEST CC1','basis_ref':'TEST','factors':{'VL':'1.05','M':'2','prefixes':{'M102':'1.5','ZZ':'2'}},'overrides':[{'row_id':'p2','value':'350000','reason':'TEST khảo sát'}]})
        rows={x['id']:x for x in self.st.library_rows(r['dataset']['id'])}
        self.assertEqual(rows['p1']['unit_price'],'105');self.assertEqual(rows['p4']['unit_price'],'3000');self.assertIsNone(rows['p3']['unit_price'],'missing stays missing')
        self.assertEqual(rows['p2']['unit_price'],'350000');self.assertEqual(rows['p2']['derived']['override_original'],'300000');self.assertEqual(rows['p1']['derived']['from_sha256'],self.src['sha256'])
        self.assertEqual(r['stats']['null_kept'],1);self.assertEqual(r['stats']['unmatched_prefixes'],['ZZ']);self.assertEqual(self.st.library_rows(self.src['id']),before)
        for body in ({'name':'X','basis_ref':'T','factors':{'VL':'0'}},{'name':'X','basis_ref':'T'},{'name':'TEST ONLY nhà nước','basis_ref':'T','factors':{'all':'1.1'}},
                     {'name':'X','basis_ref':'T','overrides':[{'row_id':'p1','value':'1'}]},{'name':'X','factors':{'all':'1.1'}}):
            with self.assertRaises(ValueError):libtools.derive(self.st,self.src['id'],body)
    def test_activate_old_version(self):
        v2=self.st.library_add('PRICES','TEST ONLY nhà nước','f'*64,'t.csv','',[]);self.st.library_archive(self.src['id'])
        self.st.library_activate(self.src['id']);by={d['id']:d['status'] for d in self.st.library_list()}
        self.assertEqual((by[self.src['id']],by[v2['id']]),('ACTIVE','ARCHIVED'))
    def test_dictionary_export_and_versioned_import(self):
        wb=load_workbook(io.BytesIO(libtools.dictionaries_workbook(self.st,multi())),read_only=True)
        for n in ('00_HUONG_DAN','WORK_RESULT','CBS_SCHEMES','CBS_CC1','LBS','FAMILIES','BIEN_PHAP','MO_HINH_QLDA','P_WBS','P_MAPPING','P_GOI_THAU'):self.assertIn(n,wb.sheetnames)
        self.assertEqual(len(list(wb['LBS'].iter_rows()))-1,44)
        csvdata=lambda rows:base64.b64encode(('code,name,work_group,norm_param,applies_to,note,status\n'+'\n'.join(rows)).encode()).decode()
        d=libtools.import_dictionary(self.st,{'sheet':'BIEN_PHAP','filename':'b.csv','data':csvdata(['THU_CONG,Thủ công,BT,,AF.1,,DRAFT','BOM,Bơm,BT,,AF.3,,DRAFT'])})
        self.assertEqual(d['counts']['added'],2);self.assertFalse([x for x in self.st.library_list() if x['name']==libtools.METHOD_NAME],'dry run never writes')
        libtools.import_dictionary(self.st,{'sheet':'BIEN_PHAP','filename':'b.csv','dry_run':False,'data':csvdata(['THU_CONG,Thủ công,BT,,AF.1,,DRAFT','BOM,Bơm,BT,,AF.3,,DRAFT'])})
        r=libtools.import_dictionary(self.st,{'sheet':'BIEN_PHAP','filename':'b.csv','dry_run':False,'data':csvdata(['THU_CONG,Thủ công sửa,BT,,AF.1,,DRAFT','CAN_CAU,Cẩu,BT,,AF.2,,DRAFT'])})
        self.assertEqual(r['counts'],{'added':1,'removed':1,'changed':1,'rows':2})
        vs=sorted((d['version'],d['status']) for d in self.st.library_list() if d['name']==libtools.METHOD_NAME);self.assertEqual(vs,[(1,'ARCHIVED'),(2,'ACTIVE')])
        dup=libtools.import_dictionary(self.st,{'sheet':'BIEN_PHAP','filename':'b.csv','data':csvdata(['A,x,,,,,','A,y,,,,,'])});self.assertTrue(dup['error_count'])
        with self.assertRaises(ValueError):libtools.import_dictionary(self.st,{'sheet':'WORK_RESULT','filename':'b.csv','data':csvdata(['x'])})
        with self.assertRaises(ValueError):libtools.import_dictionary(self.st,{'sheet':'BIEN_PHAP','filename':'b.csv','dry_run':False,'data':csvdata(['A,x,,,,,','A,y,,,,,'])})
    def test_price_upload_validation(self):
        ok=[dict(id='1',resource_code='VL.X',spec='',uom='m3',unit_price='10',location='DN',period='2026-10',currency='VND')]
        self.assertEqual(libtools.validate_prices(ok),[])
        e=libtools.validate_prices(ok+[dict(ok[0],id='2'),dict(ok[0],id='3',resource_code='VL.Y',unit_price='',period='10/2026')])
        self.assertTrue(any('trùng khóa' in x for x in e) and any('số dương' in x for x in e) and any('YYYY-MM' in x for x in e))
        self.assertIn('Thiếu cột',libtools.validate_prices([{'id':'1','resource_code':'X'}])[0])
class DrawingsAndModels(unittest.TestCase):
    def dxf(self):
        def e(*pairs):return ''.join(f'{c}\n{v}\n' for c,v in pairs)
        body=e((0,'SECTION'),(2,'HEADER'),(9,'$INSUNITS'),(70,6),(0,'ENDSEC'),(0,'SECTION'),(2,'ENTITIES'),
               (0,'LINE'),(5,'A1'),(8,'COC'),(10,0),(20,0),(11,10),(21,0),
               (0,'LWPOLYLINE'),(5,'B1'),(8,'MONG'),(90,4),(70,1),(10,0),(20,0),(10,2),(20,0),(10,2),(20,2),(10,0),(20,2))
        for h in ('C1','C2','C3'):body+=e((0,'INSERT'),(5,h),(8,'COC'),(2,'COC_D1000'),(10,1),(20,1))
        body+=e((0,'TEXT'),(5,'D1'),(8,'GHICHU'),(10,0),(20,0),(1,'Bê tông lót 12,5 m3'),(0,'ENDSEC'),(0,'EOF'))
        return body.encode()
    def test_dxf_candidates(self):
        c=drawings.dxf_candidates(self.dxf());k={(x['kind'],x['source'].get('layer')):x for x in c}
        self.assertEqual(Decimal(k[('DXF_LENGTH','COC')]['value']),10);self.assertEqual(Decimal(k[('DXF_LENGTH','MONG')]['value']),8);self.assertEqual(Decimal(k[('DXF_AREA','MONG')]['value']),4)
        self.assertEqual(k[('DXF_BLOCK','COC')]['value'],'3');self.assertEqual((k[('DXF_TEXT','GHICHU')]['value'],k[('DXF_TEXT','GHICHU')]['uom']),('12.5','m3'));self.assertIn('(m)',k[('DXF_LENGTH','COC')]['uom'])
        with self.assertRaises(ValueError):drawings.kind_of('x.dwg')
    def test_pdf_assign_compare_never_changes_boq(self):
        import fitz
        doc=fitz.open();page=doc.new_page();page.insert_text((72,72),'Be tong lot mong 150,5 m3',fontsize=11);data=doc.tobytes()
        s=multi();before=copy.deepcopy(s['boq']);s,res=drawings.add(s,{'filename':'mong.pdf','label':'TEST'},hashlib.sha256(data).hexdigest(),data)
        t=next(x for x in s['takeoffs'] if x['kind']=='PDF_TEXT');self.assertEqual((t['value'],t['uom'],t['source']['page']),('150.5','m3',1))
        with self.assertRaises(ValueError):drawings.assign(s,{'takeoff_id':t['id'],'status':'CONFIRMED','ab_line_id':'test-source','factor':'1'})
        with self.assertRaises(ValueError):drawings.assign(s,{'takeoff_id':t['id'],'status':'CONFIRMED','ab_line_id':'h1','factor':'1','factor_ref':'x'})
        s=drawings.assign(s,{'takeoff_id':t['id'],'status':'CONFIRMED','ab_line_id':'test-source','factor':'1','factor_ref':'TEST cùng đơn vị'})
        cmp=drawings.compare(s)[0];self.assertEqual(cmp['deviation_pct'],'-24.8');self.assertEqual(s['boq'],before)
    def test_pm_model_gate_pmbok(self):
        """D71: a project still on retired PMBOK 7 must re-choose; PMBOK 8 without a source must not pass silently."""
        ms=pmmodels.models();base=sc.save(multi(),{'packages':[pk('GT-01'),pk('GT-02'),pk('CC1','CC1 tự làm')],'pm_model':'VN_ND175'})
        for pm,needle in (('PMBOK7_ISO21502','không còn dùng'),('PMBOK8','CẦN BỔ SUNG')):
            s=copy.deepcopy(base);s['context']['pm_model']=pm
            q=[x for x in sc.questions(s,ms) if x['field']=='context.pm_model'];self.assertTrue(q and needle in q[0]['question'],(pm,q))
        s=copy.deepcopy(base);s['context']['pm_model']='PMBOK6'
        self.assertFalse([x for x in sc.questions(s,ms) if x['field']=='context.pm_model'])
    def test_models_are_draft(self):
        ms=pmmodels.models()
        self.assertEqual(set(ms),{'VN_ND175','FIDIC_RED','FIDIC_YELLOW','PRINCE2','PMBOK6','PMBOK8'})  # D71: PMBOK 6 + 8, not PMBOK 7
        for k,m in ms.items():
            self.assertEqual(m['status'],'DRAFT_NEEDS_LEGAL_REVIEW');self.assertTrue(m['checklist'])
            if k!='PMBOK8':self.assertTrue(m['procurement_methods'])
        self.assertEqual(ms['PMBOK6']['stages'],['Khởi tạo (Initiating)','Lập kế hoạch (Planning)','Thực hiện (Executing)','Giám sát và kiểm soát (Monitoring and Controlling)','Kết thúc (Closing)']);self.assertEqual(len(ms['PMBOK6']['knowledge_areas']),10)
        self.assertEqual(ms['PMBOK8']['content'],'NEEDS_SOURCE');self.assertEqual(ms['PMBOK8']['procurement_methods'],{})  # not guessed
        self.assertNotIn('PMBOK 7',json.dumps({k:v['name'] for k,v in ms.items()}))
        s=multi();cl=pmmodels.checklist(s,ms['VN_ND175'],status(s));st={c['code']:c['state'] for c in cl}
        self.assertEqual((st['HS-03'],st['HS-06'],st['GT-01']),('DONE','MISSING','MISSING'))
class SubcontractHttp(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.app=App(('127.0.0.1',0),Path(self.tmp.name)/'db.sqlite');threading.Thread(target=self.app.serve_forever,daemon=True).start()
        self.base=f'http://127.0.0.1:{self.app.server_port}';self.call('/api/bootstrap')
        f=multi();s=self.app.store.create('TEST ONLY subcontract');s.update({k:f[k] for k in ('context','boq','wbs','mapping','applications','norms','resources','prices','adjustments','export','confirmed')})
        self.s=self.app.store.put(s,s['revision'],'TEST_FIXTURE')
    def tearDown(self):self.app.shutdown();self.app.server_close();self.tmp.cleanup()
    def call(self,path,body=None,raw=False):
        req=urllib.request.Request(self.base+path,data=None if body is None else json.dumps(body).encode(),headers={} if body is None else {'Content-Type':'application/json','X-CC1-Session':self.app.token})
        with urllib.request.urlopen(req) as r:return r.read() if raw else json.load(r)
    def test_full_flow_and_packs(self):
        pid=self.s['id'];v=self.call(f'/api/projects/{pid}/subcontract');self.assertTrue(v['wbs_ready']);self.assertIn('VN_ND175',v['models'])
        with self.assertRaises(urllib.error.HTTPError) as e:self.call(f'/api/projects/{pid}/packages',{'revision':v['revision']-1,'packages':[]})
        self.assertEqual(e.exception.code,409)
        v=self.call(f'/api/projects/{pid}/packages',{'revision':v['revision'],'pm_model':'VN_ND175','packages':[pk('GT-01',procurement_method='CHI_DINH_THAU',delivery=dict(GOP),pricing={'k_all':'0.9'}),pk('CC1','CC1 tự làm')]})
        v=self.call(f'/api/projects/{pid}/packages/add-lines',{'revision':v['revision'],'package_id':'id-GT-01','by':'WBS','wbs_prefix':'1','alloc_pct':'100'});self.assertEqual(v['result']['added'],['GT-01.1','GT-01.2'])
        v=self.call(f'/api/projects/{pid}/packages/add-lines',{'revision':v['revision'],'package_id':'id-CC1','by':'IDS','ids':['ab3'],'alloc':'REMAINING'})
        self.assertEqual(v['summary']['coverage'],{'OK':3,'UNDER':0,'OVER':0,'INVALID':0});self.assertFalse(v['summary']['questions'],v['summary']['questions'])
        bb=load_workbook(io.BytesIO(self.call(f'/api/projects/{pid}/bb/GT-01.xlsx',raw=True)))['BOQ_BB'];hdr=[c.value for c in bb[1]]
        self.assertEqual(hdr[:8],['Mã dòng BB','Mã WBS','Work Result','LBS','Mô tả','Đơn vị','Khối lượng','Đơn giá']);self.assertNotIn('Giá gợi ý A (AB×K)',hdr)
        self.assertEqual([bb.cell(2,1).value,bb.cell(2,2).value],['GT-01.1','1.1'])
        out=io.BytesIO();w=load_workbook(io.BytesIO(self.call(f'/api/projects/{pid}/bb/GT-01.xlsx',raw=True)));ws=w['BOQ_BB'];ws.cell(2,8).value=88;ws.cell(3,8).value=45;w.save(out)
        v=self.call(f'/api/projects/{pid}/quotes',{'revision':v['revision'],'package_id':'id-GT-01','vendor':'TEST NTP','filename':'q.xlsx','data':base64.b64encode(out.getvalue()).decode()})
        rv=v['summary']['packages'][0]['reviews'][0];self.assertTrue(rv['valid'],rv);qid=v['quotes'][0]['id']
        v=self.call(f'/api/projects/{pid}/award',{'revision':v['revision'],'package_id':'id-GT-01','quote_id':qid,'reason':'TEST','proposed_by':'TEST'})
        v=self.call(f'/api/projects/{pid}/progress',{'revision':v['revision'],'bb_code':'GT-01.2','cum_qty':'10','date':'2026-10-20','ref':'TEST BB'})
        self.assertEqual({w['wbs']:w['state'] for w in v['summary']['progress']['wbs']}['1.2'],'XONG')
        z=zipfile.ZipFile(io.BytesIO(self.call(f'/api/projects/{pid}/procurement-pack.zip',raw=True)));man=json.loads(z.read('MANIFEST.json'))
        self.assertIn('BOQ_BB/GT-01_GUI_NHA_THAU_PHU.xlsx',z.namelist());self.assertNotIn('BOQ_BB/CC1_GUI_NHA_THAU_PHU.xlsx',z.namelist())
        for f in man['files']:self.assertEqual(hashlib.sha256(z.read(f['file'])).hexdigest(),f['sha256'])
        wb=load_workbook(io.BytesIO(z.read('01_HO_SO_CHON_THAU_GIAO_THAU.xlsx')))
        for n in ('04_PROCUREMENT_PACKAGE_PLAN','GOI_THAU','SO_SANH_GIA','THAM_DINH_BAO_GIA','DE_XUAT_GIAO_THAU','PHU_KHOI_LUONG','DOI_CHIEU_WBS_AB_BB','HOAN_THANH_THEO_WBS','CHECKLIST_QLDA'):self.assertIn(n,wb.sheetnames)
        dw={r[0]:r for r in wb['DOI_CHIEU_WBS_AB_BB'].iter_rows(min_row=2,values_only=True)};self.assertEqual(dw['2.1'][5],0)
        t=zipfile.ZipFile(io.BytesIO(self.call(f'/api/projects/{pid}/tender-pack.zip',raw=True)));self.assertIn('01_HO_SO_LAP_DU_TOAN_TONG_HOP.xlsx',t.namelist())
        self.assertIn('DỰ THẢO',str(list(load_workbook(io.BytesIO(t.read('01_HO_SO_LAP_DU_TOAN_TONG_HOP.xlsx')))['TRANG_THAI'].values)))
        d=load_workbook(io.BytesIO(self.call(f'/api/dictionaries.xlsx?project={pid}',raw=True)),read_only=True);self.assertIn('P_GOI_THAU',d.sheetnames)
        self.assertEqual(self.call('/api/pm-models')['status'],'DRAFT_NEEDS_LEGAL_REVIEW')
        audit=[e['action'] for e in self.app.store.events(pid)];self.assertIn('SAVE_AWARD_DRAFT',audit);self.assertIn('ADD_SUBCONTRACT_PROGRESS',audit)
    def test_library_routes(self):
        src=self.app.store.library_add('PRICES','TEST ONLY nhà nước','e'*64,'t.csv','',[dict(id='p1',resource_code='VL.X',uom='m3',unit_price='10')])
        r=self.call(f'/api/library/{src["id"]}/derive',{'name':'TEST CC1','basis_ref':'TEST','factors':{'all':'1.2'}});self.assertEqual(r['stats']['changed'],1)
        r=self.call(f'/api/library/{src["id"]}/activate',{});self.assertEqual(r['dataset']['status'],'ACTIVE')
if __name__=='__main__':unittest.main()
